"""Filter → group → split export engine.

Consumes completed child result.json canonical lines. It never opens source
workbooks and never runs Reader / Mapping / Normalize / Business / Contra /
Anomaly.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
import zipfile
from collections import defaultdict
from dataclasses import dataclass, replace
from datetime import date
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .export import HEADER_FILL, _write_sheet, safe_excel_text
from .summary_structure import (
    SemanticType,
    SummaryStructureRule,
    extract_derived_fields,
)

SPLIT_SCHEMA_VERSION = 1

COLUMN_GROUPS = (
    ("core", "财务核心字段"),
    ("aux", "往来与辅助核算"),
    ("other", "其他业务字段"),
    ("source", "来源与系统追溯"),
    ("advanced", "高级系统字段"),
)
COLUMN_DEFINITIONS = (
    {"id": "date", "label": "日期", "group": "core", "default": True, "kind": "date"},
    {"id": "voucher_id", "label": "凭证号", "group": "core", "default": True, "kind": "text"},
    {"id": "description", "label": "摘要", "group": "core", "default": True, "kind": "text"},
    {"id": "account_code", "label": "科目编码", "group": "core", "default": True, "kind": "text"},
    {"id": "account_name", "label": "科目名称", "group": "core", "default": True, "kind": "text"},
    {"id": "detail_account", "label": "明细科目", "group": "core", "default": True, "kind": "text"},
    {"id": "debit", "label": "借方金额", "group": "core", "default": True, "kind": "amount"},
    {"id": "credit", "label": "贷方金额", "group": "core", "default": True, "kind": "amount"},
    {"id": "customer_code", "label": "客户编码", "group": "aux", "default_if_data": True, "kind": "text"},
    {"id": "customer_name", "label": "客户名称", "group": "aux", "default_if_data": True, "kind": "text"},
    {"id": "supplier_code", "label": "供应商编码", "group": "aux", "default_if_data": True, "kind": "text"},
    {"id": "supplier_name", "label": "供应商名称", "group": "aux", "default_if_data": True, "kind": "text"},
    {"id": "counterparty_code", "label": "往来单位编码", "group": "aux", "default_if_data": True, "kind": "text"},
    {"id": "counterparty_name", "label": "往来单位名称", "group": "aux", "default_if_data": True, "kind": "text"},
    {"id": "entity", "label": "主体", "group": "other", "default_if_data": True, "kind": "text"},
    {"id": "period", "label": "会计期间", "group": "other", "default_if_data": True, "kind": "text"},
    {"id": "currency", "label": "币种", "group": "other", "default_if_data": True, "kind": "text"},
    {"id": "original_amount", "label": "原币金额", "group": "other", "default_if_data": True, "kind": "amount"},
    {"id": "preparer", "label": "制单人", "group": "other", "default_if_data": True, "kind": "text"},
    {"id": "original_filename", "label": "原始文件", "group": "source", "kind": "text"},
    {"id": "sheet_name", "label": "源 Sheet", "group": "source", "kind": "text"},
    {"id": "source_row", "label": "源行号", "group": "source", "kind": "number"},
    {"id": "line_id", "label": "外部行 ID", "group": "advanced", "advanced": True, "kind": "text"},
    {"id": "internal_line_key", "label": "内部行标识（系统追溯）", "group": "advanced", "advanced": True, "kind": "key"},
    {"id": "job_id", "label": "任务 ID（系统字段）", "group": "advanced", "advanced": True, "kind": "key"},
    {"id": "voucher_key", "label": "凭证内部标识（系统字段）", "group": "advanced", "advanced": True, "kind": "key"},
)
COLUMN_BY_ID = {item["id"]: item for item in COLUMN_DEFINITIONS}
DEFAULT_COLUMN_ORDER = (
    "date", "voucher_id", "description", "account_code", "account_name",
    "detail_account", "customer_code", "customer_name", "supplier_code",
    "supplier_name", "counterparty_code", "counterparty_name", "debit",
    "credit", "currency", "original_amount", "preparer", "entity", "period",
    "original_filename", "sheet_name", "source_row", "line_id",
    "internal_line_key", "job_id", "voucher_key",
)
DEFAULT_MAX_GROUPS_MULTI_SHEET = 200
DEFAULT_MAX_GROUPS_ZIP = 500
DEFAULT_MAX_ROWS_PER_SHEET = 1_048_000
MAX_EXPLICIT_ROWS = 5000
DEFAULT_PAGE_SIZE = 50
MAX_PAGE_SIZE = 200
ROW_SELECTION_MODES = ("all_except", "only_selected")
BLANK_LABELS = {
    "account": "未指定科目",
    "customer": "未指定客户",
    "supplier": "未指定供应商",
    "counterparty": "未指定往来单位",
    "month": "未指定月份",
    "period": "未指定会计期间",
    "entity": "未指定主体",
    "source_file": "未指定来源文件",
}
GROUP_BY_VALUES = (
    "account", "customer", "supplier", "counterparty",
    "month", "period", "entity", "source_file",
)
GROUP_BY_LABELS = {
    "account": "科目", "customer": "客户", "supplier": "供应商",
    "counterparty": "往来单位", "month": "月份", "period": "会计期间",
    "entity": "主体", "source_file": "来源文件",
}
ILLEGAL_SHEET = re.compile(r"[\[\]:*?/\\]")
ILLEGAL_FILENAME = re.compile(r'[\u0000-\u001f\u007f/\\:*?"<>|]+')


class SplitExportError(Exception):
    pass


@dataclass
class Record:
    job_id: str
    filename: str
    position: int
    voucher_key: str
    voucher_id: str
    date: Optional[str]
    period: Optional[str]
    entity: Optional[str]
    line: Dict[str, Any]


def _decimal(value: Any) -> Decimal:
    if value in (None, ""):
        return Decimal("0.00")
    try:
        return Decimal(str(value)).quantize(Decimal("0.01"))
    except (InvalidOperation, ValueError):
        return Decimal("0.00")


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _safe_excel(value: Any) -> Any:
    return safe_excel_text(value)


def _parse_date(value: Any) -> Optional[date]:
    text = _text(value)
    if not text:
        return None
    try:
        return date.fromisoformat(text[:10])
    except ValueError:
        return None


def load_records(
    jobs: Sequence[Dict[str, Any]], collect_rules: bool = False
):
    records: List[Record] = []
    rules: Dict[Tuple[str, str], SummaryStructureRule] = {}
    for job in jobs:
        path = Path(str(job.get("result_path") or ""))
        if not path.exists():
            raise SplitExportError(f"结果文件不存在：{path.name}")
        try:
            payload = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, ValueError) as exc:
            raise SplitExportError(f"结果文件无法读取：{path.name}") from exc
        if not isinstance(payload, dict) or not isinstance(payload.get("vouchers"), list):
            raise SplitExportError(f"结果文件结构不完整：{path.name}")
        job_id = str(job.get("job_id") or "")
        filename = str(job.get("original_filename") or "")
        position = int(job.get("position") or 0)
        summary_structure = payload.get("summary_structure") or {}
        for item in summary_structure.get("rules") or []:
            if not isinstance(item, dict):
                continue
            rule = SummaryStructureRule.from_dict(item)
            if rule.account_code:
                rules[(job_id, rule.account_code)] = rule
        for voucher in payload.get("vouchers") or []:
            if not isinstance(voucher, dict):
                continue
            voucher_key = str(voucher.get("voucher_key") or voucher.get("voucher_id") or "")
            voucher_id = str(voucher.get("voucher_id") or "")
            for line in voucher.get("lines") or []:
                if not isinstance(line, dict):
                    continue
                records.append(Record(
                    job_id=job_id,
                    filename=filename,
                    position=position,
                    voucher_key=voucher_key,
                    voucher_id=voucher_id,
                    date=_text(line.get("date") or voucher.get("date")) or None,
                    period=_text(line.get("period") or voucher.get("period")) or None,
                    entity=_text(line.get("entity") or voucher.get("entity")) or None,
                    line=line,
                ))
    if collect_rules:
        return records, rules
    return records


def _list(value: Any) -> List[str]:
    if value in (None, ""):
        return []
    if isinstance(value, (list, tuple)):
        return [str(item).strip() for item in value if str(item).strip()]
    return [str(value).strip()] if str(value).strip() else []


def _matches(value: Any, values: Sequence[str]) -> bool:
    if not values:
        return True
    text = _text(value)
    return text in values if text else False


def _matches_pair(code: Any, name: Any, values: Sequence[str]) -> bool:
    if not values:
        return True
    return _text(code) in values or _text(name) in values


def _min_max_ok(value: Decimal, minimum: Any, maximum: Any) -> bool:
    if minimum not in (None, ""):
        try:
            if value < Decimal(str(minimum)):
                return False
        except InvalidOperation:
            raise SplitExportError("金额筛选值无效。")
    if maximum not in (None, ""):
        try:
            if value > Decimal(str(maximum)):
                return False
        except InvalidOperation:
            raise SplitExportError("金额筛选值无效。")
    return True


def apply_filters(records: Sequence[Record], filters: Dict[str, Any]) -> List[Record]:
    filters = filters or {}
    source_files = _list(filters.get("source_files"))
    entities = _list(filters.get("entities"))
    months = _list(filters.get("months"))
    periods = _list(filters.get("periods"))
    account_codes = _list(filters.get("account_codes"))
    account_names = _list(filters.get("account_names"))
    account_values = _list(filters.get("account_values"))
    customer_codes = _list(filters.get("customer_codes"))
    customer_names = _list(filters.get("customer_names"))
    customer_values = _list(filters.get("customer_values"))
    supplier_codes = _list(filters.get("supplier_codes"))
    supplier_names = _list(filters.get("supplier_names"))
    supplier_values = _list(filters.get("supplier_values"))
    counterparty_codes = _list(filters.get("counterparty_codes"))
    counterparty_names = _list(filters.get("counterparty_names"))
    counterparty_values = _list(filters.get("counterparty_values"))
    date_from = _parse_date(filters.get("date_from"))
    date_to = _parse_date(filters.get("date_to"))
    if filters.get("date_from") and date_from is None:
        raise SplitExportError("起始日期无效。")
    if filters.get("date_to") and date_to is None:
        raise SplitExportError("结束日期无效。")

    result = []
    for record in records:
        line = record.line
        if source_files and record.filename not in source_files:
            continue
        if entities and _text(record.entity) not in entities:
            continue
        record_date = _parse_date(record.date)
        if (date_from or date_to) and record_date is None:
            continue
        if date_from and record_date < date_from:
            continue
        if date_to and record_date > date_to:
            continue
        month = record_date.isoformat()[:7] if record_date else ""
        if months and month not in months:
            continue
        if periods and _text(record.period) not in periods:
            continue
        if account_codes and not _matches(line.get("account_code"), account_codes):
            continue
        if account_names and not _matches(line.get("account_name"), account_names):
            continue
        if account_values and not (
            _matches(line.get("account_code"), account_values)
            or _matches(line.get("account_name"), account_values)
        ):
            continue
        if customer_codes and not _matches(line.get("customer_code"), customer_codes):
            continue
        if customer_names and not _matches(line.get("customer_name"), customer_names):
            continue
        if customer_values and not _matches_pair(
            line.get("customer_code"), line.get("customer_name"), customer_values
        ):
            continue
        if supplier_codes and not _matches(line.get("supplier_code"), supplier_codes):
            continue
        if supplier_names and not _matches(line.get("supplier_name"), supplier_names):
            continue
        if supplier_values and not _matches_pair(
            line.get("supplier_code"), line.get("supplier_name"), supplier_values
        ):
            continue
        if counterparty_codes and not _matches(line.get("counterparty_code"), counterparty_codes):
            continue
        if counterparty_names and not _matches(line.get("counterparty_name"), counterparty_names):
            continue
        if counterparty_values and not _matches_pair(
            line.get("counterparty_code"), line.get("counterparty_name"), counterparty_values
        ):
            continue
        if not _min_max_ok(_decimal(line.get("debit")), filters.get("debit_min"), filters.get("debit_max")):
            continue
        if not _min_max_ok(_decimal(line.get("credit")), filters.get("credit_min"), filters.get("credit_max")):
            continue
        result.append(record)
    return result


def row_identity(record: Record) -> Tuple[str, str]:
    key = _text(record.line.get("internal_line_key"))
    if not key:
        raise SplitExportError("结果缺少稳定行标识，无法进行行选择。")
    return record.job_id, key


def _row_ref(value: Any) -> Tuple[str, str]:
    if not isinstance(value, dict):
        raise SplitExportError("行选择格式不正确。")
    job_id = _text(value.get("job_id") or value.get("jobId"))
    key = _text(value.get("internal_line_key") or value.get("internalLineKey"))
    if not key:
        raise SplitExportError("行选择缺少内部行标识。")
    if len(key) > 200:
        raise SplitExportError("行标识过长。")
    if len(job_id) > 100:
        raise SplitExportError("行选择任务标识过长。")
    return job_id, key


def normalize_row_selection(request: Dict[str, Any]) -> Dict[str, Any]:
    mode = _text(
        request.get("row_selection_mode")
        or request.get("selection_mode")
        or "all_except"
    ) or "all_except"
    if mode not in ROW_SELECTION_MODES:
        raise SplitExportError("行选择模式无效。")

    def collect(raw: Any, field: str) -> List[Tuple[str, str]]:
        if raw in (None, ""):
            return []
        if not isinstance(raw, (list, tuple)):
            raise SplitExportError(f"{field} 必须是数组。")
        if len(raw) > MAX_EXPLICIT_ROWS:
            raise SplitExportError(f"手工{field}超过 {MAX_EXPLICIT_ROWS} 行上限，请改用筛选条件。")
        refs: List[Tuple[str, str]] = []
        seen = set()
        for item in raw:
            ref = _row_ref(item)
            if ref in seen:
                continue
            seen.add(ref)
            refs.append(ref)
        return refs

    excluded = collect(
        request.get("excluded_rows") or request.get("excludedRows"),
        "排除行",
    )
    included = collect(
        request.get("included_rows") or request.get("includedRows"),
        "保留行",
    )
    if mode == "all_except" and included:
        raise SplitExportError("全部保留模式下不能提交保留行清单。")
    if mode == "only_selected" and excluded:
        raise SplitExportError("仅保留选中模式下不能提交排除行清单。")
    if mode == "only_selected" and not included:
        raise SplitExportError("仅保留选中行时必须选择至少 1 行。")
    return {
        "mode": mode,
        "excluded": excluded,
        "included": included,
    }


def resolve_row_selection(
    all_records: Sequence[Record],
    filtered: Sequence[Record],
    request: Dict[str, Any],
) -> Tuple[List[Record], Dict[str, Any]]:
    selection = normalize_row_selection(request)
    scope_jobs = {record.job_id for record in all_records}
    if not scope_jobs:
        raise SplitExportError("数据范围为空，无法进行行选择。")
    if len(scope_jobs) == 1:
        only_job = next(iter(scope_jobs))
    else:
        only_job = ""

    known = {row_identity(record) for record in all_records}
    candidate_keys = {row_identity(record) for record in filtered}

    def resolve_refs(refs: Sequence[Tuple[str, str]], field: str) -> set:
        resolved = set()
        for job_id, key in refs:
            if not job_id:
                if not only_job:
                    raise SplitExportError(f"{field}必须包含来源任务标识。")
                job_id = only_job
            if job_id not in scope_jobs:
                raise SplitExportError(f"{field}不属于当前数据范围。")
            identity = (job_id, key)
            if identity not in known:
                raise SplitExportError(f"{field}包含无效行标识。")
            resolved.add(identity)
        return resolved

    excluded = resolve_refs(selection["excluded"], "排除行")
    included = resolve_refs(selection["included"], "保留行")
    if selection["mode"] == "only_selected":
        missing = included - candidate_keys
        if missing:
            raise SplitExportError("所选行不在当前筛选结果中，请重新选择。")

    final: List[Record] = []
    for record in filtered:
        identity = row_identity(record)
        keep = (
            identity not in excluded
            if selection["mode"] == "all_except"
            else identity in included
        )
        if keep:
            final.append(record)

    return final, {
        "mode": selection["mode"],
        "excluded_rows": [
            {"job_id": job_id, "internal_line_key": key}
            for job_id, key in selection["excluded"]
        ],
        "included_rows": [
            {"job_id": job_id, "internal_line_key": key}
            for job_id, key in selection["included"]
        ],
    }


def _dimension_identity(
    record: Record,
    dimension: str,
    group_options: Dict[str, Any],
) -> Tuple[str, str]:
    line = record.line
    if dimension == "account":
        code = _text(line.get("account_code"))
        name = _text(line.get("account_name"))
        prefix = group_options.get("account_prefix_length")
        if prefix not in (None, ""):
            try:
                prefix = int(prefix)
            except (TypeError, ValueError):
                raise SplitExportError("科目级次必须为整数。")
            if prefix < 1 or prefix > 20:
                raise SplitExportError("科目级次必须在 1–20 之间。")
            if code:
                return code[:prefix], code[:prefix]
        value = code or name
        display = f"{code} {name}".strip() if code and name else value
        return (value, display) if value else ("__blank__", BLANK_LABELS["account"])
    if dimension in ("customer", "supplier", "counterparty"):
        code = _text(line.get(f"{dimension}_code"))
        name = _text(line.get(f"{dimension}_name"))
        value = code or name
        display = name or code
        return (value, display) if value else ("__blank__", BLANK_LABELS[dimension])
    if dimension == "month":
        parsed = _parse_date(record.date)
        if not parsed:
            return "__blank__", BLANK_LABELS["month"]
        month = parsed.isoformat()[:7]
        return month, month
    if dimension == "period":
        value = _text(record.period)
        return (value, value) if value else ("__blank__", BLANK_LABELS["period"])
    if dimension == "entity":
        value = _text(record.entity)
        return (value, value) if value else ("__blank__", BLANK_LABELS["entity"])
    if dimension == "source_file":
        value = _text(record.filename)
        return (f"{record.job_id}:{value}", value) if value else ("__blank__", BLANK_LABELS["source_file"])
    raise SplitExportError("不支持的拆分维度。")


def _group_identity(
    record: Record,
    dimensions: Sequence[str],
    group_options: Dict[str, Any],
) -> Tuple[Tuple[Tuple[str, str], ...], Dict[str, str], Dict[str, str], str]:
    key_parts = []
    values = {}
    raw_values = {}
    labels = []
    for dimension in dimensions:
        raw, label = _dimension_identity(record, dimension, group_options or {})
        key_parts.append((dimension, raw))
        raw_values[dimension] = raw
        values[dimension] = label
        labels.append(label)
    return tuple(key_parts), values, raw_values, " / ".join(labels)


def group_records(
    records: Sequence[Record],
    dimensions: Sequence[str],
    group_options: Dict[str, Any],
) -> List[Dict[str, Any]]:
    groups: Dict[Tuple[Tuple[str, str], ...], Dict[str, Any]] = {}
    for record in records:
        key, values, raw_values, display = _group_identity(
            record, dimensions, group_options or {}
        )
        group = groups.setdefault(key, {
            "group_key": json.dumps(
                [{"dimension": dimension, "value": raw_values[dimension]} for dimension in dimensions],
                ensure_ascii=False, separators=(",", ":"),
            ),
            "display_name": display,
            "values": values,
            "raw_values": raw_values,
            "rows": [],
            "vouchers": set(),
            "debit": Decimal("0.00"),
            "credit": Decimal("0.00"),
        })
        group["rows"].append(record)
        group["vouchers"].add((record.job_id, record.voucher_key))
        group["debit"] += _decimal(record.line.get("debit"))
        group["credit"] += _decimal(record.line.get("credit"))
    result = []
    for group in groups.values():
        result.append({
            **group,
            "line_count": len(group["rows"]),
            "voucher_count": len(group["vouchers"]),
            "debit_total": group["debit"].quantize(Decimal("0.01")),
            "credit_total": group["credit"].quantize(Decimal("0.01")),
        })
    result.sort(key=lambda item: tuple(str(value) for value in item["raw_values"].values()))
    return result


def _normalize_dimensions(raw: Any) -> List[str]:
    if isinstance(raw, str):
        dimensions = [raw] if raw else []
    elif isinstance(raw, (list, tuple)):
        dimensions = [str(item) for item in raw if str(item)]
    elif raw in (None, ""):
        dimensions = []
    else:
        raise SplitExportError("拆分维度无效。")
    if not 1 <= len(dimensions) <= 3:
        raise SplitExportError("拆分维度数量必须为 1–3 个。")
    if len(set(dimensions)) != len(dimensions):
        raise SplitExportError("拆分维度不能重复。")
    for dimension in dimensions:
        if dimension not in GROUP_BY_VALUES:
            raise SplitExportError("不支持的拆分维度。")
    return dimensions


def _limits(request: Dict[str, Any]) -> Dict[str, int]:
    raw = request.get("limits") or {}
    mode = request.get("output_mode") or "multi_sheet"
    default_groups = (
        DEFAULT_MAX_GROUPS_ZIP if mode == "zip" else DEFAULT_MAX_GROUPS_MULTI_SHEET
    )
    return {
        "max_groups": max(1, int(raw.get("max_groups") or default_groups)),
        "max_rows_per_sheet": max(
            1, min(1_048_575, int(raw.get("max_rows_per_sheet") or DEFAULT_MAX_ROWS_PER_SHEET))
        ),
    }


def _preview_warnings(
    records: Sequence[Record],
    groups: Sequence[Dict[str, Any]],
    dimensions: Sequence[str],
    limits: Dict[str, int],
) -> List[str]:
    warnings = []
    for dimension in dimensions:
        if dimension in ("customer", "supplier", "counterparty"):
            field = f"{dimension}_name"
            code_field = f"{dimension}_code"
            if not any(
                _text(record.line.get(field)) or _text(record.line.get(code_field))
                for record in records
            ):
                warnings.append(f"{dimension}_not_mapped")
    if any(
        any(value in BLANK_LABELS.values() for value in group["values"].values())
        for group in groups
    ):
        warnings.append("blank_group_included")
    if len(groups) > limits["max_groups"]:
        warnings.append("group_limit_exceeded")
    return warnings


def preview_split(request: Dict[str, Any]) -> Dict[str, Any]:
    dimensions = _normalize_dimensions(request.get("group_by"))
    if request.get("row_mode") not in (None, "", "matched_only"):
        raise SplitExportError("FULL VOUCHER MODE NOT IMPLEMENTED")
    jobs = list(request.get("jobs") or [])
    records = load_records(jobs)
    filtered = apply_filters(records, request.get("filters") or {})
    final, selection = resolve_row_selection(records, filtered, request)
    groups = group_records(final, dimensions, request.get("group_options") or {})
    stats = _column_statistics(final)
    selected_columns = _normalize_columns(request.get("columns"), stats)
    limits = _limits(request)
    warnings = _preview_warnings(records, groups, dimensions, limits)
    group_count = len(groups)
    limit_exceeded = group_count > limits["max_groups"]
    vouchers = {(record.job_id, record.voucher_key) for record in final}
    return {
        "schema_version": SPLIT_SCHEMA_VERSION,
        "matched_lines": len(final),
        "matched_vouchers": len(vouchers),
        "candidate_lines": len(filtered),
        "candidate_vouchers": len({(record.job_id, record.voucher_key) for record in filtered}),
        "excluded_lines": len(filtered) - len(final),
        "row_selection_mode": selection["mode"],
        "excluded_rows": selection["excluded_rows"],
        "included_rows": selection["included_rows"],
        "completed_files": len(jobs),
        "source_file_count": len({record.filename for record in final}),
        "group_count": group_count,
        "group_by": dimensions,
        "dimensions": dimensions,
        "group_options": request.get("group_options") or {},
        "output_mode": request.get("output_mode") or "multi_sheet",
        "row_mode": request.get("row_mode") or "matched_only",
        "limits": limits,
        "limit_exceeded": limit_exceeded,
        "warnings": warnings,
        "column_stats": stats,
        "selected_columns": selected_columns,
        "column_order": selected_columns,
        "included_job_ids": [str(job.get("job_id")) for job in jobs],
        "groups": _group_payload(groups, request, limits),
        "data_sheet_count": (
            len(_multi_sheet_layout(groups, limits["max_rows_per_sheet"]))
            if (request.get("output_mode") or "multi_sheet") == "multi_sheet"
            else 0
        ),
        "aux_sheet_count": (
            2 if (request.get("output_mode") or "multi_sheet") == "multi_sheet" else 0
        ),
    }


def _display_value(record: Record, column_id: str) -> Any:
    value = _column_value(record, column_id)
    definition = COLUMN_BY_ID[column_id]
    if _is_empty(value):
        return None
    if definition["kind"] == "amount":
        try:
            return format(Decimal(str(value)).quantize(Decimal("0.01")), ".2f")
        except (InvalidOperation, ValueError):
            return _text(value) or None
    if definition["kind"] == "date":
        return _text(value)[:10] or None
    if definition["kind"] == "number":
        try:
            return int(value)
        except (TypeError, ValueError):
            return _text(value) or None
    return _text(value) or None


def _page_numbers(request: Dict[str, Any]) -> Tuple[int, int]:
    try:
        page = int(request.get("page") or 1)
    except (TypeError, ValueError):
        page = 1
    try:
        page_size = int(request.get("page_size") or request.get("pageSize") or DEFAULT_PAGE_SIZE)
    except (TypeError, ValueError):
        page_size = DEFAULT_PAGE_SIZE
    return max(1, page), max(1, min(MAX_PAGE_SIZE, page_size))


def _group_payload(
    groups: Sequence[Dict[str, Any]],
    request: Dict[str, Any],
    limits: Dict[str, int],
) -> List[Dict[str, Any]]:
    dimensions = _normalize_dimensions(request.get("group_by"))
    sheets_by_group: Dict[str, List[Dict[str, Any]]] = {}
    if (request.get("output_mode") or "multi_sheet") == "multi_sheet":
        for item in _multi_sheet_layout(groups, limits["max_rows_per_sheet"]):
            sheets_by_group.setdefault(item["group"]["group_key"], []).append(item)
    payload = []
    for group in groups:
        sheets = [
            {
                "sheet_name": item["sheet_name"],
                "part": item["part"],
                "part_count": item["part_count"],
                "line_count": item["line_count"],
                "voucher_count": item["voucher_count"],
                "debit_total": format(item["debit_total"], ".2f"),
                "credit_total": format(item["credit_total"], ".2f"),
            }
            for item in sheets_by_group.get(group["group_key"], [])
        ]
        entry = {
            "group_key": group["group_key"],
            "group_ref": {
                "dimensions": [
                    {
                        "dimension": dimension,
                        "value": group["raw_values"][dimension],
                    }
                    for dimension in dimensions
                ]
            },
            "display_name": group["display_name"],
            "values": group["values"],
            "raw_values": group["raw_values"],
            "line_count": group["line_count"],
            "voucher_count": group["voucher_count"],
            "debit_total": format(group["debit_total"], ".2f"),
            "credit_total": format(group["credit_total"], ".2f"),
            "sheets": sheets,
        }
        if sheets:
            entry["sheet_name"] = sheets[0]["sheet_name"]
        if "account" in dimensions:
            account_code, account_name = _account_index_values(group, request)
            entry["account_code"] = account_code
            entry["account_name"] = account_name
        payload.append(entry)
    return payload


def _normalize_summary_overrides(raw: Any) -> Dict[str, Dict[str, Any]]:
    if not isinstance(raw, dict):
        return {}
    result = {}
    for account_code, value in raw.items():
        if not isinstance(value, dict):
            continue
        code = str(account_code or "").strip()
        if not code:
            continue
        fields = {}
        for key, field_value in (value.get("fields") or {}).items():
            if not isinstance(field_value, dict):
                continue
            field_key = str(key or "").strip()
            if not field_key:
                continue
            fields[field_key] = field_value
        result[code] = {
            "disabled": bool(value.get("disabled")),
            "confirmed": bool(value.get("confirmed")),
            "fields": fields,
        }
    return result


def _effective_summary_rule(
    rule: SummaryStructureRule | None,
    overrides: Dict[str, Dict[str, Any]],
) -> SummaryStructureRule | None:
    if rule is None or rule.status not in ("auto", "review"):
        return None
    override = overrides.get(rule.account_code) or {}
    if override.get("disabled"):
        return None
    if rule.status == "review" and not override.get("confirmed"):
        return None
    override_fields = override.get("fields") or {}
    fields = []
    for field in rule.fields:
        field_override = override_fields.get(field.key) or {}
        if field_override.get("selected") is False:
            continue
        semantic_type = field.semantic_type
        if field_override.get("semantic_type"):
            try:
                semantic_type = SemanticType(str(field_override["semantic_type"]))
            except ValueError:
                semantic_type = field.semantic_type
        label = str(field_override.get("label") or field.label).strip()[:40] or field.label
        fields.append(replace(field, label=label, semantic_type=semantic_type))
    if not fields:
        return None
    return replace(rule, fields=fields, status="auto")


def _group_summary_rule(
    group: Dict[str, Any],
    dimensions: Sequence[str],
    rules: Dict[Tuple[str, str], SummaryStructureRule],
    overrides: Dict[str, Dict[str, Any]],
) -> SummaryStructureRule | None:
    if "account" not in dimensions:
        return None
    account_code = _text((group.get("raw_values") or {}).get("account"))
    if not account_code or account_code == "__blank__":
        return None
    job_ids = {record.job_id for record in group.get("rows") or []}
    candidates = [
        rules.get((job_id, account_code))
        for job_id in sorted(job_ids)
    ]
    candidates = [rule for rule in candidates if rule is not None]
    if not candidates:
        return None
    signatures = {rule.signature() for rule in candidates}
    if len(signatures) != 1:
        return None
    return _effective_summary_rule(candidates[0], overrides)



def _group_source_rule(
    group: Dict[str, Any],
    dimensions: Sequence[str],
    rules: Dict[Tuple[str, str], SummaryStructureRule],
) -> SummaryStructureRule | None:
    if "account" not in dimensions:
        return None
    account_code = _text((group.get("raw_values") or {}).get("account"))
    if not account_code or account_code == "__blank__":
        return None
    job_ids = {record.job_id for record in group.get("rows") or []}
    candidates = [
        rules.get((job_id, account_code))
        for job_id in sorted(job_ids)
    ]
    candidates = [rule for rule in candidates if rule is not None]
    if not candidates:
        return None
    signatures = {rule.signature() for rule in candidates}
    if len(signatures) != 1:
        return None
    return candidates[0]


def _summary_structure_counts(
    records: Sequence[Record],
    rules: Dict[Tuple[str, str], SummaryStructureRule],
) -> Dict[str, Any]:
    values = list(rules.values())
    covered = 0
    for record in records:
        code = _text(record.line.get("account_code"))
        rule = rules.get((record.job_id, code))
        if rule is not None and rule.status == "auto":
            covered += 1
    return {
        "version": 1,
        "account_count": len({(key[0], key[1]) for key in rules}),
        "auto_rule_count": sum(1 for rule in values if rule.status == "auto"),
        "review_rule_count": sum(1 for rule in values if rule.status == "review"),
        "none_rule_count": sum(1 for rule in values if rule.status == "none"),
        "covered_line_count": covered,
        "total_line_count": len(records),
    }

def _derived_fields(rule: SummaryStructureRule | None) -> List[Dict[str, Any]]:
    if rule is None:
        return []
    return [
        {
            "key": field.key,
            "label": field.label,
            "semantic_type": field.semantic_type.value,
            "part_index": field.part_index,
            "confidence": field.confidence,
        }
        for field in rule.fields
    ]


def _derived_field_map(rule: SummaryStructureRule | None) -> Dict[str, Any]:
    if rule is None:
        return {}
    return {field.key: field for field in rule.fields}


def _derived_value_for_json(value: Any) -> Any:
    if isinstance(value, date):
        return value.isoformat()
    return value


def _sheet_columns_with_derived(
    columns: Sequence[str], rule: SummaryStructureRule | None
) -> List[str]:
    base = list(columns)
    derived = [field.key for field in (rule.fields if rule else [])]
    if not derived:
        return base
    if "description" in base:
        index = base.index("description") + 1
        return base[:index] + derived + base[index:]
    return base + derived


def _standard_headers_with_derived(
    base_columns: Sequence[str], rule: SummaryStructureRule | None
) -> List[str]:
    headers = [COLUMN_BY_ID[column_id]["label"] for column_id in base_columns]
    if rule is None:
        return headers
    derived = [field.label for field in rule.fields]
    if "description" in base_columns:
        index = list(base_columns).index("description") + 1
        return headers[:index] + derived + headers[index:]
    return headers + derived


def _standard_row_with_derived(
    record: Record,
    base_columns: Sequence[str],
    rule: SummaryStructureRule | None,
) -> List[Any]:
    base_values = _standard_row(record, base_columns)
    if rule is None:
        return base_values
    extracted = extract_derived_fields(
        _text(record.line.get("description")),
        rule,
    )
    derived_values = []
    for field in rule.fields:
        value = extracted.get(field.key)
        if field.semantic_type != SemanticType.DATE and isinstance(value, str):
            value = safe_excel_text(value)
        derived_values.append(value)
    if "description" in base_columns:
        index = list(base_columns).index("description") + 1
        return base_values[:index] + derived_values + base_values[index:]
    return base_values + derived_values


def _style_sheet_with_derived(
    sheet, base_columns: Sequence[str], rule: SummaryStructureRule | None
) -> None:
    _style_data_sheet(sheet, base_columns)
    if rule is None:
        return
    offset = (
        list(base_columns).index("description") + 1
        if "description" in base_columns else len(base_columns)
    )
    for index, field in enumerate(rule.fields, start=1):
        column_index = offset + index
        letter = get_column_letter(column_index)
        width = 24 if field.semantic_type != SemanticType.DATE else 14
        sheet.column_dimensions[letter].width = width
        if field.semantic_type == SemanticType.DATE:
            for cell in sheet[letter][1:]:
                cell.number_format = "yyyy-mm-dd"


def _normalize_preview_group(
    raw: Any, dimensions: Sequence[str]
) -> Optional[Tuple[Tuple[str, str], ...]]:
    if raw in (None, ""):
        return None
    if isinstance(raw, dict):
        items = raw.get("dimensions")
    else:
        items = raw
    if not isinstance(items, (list, tuple)) or len(items) != len(dimensions):
        raise SplitExportError("预览 Sheet 标识无效。")
    values = []
    for expected, item in zip(dimensions, items):
        if not isinstance(item, dict):
            raise SplitExportError("预览 Sheet 标识无效。")
        dimension = str(item.get("dimension") or "")
        value = str(item.get("value") or "")
        if dimension != expected:
            raise SplitExportError("预览 Sheet 维度与当前拆分方式不一致。")
        values.append((dimension, value))
    return tuple(values)


def _find_preview_group(
    groups: Sequence[Dict[str, Any]],
    dimensions: Sequence[str],
    selected: Tuple[Tuple[str, str], ...],
) -> Dict[str, Any]:
    target = json.dumps(
        [{"dimension": dimension, "value": value} for dimension, value in selected],
        ensure_ascii=False,
        separators=(",", ":"),
    )
    for group in groups:
        if group["group_key"] == target:
            return group
    raise SplitExportError("预览 Sheet 不存在或已被筛选排除。")


def preview_rows(request: Dict[str, Any]) -> Dict[str, Any]:
    """Server-side paged export preview with manual row selection."""

    dimensions = _normalize_dimensions(request.get("group_by"))
    if request.get("row_mode") not in (None, "", "matched_only"):
        raise SplitExportError("FULL VOUCHER MODE NOT IMPLEMENTED")
    jobs = list(request.get("jobs") or [])
    records, summary_rules = load_records(jobs, collect_rules=True)
    summary_overrides = _normalize_summary_overrides(
        request.get("summary_overrides") or request.get("summaryOverrides")
    )
    filtered = apply_filters(records, request.get("filters") or {})
    final, selection = resolve_row_selection(records, filtered, request)
    groups = group_records(final, dimensions, request.get("group_options") or {})
    stats = _column_statistics(final)
    selected_columns = _normalize_columns(request.get("columns"), stats)
    limits = _limits(request)
    warnings = _preview_warnings(records, groups, dimensions, limits)

    page, page_size = _page_numbers(request)
    group_payload = _group_payload(groups, request, limits)
    account_dimension = "account" in dimensions
    for entry, group in zip(group_payload, groups):
        source_rule = _group_source_rule(group, dimensions, summary_rules)
        effective = _effective_summary_rule(source_rule, summary_overrides)
        entry["summary_rule"] = {
            "status": source_rule.status if source_rule else "none",
            "applied": effective is not None,
            "confidence": source_rule.overall_confidence if source_rule else 0.0,
            "delimiter": source_rule.delimiter if source_rule else None,
            "fields": _derived_fields(source_rule),
            "applied_fields": _derived_fields(effective),
        }
    is_multi_sheet = (request.get("output_mode") or "multi_sheet") == "multi_sheet"
    layout = (
        _multi_sheet_layout(groups, limits["max_rows_per_sheet"])
        if is_multi_sheet else []
    )
    selected_group = _normalize_preview_group(
        request.get("preview_group") or request.get("previewGroup"),
        dimensions,
    )
    active_sheet = None
    if selected_group is None:
        scope_records = filtered
        scope_line_count = len(filtered)
        scope_voucher_count = len({(record.job_id, record.voucher_key) for record in filtered})
    else:
        group = _find_preview_group(groups, dimensions, selected_group)
        group_sheets = [
            item for item in layout if item["group"] is group
        ]
        try:
            preview_part = int(request.get("preview_part") or request.get("previewPart") or 1)
        except (TypeError, ValueError):
            raise SplitExportError("预览 Sheet 分片无效。")
        if group_sheets and not 1 <= preview_part <= len(group_sheets):
            raise SplitExportError("预览 Sheet 分片不存在。")
        if group_sheets:
            selected_sheet = group_sheets[preview_part - 1]
            if len(group_sheets) == 1:
                scope_records = [
                    record for record in filtered
                    if _group_identity(
                        record,
                        dimensions,
                        request.get("group_options") or {},
                    )[0] == selected_group
                ]
                included_scope = selected_sheet["line_count"]
            else:
                scope_records = selected_sheet["rows"]
                included_scope = selected_sheet["line_count"]
            active_sheet = {
                "group_key": group["group_key"],
                "group_ref": {
                    "dimensions": [
                        {"dimension": dimension, "value": group["raw_values"][dimension]}
                        for dimension in dimensions
                    ]
                },
                "sheet_name": selected_sheet["sheet_name"],
                "part": selected_sheet["part"],
                "part_count": selected_sheet["part_count"],
                "line_count": selected_sheet["line_count"],
                "voucher_count": selected_sheet["voucher_count"],
                "debit_total": format(selected_sheet["debit_total"], ".2f"),
                "credit_total": format(selected_sheet["credit_total"], ".2f"),
            }
        else:
            scope_records = group["rows"]
            active_sheet = {
                "group_key": group["group_key"],
                "group_ref": {
                    "dimensions": [
                        {"dimension": dimension, "value": group["raw_values"][dimension]}
                        for dimension in dimensions
                    ]
                },
                "sheet_name": group["display_name"],
                "part": 1,
                "part_count": 1,
                "line_count": group["line_count"],
                "voucher_count": group["voucher_count"],
                "debit_total": format(group["debit_total"], ".2f"),
                "credit_total": format(group["credit_total"], ".2f"),
            }
        scope_line_count = len(scope_records)
        scope_included_line_count = included_scope
        scope_voucher_count = len({
            (record.job_id, record.voucher_key) for record in group["rows"]
        })

    if selected_group is None:
        scope_included_line_count = len(final)
    active_rule = None
    if selected_group is not None and account_dimension:
        active_group = _find_preview_group(groups, dimensions, selected_group)
        active_rule = _group_summary_rule(
            active_group, dimensions, summary_rules, summary_overrides
        )
    derived_columns = _derived_fields(active_rule)
    total = scope_line_count
    total_pages = (total + page_size - 1) // page_size if total else 0
    start = (page - 1) * page_size
    page_records = scope_records[start:start + page_size]

    selection_keys = (
        {(item["job_id"], item["internal_line_key"]) for item in selection["excluded_rows"]}
        if selection["mode"] == "all_except"
        else {(item["job_id"], item["internal_line_key"]) for item in selection["included_rows"]}
    )
    rows = []
    for record in page_records:
        identity = row_identity(record)
        included = (
            identity not in selection_keys
            if selection["mode"] == "all_except"
            else identity in selection_keys
        )
        values = {
            column_id: _display_value(record, column_id)
            for column_id in COLUMN_BY_ID
        }
        if active_rule is not None:
            extracted = extract_derived_fields(
                _text(record.line.get("description")), active_rule
            )
            for field in active_rule.fields:
                values[field.key] = _derived_value_for_json(extracted.get(field.key))
        rows.append({
            "job_id": record.job_id,
            "internal_line_key": identity[1],
            "voucher_key": record.voucher_key,
            "included": included,
            "values": values,
        })

    return {
        "schema_version": SPLIT_SCHEMA_VERSION,
        "candidate_line_count": len(filtered),
        "included_line_count": len(final),
        "excluded_line_count": len(filtered) - len(final),
        "candidate_voucher_count": len({(record.job_id, record.voucher_key) for record in filtered}),
        "included_voucher_count": len({(record.job_id, record.voucher_key) for record in final}),
        "page": page,
        "page_size": page_size,
        "total_pages": total_pages,
        "columns": selected_columns,
        "column_order": selected_columns,
        "column_stats": stats,
        "rows": rows,
        "row_selection_mode": selection["mode"],
        "excluded_rows": selection["excluded_rows"],
        "included_rows": selection["included_rows"],
        "group_summary": {
            "group_count": len(groups),
            "groups": group_payload,
        },
        "group_count": len(groups),
        "groups": group_payload,
        "derived_columns": derived_columns,
        "summary_rule": (
            {
                **next(
                    (
                        entry["summary_rule"]
                        for entry in group_payload
                        if selected_group is not None
                        and entry["group_key"] == _find_preview_group(groups, dimensions, selected_group)["group_key"]
                    ),
                    {},
                ),
                "applied_fields": derived_columns,
            }
            if selected_group is not None and account_dimension else None
        ),
        "summary_structure": _summary_structure_counts(records, summary_rules),
        "data_sheet_count": len(layout) if is_multi_sheet else 0,
        "aux_sheet_count": 2 if is_multi_sheet else 0,
        "scope_line_count": scope_line_count,
        "scope_included_line_count": scope_included_line_count,
        "scope_excluded_line_count": max(0, scope_line_count - scope_included_line_count),
        "scope_voucher_count": scope_voucher_count,
        "active_sheet": active_sheet,
        "limit_exceeded": len(groups) > limits["max_groups"],
        "limits": limits,
        "warnings": warnings,
        "dimensions": dimensions,
        "group_by": dimensions,
        "included_job_ids": [str(job.get("job_id")) for job in jobs],
    }


def dimension_options(request: Dict[str, Any]) -> Dict[str, Any]:
    dimension = str(request.get("dimension") or "")
    if dimension not in GROUP_BY_VALUES and dimension not in ("account",):
        raise SplitExportError("不支持的维度。")
    jobs = list(request.get("jobs") or [])
    records = load_records(jobs)
    query = _text(request.get("q")).casefold()
    limit = max(1, min(200, int(request.get("limit") or 50)))
    counts: Dict[str, Dict[str, Any]] = {}
    blank = 0
    for record in records:
        line = record.line
        if dimension == "account":
            code = _text(line.get("account_code"))
            name = _text(line.get("account_name"))
            value = code or name
            label = f"{code} {name}".strip() or value
            extra = {"code": code, "name": name}
        elif dimension in ("customer", "supplier", "counterparty"):
            code = _text(line.get(f"{dimension}_code"))
            name = _text(line.get(f"{dimension}_name"))
            value = code or name
            label = name or code
            extra = {"code": code, "name": name}
        elif dimension == "month":
            parsed = _parse_date(record.date)
            value = parsed.isoformat()[:7] if parsed else ""
            label = value
            extra = {}
        elif dimension == "period":
            value = _text(record.period)
            label = value
            extra = {}
        elif dimension == "entity":
            value = _text(record.entity)
            label = value
            extra = {}
        else:
            value = _text(record.filename)
            label = value
            extra = {}
        if not value:
            blank += 1
            continue
        entry = counts.setdefault(value, {"value": value, "label": label, "count": 0, **extra})
        entry["count"] += 1
        if label and not entry.get("label"):
            entry["label"] = label
    options = list(counts.values())
    if query:
        options = [
            option for option in options
            if query in str(option.get("value", "")).casefold()
            or query in str(option.get("label", "")).casefold()
            or query in str(option.get("code", "")).casefold()
            or query in str(option.get("name", "")).casefold()
        ]
    options.sort(key=lambda item: (-item["count"], str(item["label"] or item["value"])))
    total = len(options)
    return {
        "dimension": dimension,
        "options": options[:limit],
        "total": total,
        "truncated": total > limit,
        "blank_count": blank,
    }


def _safe_sheet_name(display: Any, used: set) -> str:
    raw = ILLEGAL_SHEET.sub("_", _text(display)).strip("'") or "未命名"
    base = raw[:31]
    candidate = base
    index = 2
    while candidate.casefold() in used:
        suffix = f"_{index}"
        candidate = f"{base[:31 - len(suffix)]}{suffix}"
        index += 1
    used.add(candidate.casefold())
    return candidate


def _multi_sheet_layout(
    groups: Sequence[Dict[str, Any]], max_rows: int
) -> List[Dict[str, Any]]:
    used = {"拆分说明", "分组索引"}
    layout: List[Dict[str, Any]] = []
    for group in groups:
        source_rows = group["rows"]
        parts = [
            source_rows[index:index + max_rows]
            for index in range(0, len(source_rows), max_rows)
        ] or [[]]
        part_count = len(parts)
        for part_index, part in enumerate(parts, start=1):
            name = (
                group["display_name"]
                if part_count == 1
                else f"{group['display_name']}_{part_index}"
            )
            sheet_name = _safe_sheet_name(name, used)
            layout.append({
                "group": group,
                "rows": part,
                "part": part_index,
                "part_count": part_count,
                "sheet_name": sheet_name,
                "line_count": len(part),
                "voucher_count": len({
                    (record.job_id, record.voucher_key) for record in part
                }),
                "debit_total": sum(
                    (_decimal(record.line.get("debit")) for record in part),
                    Decimal("0.00"),
                ),
                "credit_total": sum(
                    (_decimal(record.line.get("credit")) for record in part),
                    Decimal("0.00"),
                ),
            })
    return layout


def _safe_file_name(display: Any, fallback: str = "group") -> str:
    name = _text(display).replace("\\", "/").split("/")[-1]
    name = ILLEGAL_FILENAME.sub("_", name).strip(" .")
    return (name or fallback)[:80]


def _column_value(record: Record, column_id: str) -> Any:
    line = record.line
    if column_id == "date":
        return record.date
    if column_id == "voucher_id":
        return record.voucher_id
    if column_id == "entity":
        return record.entity
    if column_id == "period":
        return record.period
    if column_id == "original_filename":
        return record.filename
    if column_id == "job_id":
        return record.job_id
    if column_id == "voucher_key":
        return record.voucher_key
    if column_id == "debit":
        return _decimal(line.get("debit"))
    if column_id == "credit":
        return _decimal(line.get("credit"))
    if column_id == "original_amount":
        return _decimal(line.get("original_amount")) if line.get("original_amount") not in (None, "") else None
    return line.get(column_id)


def _is_empty(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, str):
        return not value.strip()
    return False


def _column_statistics(records: Sequence[Record]) -> Dict[str, Dict[str, Any]]:
    total = len(records)
    stats: Dict[str, Dict[str, Any]] = {}
    for definition in COLUMN_DEFINITIONS:
        column_id = definition["id"]
        non_empty = sum(0 if _is_empty(_column_value(record, column_id)) else 1 for record in records)
        stats[column_id] = {
            "id": column_id,
            "label": definition["label"],
            "group": definition["group"],
            "advanced": bool(definition.get("advanced")),
            "non_empty_count": non_empty,
            "total_count": total,
            "entirely_empty": non_empty == 0,
        }
    return stats


def _normalize_columns(raw: Any, stats: Dict[str, Dict[str, Any]]) -> List[str]:
    explicit = raw is not None
    if not explicit:
        return [
            column_id for column_id in DEFAULT_COLUMN_ORDER
            if (
                COLUMN_BY_ID[column_id].get("default")
                or (
                    COLUMN_BY_ID[column_id].get("default_if_data")
                    and not stats[column_id]["entirely_empty"]
                )
            )
        ]
    if not isinstance(raw, (list, tuple)) or not raw:
        raise SplitExportError("至少选择 1 个导出字段。")
    selected: List[str] = []
    for value in raw:
        column_id = str(value or "").strip()
        if column_id not in COLUMN_BY_ID:
            raise SplitExportError("导出字段无效。")
        if column_id in selected:
            raise SplitExportError("导出字段不能重复。")
        selected.append(column_id)
    if not selected:
        raise SplitExportError("至少选择 1 个导出字段。")
    return selected


def _style_data_sheet(sheet, columns: Sequence[str]) -> None:
    for index, column_id in enumerate(columns, start=1):
        definition = COLUMN_BY_ID[column_id]
        letter = get_column_letter(index)
        if definition["kind"] == "amount":
            for cell in sheet[letter][1:]:
                cell.number_format = "#,##0.00"
        if definition["kind"] == "date":
            for cell in sheet[letter][1:]:
                cell.number_format = "yyyy-mm-dd"
        width = {
            "date": 12, "description": 52, "account_name": 28,
            "customer_name": 24, "supplier_name": 24,
            "counterparty_name": 24, "account_code": 18,
            "customer_code": 18, "supplier_code": 18, "counterparty_code": 18,
            "original_filename": 28, "sheet_name": 18, "source_row": 10,
            "job_id": 24, "line_id": 18, "internal_line_key": 28,
            "voucher_key": 32, "voucher_id": 20, "detail_account": 24,
            "entity": 18, "period": 14, "currency": 10, "preparer": 14,
        }.get(column_id, 20)
        sheet.column_dimensions[letter].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def _standard_headers(columns: Sequence[str]) -> List[str]:
    return [COLUMN_BY_ID[column_id]["label"] for column_id in columns]


def _standard_row(record: Record, columns: Sequence[str]) -> List[Any]:
    values = []
    for column_id in columns:
        value = _column_value(record, column_id)
        if COLUMN_BY_ID[column_id]["kind"] in ("text", "key"):
            value = _safe_excel(value or "")
        values.append(value)
    return values


def _description_rows(request: Dict[str, Any], preview: Dict[str, Any], group: Optional[Dict[str, Any]] = None) -> List[List[Any]]:
    scope = request.get("scope") or {}
    rows = [
        ["生成时间", request.get("generated_at") or ""],
        ["数据范围", f"{scope.get('type', '')}:{scope.get('id', '')}"],
        ["筛选条件", json.dumps(request.get("filters") or {}, ensure_ascii=False, sort_keys=True)[:4000]],
        ["拆分维度", json.dumps(preview.get("group_by"), ensure_ascii=False)],
        ["输出模式", request.get("output_mode")],
        ["输出范围", request.get("row_mode") or "matched_only"],
        ["候选行数", preview.get("candidate_lines", preview.get("matched_lines"))],
        ["导出保留行数", preview.get("matched_lines")],
        ["手工排除行数", preview.get("excluded_lines", 0)],
        ["匹配凭证数", preview.get("matched_vouchers")],
        ["行选择模式", preview.get("row_selection_mode", "all_except")],
        ["分组数量", preview.get("group_count")],
        ["来源文件数", preview.get("source_file_count")],
        ["导出字段", json.dumps(preview.get("selected_columns"), ensure_ascii=False)],
    ]
    if group:
        rows.extend([
            ["分组名称", group["display_name"]],
            ["分组行数", group["line_count"]],
            ["分组凭证数", group["voucher_count"]],
        ])
    return rows


def _write_description(workbook, request, preview, group=None):
    sheet = workbook.create_sheet("拆分说明")
    sheet.append(["项目", "内容"])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
    for label, value in _description_rows(request, preview, group):
        sheet.append([_safe_excel(label), _safe_excel(value)])
    sheet.column_dimensions["A"].width = 18
    sheet.column_dimensions["B"].width = 90


def _account_index_values(
    group: Dict[str, Any], request: Dict[str, Any]
) -> Tuple[str, str]:
    raw = _text((group.get("raw_values") or {}).get("account"))
    if raw == "__blank__":
        return "未指定", "未指定科目"
    prefix = (request.get("group_options") or {}).get("account_prefix_length")
    if prefix not in (None, ""):
        return raw, f"按前 {prefix} 位归并"
    names: List[str] = []
    for record in group.get("rows") or []:
        name = _text(record.line.get("account_name"))
        if name and name not in names:
            names.append(name)
    if not names:
        name = ""
    elif len(names) == 1:
        name = names[0]
    else:
        name = "、".join(names[:2]) + (" 等" if len(names) > 2 else "")
    return raw, name


def _build_multi_sheet(path: Path, request: Dict[str, Any], preview: Dict[str, Any], groups: Sequence[Dict[str, Any]], summary_rules: Dict[Tuple[str, str], SummaryStructureRule] | None = None, summary_overrides: Dict[str, Dict[str, Any]] | None = None) -> None:
    workbook = Workbook()
    default = workbook.active
    workbook.remove(default)
    _write_description(workbook, request, preview)
    index = workbook.create_sheet("分组索引")
    dimensions = preview.get("dimensions") or []
    dimension_headers = []
    for dimension in dimensions:
        if dimension == "account":
            dimension_headers.extend(["科目编码", "科目名称"])
        else:
            dimension_headers.append(GROUP_BY_LABELS.get(dimension, dimension))
    index.append(dimension_headers + ["Sheet 名称", "部分", "行数", "凭证数", "借方合计", "贷方合计"])
    for cell in index[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
    max_rows = _limits(request)["max_rows_per_sheet"]
    columns = preview["selected_columns"]
    summary_rules = summary_rules or {}
    summary_overrides = summary_overrides or {}
    for item in _multi_sheet_layout(groups, max_rows):
        group = item["group"]
        rule = _group_summary_rule(
            group, preview.get("dimensions") or [], summary_rules, summary_overrides
        )
        sheet_columns = _sheet_columns_with_derived(columns, rule)
        sheet = workbook.create_sheet(item["sheet_name"])
        _write_sheet(
            sheet,
            _standard_headers_with_derived(columns, rule),
            [_standard_row_with_derived(record, columns, rule) for record in item["rows"]],
        )
        _style_sheet_with_derived(sheet, columns, rule)
        dimension_values = []
        for dimension in dimensions:
            if dimension == "account":
                account_code, account_name = _account_index_values(group, request)
                dimension_values.extend([
                    _safe_excel(account_code),
                    _safe_excel(account_name),
                ])
            else:
                dimension_values.append(
                    _safe_excel(group["values"].get(dimension, ""))
                )
        index.append(
            dimension_values + [
                item["sheet_name"], item["part"], item["line_count"],
                item["voucher_count"],
                item["debit_total"].quantize(Decimal("0.01")),
                item["credit_total"].quantize(Decimal("0.01")),
            ]
        )
    index.freeze_panes = "A2"
    index.column_dimensions["A"].width = 34
    index.column_dimensions["B"].width = 34
    workbook.save(path)
    try:
        path.chmod(0o600)
    except OSError:
        pass


def _build_group_workbook(path: Path, request: Dict[str, Any], preview: Dict[str, Any], group: Dict[str, Any], rows: Sequence[Record], part: int, part_count: int, summary_rules: Dict[Tuple[str, str], SummaryStructureRule] | None = None, summary_overrides: Dict[str, Dict[str, Any]] | None = None) -> None:
    workbook = Workbook()
    default = workbook.active
    workbook.remove(default)
    _write_description(workbook, request, preview, group)
    sheet = workbook.create_sheet("序时账")
    columns = preview["selected_columns"]
    summary_rules = summary_rules or {}
    summary_overrides = summary_overrides or {}
    rule = _group_summary_rule(
        group, preview.get("dimensions") or [], summary_rules, summary_overrides
    )
    _write_sheet(
        sheet,
        _standard_headers_with_derived(columns, rule),
        [_standard_row_with_derived(record, columns, rule) for record in rows],
    )
    _style_sheet_with_derived(sheet, columns, rule)
    workbook.save(path)
    try:
        path.chmod(0o600)
    except OSError:
        pass


def build_split_output(output_dir: Path, request: Dict[str, Any]) -> Dict[str, Any]:
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    output_mode = request.get("output_mode") or "multi_sheet"
    if output_mode not in ("multi_sheet", "zip"):
        raise SplitExportError("不支持的输出模式。")
    preview = preview_split(request)
    if preview["matched_lines"] == 0:
        raise SplitExportError("筛选结果为 0 行，无法生成。")
    if preview["limit_exceeded"]:
        raise SplitExportError(
            f"预计生成 {preview['group_count']} 个分组，超过当前输出上限 {preview['limits']['max_groups']}。"
        )
    jobs = list(request.get("jobs") or [])
    records, summary_rules = load_records(jobs, collect_rules=True)
    summary_overrides = _normalize_summary_overrides(
        request.get("summary_overrides") or request.get("summaryOverrides")
    )
    filtered = apply_filters(records, request.get("filters") or {})
    final, selection = resolve_row_selection(records, filtered, request)
    dimensions = _normalize_dimensions(request.get("group_by"))
    groups = group_records(final, dimensions, request.get("group_options") or {})
    max_rows = preview["limits"]["max_rows_per_sheet"]
    manifest_groups = []
    if output_mode == "multi_sheet":
        output_path = output_dir / "split-export.xlsx"
        _build_multi_sheet(output_path, request, preview, groups, summary_rules, summary_overrides)
        for group in groups:
            manifest_groups.append({
                "group_key": group["group_key"],
                "display_name": group["display_name"],
                "values": group["values"],
                "raw_values": group["raw_values"],
                "line_count": group["line_count"],
                "voucher_count": group["voucher_count"],
                "output_filename": output_path.name,
                "derived_fields": _derived_fields(_group_summary_rule(
                    group, dimensions, summary_rules, summary_overrides
                )),
            })
        final_vouchers = len({(record.job_id, record.voucher_key) for record in final})
        manifest = {
            "schema_version": SPLIT_SCHEMA_VERSION,
            "selected_columns": preview["selected_columns"],
            "summary_structure_version": 1,
            "summary_rule_overrides": summary_overrides,
            "scope": request.get("scope"),
            "generated_at": request.get("generated_at"),
            "filters": request.get("filters") or {},
            "group_by": dimensions,
            "dimensions": dimensions,
            "group_options": request.get("group_options") or {},
            "output_mode": output_mode,
            "row_mode": request.get("row_mode") or "matched_only",
            "row_selection_mode": selection["mode"],
            "excluded_rows": selection["excluded_rows"],
            "included_rows": selection["included_rows"],
            "candidate_line_count": len(filtered),
            "final_line_count": len(final),
            "candidate_voucher_count": len({(record.job_id, record.voucher_key) for record in filtered}),
            "matched_lines": len(final),
            "matched_vouchers": final_vouchers,
            "group_count": len(groups),
            "groups": manifest_groups,
        }
        return {
            "output_mode": output_mode,
            "output_file": output_path.name,
            "group_count": len(groups),
            "matched_lines": len(final),
            "matched_vouchers": final_vouchers,
            "selected_columns": preview["selected_columns"],
            "manifest": manifest,
        }

    zip_path = output_dir / "split-export.zip"
    used = set()
    sequence = 1
    for group in groups:
        source_rows = group["rows"]
        parts = [source_rows[i:i + max_rows] for i in range(0, len(source_rows), max_rows)] or [[]]
        for part_index, part in enumerate(parts, start=1):
            stem = _safe_file_name(group["display_name"])
            filename = f"{sequence:03d}-{stem}"
            if len(parts) > 1:
                filename += f"-part{part_index:03d}"
            filename += ".xlsx"
            while filename.casefold() in used:
                filename = f"{sequence:03d}-{stem}-{sequence}.xlsx"
            used.add(filename.casefold())
            sequence += 1
            _build_group_workbook(
                output_dir / filename, request, preview, group, part,
                part_index, len(parts), summary_rules, summary_overrides
            )
            part_vouchers = len({(record.job_id, record.voucher_key) for record in part})
            manifest_groups.append({
                "group_key": group["group_key"],
                "display_name": group["display_name"],
                "values": group["values"],
                "raw_values": group["raw_values"],
                "line_count": len(part),
                "voucher_count": part_vouchers,
                "output_filename": filename,
                "part": part_index,
                "part_count": len(parts),
                "derived_fields": _derived_fields(_group_summary_rule(
                    group, dimensions, summary_rules, summary_overrides
                )),
            })
    final_vouchers = len({(record.job_id, record.voucher_key) for record in final})
    manifest = {
        "schema_version": SPLIT_SCHEMA_VERSION,
        "scope": request.get("scope"),
        "generated_at": request.get("generated_at"),
        "filters": request.get("filters") or {},
        "group_by": dimensions,
        "dimensions": dimensions,
        "group_options": request.get("group_options") or {},
        "output_mode": output_mode,
        "row_mode": request.get("row_mode") or "matched_only",
        "row_selection_mode": selection["mode"],
        "excluded_rows": selection["excluded_rows"],
        "included_rows": selection["included_rows"],
        "candidate_line_count": len(filtered),
        "final_line_count": len(final),
        "candidate_voucher_count": len({(record.job_id, record.voucher_key) for record in filtered}),
        "matched_lines": len(final),
        "matched_vouchers": final_vouchers,
        "group_count": len(groups),
        "selected_columns": preview["selected_columns"],
        "summary_structure_version": 1,
        "summary_rule_overrides": summary_overrides,
        "groups": manifest_groups,
    }
    with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED, allowZip64=True) as archive:
        archive.writestr("manifest.json", json.dumps(manifest, ensure_ascii=False, indent=2) + "\n")
        for item in manifest_groups:
            archive.write(output_dir / item["output_filename"], arcname=item["output_filename"])
    try:
        zip_path.chmod(0o600)
    except OSError:
        pass
    return {
        "output_mode": output_mode,
        "output_file": zip_path.name,
        "group_count": len(groups),
        "matched_lines": len(final),
        "matched_vouchers": final_vouchers,
        "selected_columns": preview["selected_columns"],
        "manifest": manifest,
    }


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description="Liuxian Lab split export engine")
    parser.add_argument("--mode", choices=("preview", "preview_rows", "dimensions", "generate"), required=True)
    parser.add_argument("--request", required=True)
    parser.add_argument("--output", required=False)
    parser.add_argument("--output-dir", required=False)
    args = parser.parse_args(argv)
    try:
        request = json.loads(Path(args.request).read_text(encoding="utf-8"))
        if args.mode == "preview":
            payload = preview_split(request)
        elif args.mode == "preview_rows":
            payload = preview_rows(request)
        elif args.mode == "dimensions":
            payload = dimension_options(request)
        else:
            payload = build_split_output(Path(args.output_dir), request)
        if args.output:
            Path(args.output).write_text(
                json.dumps(payload, ensure_ascii=False), encoding="utf-8"
            )
            try:
                Path(args.output).chmod(0o600)
            except OSError:
                pass
        else:
            print(json.dumps(payload, ensure_ascii=False))
        return 0
    except SplitExportError as exc:
        print(str(exc), file=sys.stderr)
        return 1
    except Exception as exc:  # noqa: BLE001
        print(f"拆分导出失败：{type(exc).__name__}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
