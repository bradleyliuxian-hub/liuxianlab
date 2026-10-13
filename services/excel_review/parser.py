"""Unified extraction of Excel comments/notes from a workbook.

Produces one row per comment node:

- Legacy Excel Note/Comment (openpyxl ``cell.comment``)
- Microsoft 365 threaded comment root
- Microsoft 365 threaded reply

Cell values shown to the user ("Original Value") are read from a
``data_only=True`` openpyxl workbook, so formula cells expose Excel's
cached result instead of the formula string. Formula error results and
uncached formulas are normalised to an empty string.
"""

import hashlib
import uuid
import zipfile
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from openpyxl.utils import coordinate_to_tuple

from . import ooxml_threaded_comments as threaded
from .ooxml_package import OoxmlPackage
from .statistics import build_statistics, normalize_author
from .status import ReviewError

MAPPING_VERSION = "excel-review-mapping-v3"

COMMENT_TYPE_LEGACY = "legacy"
COMMENT_TYPE_THREADED = "threaded"
COMMENT_TYPE_THREADED_REPLY = "threaded_reply"

# Kept for backward compatibility with existing callers/tests.
ITEM_TYPE_COMMENT = "comment"

SOURCE_REFERENCE_SEPARATOR = " → "

# Excel error values that must never appear as a user-facing Original Value.
EXCEL_ERROR_VALUES = frozenset(
    {
        "#NULL!",
        "#DIV/0!",
        "#VALUE!",
        "#REF!",
        "#NAME?",
        "#NUM!",
        "#N/A",
        "#GETTING_DATA",
        "#SPILL!",
        "#CALC!",
        "#CONNECT!",
        "#BLOCKED!",
        "#UNKNOWN!",
        "#FIELD!",
        "#BUSY!",
    }
)


def sha256_file(file_path) -> str:
    digest = hashlib.sha256()
    with open(file_path, "rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def sha256_text(text: str) -> str:
    return hashlib.sha256((text or "").encode("utf-8")).hexdigest()


def normalize_display_value(value):
    """Return a business-readable, JSON-safe display value.

    - empty / Excel error / uncached formula -> ""
    - bool -> "TRUE" / "FALSE"
    - int / float -> number (kept numeric)
    - datetime / date -> readable string
    - str -> itself
    """
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, str):
        return "" if value in EXCEL_ERROR_VALUES else value
    if isinstance(value, (int, float)):
        return value
    if isinstance(value, Decimal):
        try:
            if value == value.to_integral_value():
                return int(value)
        except Exception:  # noqa: BLE001
            pass
        return str(value)
    if isinstance(value, datetime):
        if (
            value.hour == 0
            and value.minute == 0
            and value.second == 0
            and value.microsecond == 0
        ):
            return value.strftime("%Y-%m-%d")
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    return str(value)


# Backward-compatible alias.
def stringify_value(value) -> str:
    return str(normalize_display_value(value))


def detect_threaded_comments(file_path) -> bool:
    """True if the workbook advertises Microsoft 365 threaded comments."""
    return threaded.has_threaded_comments(file_path)


def threaded_comment_cells(file_path):
    """{(sheet_name, cell_address)} that carry a threaded comment."""
    cells = set()
    for item in threaded.extract_threaded_comments(file_path):
        cells.add((item["sheetName"], item["cellAddress"]))
    return cells


def load_source_workbook(file_path, data_only=False):
    """Open a source workbook for reading (never saved back by us)."""
    try:
        return load_workbook(
            file_path,
            data_only=data_only,
            read_only=False,
            keep_links=True,
        )
    except Exception as exc:  # noqa: BLE001 - normalise library errors
        raise ReviewError(
            "无法读取该 Excel 工作簿，文件可能已损坏或受密码保护。"
        ) from exc


def source_reference(source_filename, sheet_name, cell_address):
    return SOURCE_REFERENCE_SEPARATOR.join(
        [source_filename or "source.xlsx", sheet_name or "", cell_address or ""]
    )


def format_excel_address(sheet_name, cell_address):
    """Build an Excel Go To (F5 / Ctrl+G) reference: ``'Sheet'!B5``.

    The sheet name is always single-quoted and Excel-escaped (``'`` ->
    ``''``) so the address is valid for any sheet name and can be pasted
    straight into the Excel "定位" window.
    """
    if not sheet_name or not cell_address:
        return ""
    escaped = str(sheet_name).replace("'", "''")
    return "'%s'!%s" % (escaped, cell_address)


def _normalized_author(raw):
    return normalize_author(raw)


def _display_cell_value(value_workbook, sheet_name, cell_address):
    """Read the cached/actual value and normalise it for display."""
    try:
        worksheet = value_workbook[sheet_name]
    except KeyError:
        return ""
    try:
        value = worksheet[cell_address].value
    except (ValueError, TypeError):
        return ""
    return normalize_display_value(value)


def _extract_legacy_items(formula_workbook, value_workbook, threaded_cells):
    items = []
    for sheet_index, worksheet in enumerate(formula_workbook.worksheets):
        for row in worksheet.iter_rows():
            for cell in row:
                comment = cell.comment
                if comment is None:
                    continue
                # Skip the legacy fallback note Excel writes for a threaded
                # comment; the threaded node already represents this cell.
                if (worksheet.title, cell.coordinate) in threaded_cells:
                    continue
                comment_text = comment.text if comment.text is not None else ""
                items.append(
                    {
                        "kind": COMMENT_TYPE_LEGACY,
                        "sheetName": worksheet.title,
                        "sheetIndex": sheet_index,
                        "cellAddress": cell.coordinate,
                        "cellValue": _display_cell_value(
                            value_workbook, worksheet.title, cell.coordinate
                        ),
                        "commentText": str(comment_text),
                        "author": _normalized_author(comment.author),
                        "commentDatetime": "",
                        "threadId": "",
                        "parentThreadId": "",
                        "personId": "",
                        "resolved": False,
                        "doneRaw": "",
                        "orderKey": coordinate_to_tuple(cell.coordinate),
                        "orderInPart": 0,
                    }
                )
    return items


def _extract_threaded_items(formula_workbook, value_workbook, raw_threaded):
    items = []
    sheet_index_by_name = {
        ws.title: index for index, ws in enumerate(formula_workbook.worksheets)
    }
    for raw in raw_threaded:
        sheet_name = raw["sheetName"]
        cell_address = raw["cellAddress"]
        sheet_index = sheet_index_by_name.get(
            sheet_name, raw.get("sheetIndex", 0)
        )
        items.append(
            {
                "kind": raw["commentType"],
                "sheetName": sheet_name,
                "sheetIndex": sheet_index,
                "cellAddress": cell_address,
                "cellValue": _display_cell_value(
                    value_workbook, sheet_name, cell_address
                ),
                "commentText": raw["commentText"],
                "author": raw["author"],
                "commentDatetime": raw["commentDatetime"],
                "threadId": raw["threadId"],
                "parentThreadId": raw["parentThreadId"],
                "personId": raw["personId"],
                "resolved": bool(raw.get("resolved", False)),
                "doneRaw": raw.get("doneRaw", ""),
                "orderKey": coordinate_to_tuple(cell_address)
                if cell_address
                else (0, 0),
                "orderInPart": raw.get("orderInPart", 0),
            }
        )
    return items


def normalize_thread_status(entries):
    """Propagate the root conversation's resolved state to every node.

    Status is a *thread-level* property: a reply inherits the resolved
    state of its root threaded comment (Excel stores ``done`` on the root).
    Orphans / broken parent chains are treated conservatively as their own
    root (never mis-assigned to another thread) and never crash.
    """
    nodes = {
        entry["threadId"]: entry
        for entry in entries
        if entry.get("threadId")
        and entry.get("kind")
        in (COMMENT_TYPE_THREADED, COMMENT_TYPE_THREADED_REPLY)
    }

    def find_root(entry):
        seen = set()
        current = entry
        while current is not None:
            thread_id = current.get("threadId")
            if thread_id in seen:
                return None  # cycle guard
            seen.add(thread_id)
            parent_id = current.get("parentThreadId")
            if not parent_id:
                return current
            parent = nodes.get(parent_id)
            if parent is None:
                # Missing parent: treat the node as its own root.
                return current
            current = parent
        return None

    for entry in entries:
        if entry.get("kind") not in (
            COMMENT_TYPE_THREADED,
            COMMENT_TYPE_THREADED_REPLY,
        ):
            entry["resolved"] = False
            entry["rootThreadId"] = ""
            continue
        root = find_root(entry)
        if root is None:
            entry["resolved"] = False
            entry["rootThreadId"] = entry.get("threadId", "")
        else:
            entry["resolved"] = bool(root.get("resolved", False))
            entry["rootThreadId"] = root.get("threadId", "")


def extract(workbook_path, source_filename="source.xlsx") -> dict:
    """Read a source workbook and build the extracted.json payload."""
    path = Path(workbook_path)
    if not path.exists():
        raise ReviewError("源文件不存在。")

    has_threaded = False
    raw_threaded = []
    threaded_cells = set()
    try:
        has_threaded = threaded.has_threaded_comments(path)
        if has_threaded:
            raw_threaded = threaded.extract_threaded_comments(path)
            threaded_cells = {
                (item["sheetName"], item["cellAddress"]) for item in raw_threaded
            }
    except (OSError, zipfile.BadZipFile, ET.ParseError) as exc:
        raise ReviewError(
            "无法读取该 Excel 工作簿，文件可能已损坏或受密码保护。"
        ) from exc

    warnings = []
    formula_workbook = load_source_workbook(path, data_only=False)
    value_workbook = load_source_workbook(path, data_only=True)
    try:
        combined = _extract_legacy_items(
            formula_workbook, value_workbook, threaded_cells
        )
        combined.extend(
            _extract_threaded_items(formula_workbook, value_workbook, raw_threaded)
        )
    finally:
        formula_workbook.close()
        value_workbook.close()

    combined.sort(
        key=lambda item: (
            item["sheetIndex"],
            item["orderKey"],
            item["orderInPart"],
        )
    )

    normalize_thread_status(combined)

    items = []
    for index, entry in enumerate(combined, start=1):
        text = entry["commentText"]
        text_hash = sha256_text(text)
        comment_type = entry["kind"]
        resolved = bool(entry.get("resolved", False))
        items.append(
            {
                "itemId": uuid.uuid4().hex,
                "reviewId": "C%03d" % index,
                "commentType": comment_type,
                "sheetName": entry["sheetName"],
                "cellAddress": entry["cellAddress"],
                "cellValue": entry["cellValue"],
                "commentText": text,
                "author": entry["author"],
                "commentDatetime": entry["commentDatetime"],
                "threadId": entry["threadId"],
                "parentThreadId": entry["parentThreadId"],
                "rootThreadId": entry.get("rootThreadId", ""),
                "personId": entry["personId"],
                "resolved": resolved,
                "status": "completed" if resolved else "pending",
                "originalTextHash": text_hash,
                "legacyCommentHash": text_hash
                if comment_type == COMMENT_TYPE_LEGACY
                else "",
                "sourceFilename": source_filename,
                "sourceReference": source_reference(
                    source_filename, entry["sheetName"], entry["cellAddress"]
                ),
                "excelAddress": format_excel_address(
                    entry["sheetName"], entry["cellAddress"]
                ),
            }
        )

    if has_threaded:
        warnings.append(
            "已读取 Microsoft 365 新版批注；原生线程回复写回暂未支持。"
        )

    statistics = build_statistics(items)

    return {
        "mappingVersion": MAPPING_VERSION,
        "sourceFilename": source_filename,
        "sourceSha256": sha256_file(path),
        "hasThreaded": has_threaded,
        "itemCount": len(items),
        "statistics": statistics,
        "warnings": warnings,
        "items": items,
    }
