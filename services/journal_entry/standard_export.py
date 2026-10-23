"""Single / batch-merge standardized journal export workspace.

Consumes completed canonical result.json files. It never re-opens source
workbooks and never re-runs Reader / Mapping / Normalize / Business / Contra /
Anomaly.
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any, Dict, List, Sequence

from openpyxl import Workbook
from openpyxl.styles import Font

from .export import HEADER_FILL, _write_sheet, safe_excel_text
from .split_export import (
    DEFAULT_PAGE_SIZE,
    MAX_EXPLICIT_ROWS,
    MAX_PAGE_SIZE,
    SPLIT_SCHEMA_VERSION,
    _column_statistics,
    _display_value,
    _normalize_columns,
    _page_numbers,
    _standard_headers,
    _standard_row,
    _style_data_sheet,
    apply_filters,
    load_records,
    resolve_row_selection,
    row_identity,
)

STANDARD_SCHEMA_VERSION = 1


def _warnings(records, final) -> List[str]:
    warnings = []
    if not final:
        warnings.append("no_selected_rows")
    if len(records) and not final:
        warnings.append("all_rows_excluded")
    return warnings


def preview_standard(request: Dict[str, Any]) -> Dict[str, Any]:
    jobs = list(request.get("jobs") or [])
    records = load_records(jobs)
    filtered = apply_filters(records, request.get("filters") or {})
    final, selection = resolve_row_selection(records, filtered, request)
    stats = _column_statistics(final)
    selected_columns = _normalize_columns(request.get("columns"), stats)
    page, page_size = _page_numbers(request)
    total = len(filtered)
    total_pages = (total + page_size - 1) // page_size if total else 0
    start = (page - 1) * page_size
    page_records = filtered[start:start + page_size]

    selection_keys = (
        {(item["job_id"], item["internal_line_key"]) for item in selection["excluded_rows"]}
        if selection["mode"] == "all_except"
        else {(item["job_id"], item["internal_line_key"]) for item in selection["included_rows"]}
    )
    rows = []
    for record in page_records:
        identity = row_identity(record)
        included = (
            identity not in selection_keys
            if selection["mode"] == "all_except"
            else identity in selection_keys
        )
        rows.append({
            "job_id": record.job_id,
            "internal_line_key": identity[1],
            "voucher_key": record.voucher_key,
            "included": included,
            "values": {
                column_id: _display_value(record, column_id)
                for column_id in stats
            },
        })

    return {
        "schema_version": STANDARD_SCHEMA_VERSION,
        "scope": request.get("scope"),
        "candidate_line_count": total,
        "included_line_count": len(final),
        "excluded_line_count": total - len(final),
        "candidate_voucher_count": len({(record.job_id, record.voucher_key) for record in filtered}),
        "included_voucher_count": len({(record.job_id, record.voucher_key) for record in final}),
        "source_file_count": len({record.job_id for record in final}),
        "page": page,
        "page_size": page_size,
        "total_pages": total_pages,
        "columns": selected_columns,
        "column_order": selected_columns,
        "column_stats": stats,
        "rows": rows,
        "row_selection_mode": selection["mode"],
        "excluded_rows": selection["excluded_rows"],
        "included_rows": selection["included_rows"],
        "warnings": _warnings(records, final),
        "limits": {"max_explicit_rows": MAX_EXPLICIT_ROWS},
        "included_job_ids": [str(job.get("job_id")) for job in jobs],
    }


def _write_explanation(workbook, request, final, columns) -> None:
    sheet = workbook.create_sheet("导出说明")
    sheet.append(["项目", "内容"])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
    scope = request.get("scope") or {}
    rows = [
        ("生成时间", request.get("generated_at") or ""),
        ("数据范围", f"{scope.get('type', '')}:{scope.get('id', '')}"),
        ("筛选条件", json.dumps(request.get("filters") or {}, ensure_ascii=False, sort_keys=True)[:4000]),
        ("行选择模式", request.get("row_selection_mode") or "all_except"),
        ("导出保留行数", len(final)),
        ("导出凭证数", len({(record.job_id, record.voucher_key) for record in final})),
        ("导出字段", json.dumps(columns, ensure_ascii=False)),
    ]
    for label, value in rows:
        sheet.append([safe_excel_text(label), safe_excel_text(value)])
    sheet.column_dimensions["A"].width = 20
    sheet.column_dimensions["B"].width = 96


def build_standard_output(output_dir: Path, request: Dict[str, Any]) -> Dict[str, Any]:
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    preview = preview_standard({**request, "page": 1, "page_size": MAX_PAGE_SIZE})
    if not preview["included_line_count"]:
        raise ValueError("最终保留行为 0，无法导出。")
    jobs = list(request.get("jobs") or [])
    records = load_records(jobs)
    filtered = apply_filters(records, request.get("filters") or {})
    final, selection = resolve_row_selection(records, filtered, request)
    columns = preview["columns"]

    workbook = Workbook()
    default = workbook.active
    workbook.remove(default)
    _write_explanation(workbook, request, final, columns)
    sheet = workbook.create_sheet("标准分录")
    _write_sheet(
        sheet,
        _standard_headers(columns),
        [_standard_row(record, columns) for record in final],
    )
    _style_data_sheet(sheet, columns)
    output_path = output_dir / "standard-journal.xlsx"
    workbook.save(output_path)
    workbook.close()
    try:
        output_path.chmod(0o600)
    except OSError:
        pass

    return {
        "output_file": output_path.name,
        "included_line_count": len(final),
        "included_voucher_count": len({(record.job_id, record.voucher_key) for record in final}),
        "candidate_line_count": len(filtered),
        "selected_columns": columns,
        "row_selection_mode": selection["mode"],
        "excluded_rows": selection["excluded_rows"],
        "included_rows": selection["included_rows"],
        "included_job_ids": [str(job.get("job_id")) for job in jobs],
    }


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description="Liuxian Lab standardized journal export")
    parser.add_argument("--mode", choices=("preview", "export"), required=True)
    parser.add_argument("--request", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--output-dir", required=False)
    args = parser.parse_args(argv)
    try:
        request = json.loads(Path(args.request).read_text(encoding="utf-8"))
        if args.mode == "preview":
            payload = preview_standard(request)
        else:
            if not args.output_dir:
                raise ValueError("missing output directory")
            payload = build_standard_output(Path(args.output_dir), request)
        Path(args.output).write_text(
            json.dumps(payload, ensure_ascii=False), encoding="utf-8"
        )
        return 0
    except Exception as exc:  # noqa: BLE001
        print(f"标准序时账导出失败：{type(exc).__name__}: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
