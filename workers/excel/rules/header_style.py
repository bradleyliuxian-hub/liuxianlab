from copy import copy
from openpyxl.styles import PatternFill


def _style_signature(cell):
    return {
        "font": {
            "bold": bool(cell.font.bold),
            "color": cell.font.color.rgb if cell.font.color and cell.font.color.type == "rgb" else None,
        },
        "fill": cell.fill.fgColor.rgb if cell.fill and cell.fill.fgColor and cell.fill.fgColor.type == "rgb" else None,
        "alignment": {
            "horizontal": cell.alignment.horizontal,
            "vertical": cell.alignment.vertical,
        },
    }


def apply(workbook, audit, config):
    if not config.get("enabled", True):
        return
    fill_color = str(config.get("fill", "EAF8CC")).replace("#", "").upper()
    font_color = str(config.get("font_color", "171915")).replace("#", "").upper()
    bold = bool(config.get("bold", True))
    horizontal = config.get("horizontal", "left")
    vertical = config.get("vertical", "center")
    row_height = config.get("row_height", 24)

    for ws in workbook.worksheets:
        if ws.max_row < 1:
            continue
        header_cells = [cell for cell in ws[1] if cell.value not in (None, "")]
        if not header_cells:
            continue
        for cell in header_cells:
            before = _style_signature(cell)

            font = copy(cell.font)
            font.bold = bold
            font.color = font_color
            cell.font = font

            cell.fill = PatternFill(fill_type="solid", fgColor=fill_color)

            alignment = copy(cell.alignment)
            alignment.horizontal = horizontal
            alignment.vertical = vertical
            cell.alignment = alignment

            after = _style_signature(cell)
            if before != after:
                audit.add(
                    rule="header_style",
                    sheet=ws.title,
                    target=cell.coordinate,
                    change_type="style",
                    before=before,
                    after=after,
                    reason="统一第一行表头样式",
                )
        if row_height:
            before_height = ws.row_dimensions[1].height
            if before_height != row_height:
                ws.row_dimensions[1].height = row_height
                audit.add(
                    rule="header_style",
                    sheet=ws.title,
                    target="1:1",
                    change_type="dimension",
                    before={"height": before_height},
                    after={"height": row_height},
                    reason="统一表头行高",
                )
