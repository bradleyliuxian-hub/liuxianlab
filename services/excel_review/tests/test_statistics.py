"""Unified statistics + question-definition + web/Excel invariant tests (V1.3)."""

import json
import shutil
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.comments import Comment

from services.excel_review import review_workbook
from services.excel_review.statistics import build_statistics, normalize_author
from services.excel_review.status import Reporter
from services.excel_review.worker import run
from services.excel_review.tests.threaded_fixture import (
    build_base_workbook,
    inject_threaded_comments,
    new_guid,
)


def q(ctype, author, text="x"):
    return {"commentType": ctype, "author": author, "commentText": text}


class TestQuestionDefinition(unittest.TestCase):
    def test_threaded_root_and_legacy_are_questions(self):
        stats = build_statistics([
            q("threaded", "Alice"),
            q("legacy", "Alice"),
            q("threaded_reply", "Bob"),
            q("threaded_reply", "Carol"),
        ])
        self.assertEqual(stats["questions_total"], 2)
        self.assertEqual(stats["threaded_questions"], 1)
        self.assertEqual(stats["legacy_questions"], 1)
        self.assertEqual(stats["threaded_replies"], 2)

    def test_scenario_4_root_7_replies_5_legacy(self):
        items = [q("threaded", "A") for _ in range(4)]
        items += [q("threaded_reply", "R") for _ in range(7)]
        items += [q("legacy", "A") for _ in range(5)]
        stats = build_statistics(items)
        self.assertEqual(stats["questions_total"], 9)
        self.assertEqual(stats["threaded_questions"], 4)
        self.assertEqual(stats["legacy_questions"], 5)

    def test_replies_only_is_zero_questions(self):
        stats = build_statistics([q("threaded_reply", "A") for _ in range(3)])
        self.assertEqual(stats["questions_total"], 0)
        self.assertEqual(stats["threaded_replies"], 3)
        self.assertEqual(stats["authors"], [])


class TestAuthorStatistics(unittest.TestCase):
    def test_scenario_alice_bob_carol(self):
        items = []
        items += [q("threaded", "Alice") for _ in range(3)]
        items += [q("legacy", "Alice") for _ in range(2)]
        items += [q("threaded_reply", "Alice") for _ in range(5)]
        items += [q("threaded", "Bob") for _ in range(1)]
        items += [q("legacy", "Bob") for _ in range(3)]
        items += [q("threaded_reply", "Bob") for _ in range(4)]
        items += [q("threaded_reply", "Carol") for _ in range(7)]
        stats = build_statistics(items)
        by_author = {row["author"]: row["questions"] for row in stats["authors"]}
        self.assertEqual(by_author.get("Alice"), 5)
        self.assertEqual(by_author.get("Bob"), 4)
        self.assertNotIn("Carol", by_author)
        self.assertEqual(stats["authors_total"], 2)

    def test_blank_author_normalised(self):
        for value in (None, "", "   ", "None"):
            self.assertEqual(normalize_author(value), "未知作者")
        stats = build_statistics([q("legacy", None), q("threaded", "   "), q("legacy", "")])
        by_author = {row["author"]: row["questions"] for row in stats["authors"]}
        self.assertEqual(by_author, {"未知作者": 3})

    def test_unicode_and_case_preserved(self):
        stats = build_statistics([q("legacy", "张三"), q("legacy", "Alice"), q("legacy", "alice")])
        names = {row["author"] for row in stats["authors"]}
        self.assertEqual(names, {"张三", "Alice", "alice"})

    def test_sorted_by_count_desc(self):
        items = [q("legacy", "Alice") for _ in range(3)] + [q("legacy", "Bob")]
        stats = build_statistics(items)
        self.assertEqual([r["author"] for r in stats["authors"]], ["Alice", "Bob"])


class TempBase(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def run_worker(self, source_path, name="job"):
        job_dir = self.tmp / name
        job_dir.mkdir(parents=True, exist_ok=True)
        shutil.copy2(str(source_path), str(job_dir / "source.xlsx"))
        run(job_dir, Reporter(job_dir))
        return job_dir


class TestWebExcelInvariant(TempBase):
    def _mixed_source(self, path):
        workbook = build_base_workbook([("S", {"A1": 1, "A2": 2, "B1": 10, "B2": 20, "B3": 30})])
        workbook["S"]["A1"].comment = Comment("legacy A", "Alice")
        workbook["S"]["A2"].comment = Comment("legacy B", "Bob")
        workbook.save(path)
        alice = new_guid()
        bob = new_guid()
        root1 = new_guid()
        root2 = new_guid()
        root3 = new_guid()
        inject_threaded_comments(
            str(path),
            "S",
            [
                {"ref": "B1", "id": root1, "personId": alice, "dT": "2026-10-13T06:00:00", "text": "t-root-1"},
                {"ref": "B1", "id": new_guid(), "personId": bob, "parentId": root1, "dT": "2026-10-13T06:01:00", "text": "reply"},
                {"ref": "B2", "id": root2, "personId": bob, "dT": "2026-10-13T06:02:00", "text": "t-root-2"},
                {"ref": "B3", "id": root3, "personId": alice, "dT": "2026-10-13T06:03:00", "text": "t-root-3"},
            ],
            [
                {"id": alice, "displayName": "Alice", "email": "alice@example.com"},
                {"id": bob, "displayName": "Bob", "email": "bob@example.com"},
            ],
        )

    def test_api_matches_workbook_dashboard_numbers(self):
        path = self.tmp / "mix.xlsx"
        self._mixed_source(path)
        job_dir = self.run_worker(path)
        data = json.loads((job_dir / "extracted.json").read_text(encoding="utf-8"))
        stats = data["statistics"]
        # 3 threaded roots + 2 legacy = 5 questions; 1 threaded reply.
        self.assertEqual(stats["questions_total"], 5)
        self.assertEqual(stats["threaded_questions"], 3)
        self.assertEqual(stats["legacy_questions"], 2)
        self.assertEqual(stats["threaded_replies"], 1)

        workbook = load_workbook(job_dir / "review.xlsx")
        threaded = workbook[review_workbook.THREADED_SHEET]
        legacy = workbook[review_workbook.LEGACY_SHEET]
        t_header = [c.value for c in threaded[1]]
        t_type = t_header.index("类型") + 1
        root_rows = 0
        for r in range(2, threaded.max_row + 1):
            if threaded.cell(row=r, column=t_type).value == "新版批注":
                root_rows += 1
        legacy_rows = sum(
            1 for r in range(2, legacy.max_row + 1) if legacy.cell(row=r, column=1).value is not None
        )
        self.assertEqual(root_rows, stats["threaded_questions"])
        self.assertEqual(legacy_rows, stats["legacy_questions"])
        self.assertEqual(root_rows + legacy_rows, stats["questions_total"])

    def test_tables_present_when_data(self):
        path = self.tmp / "mix.xlsx"
        self._mixed_source(path)
        job_dir = self.run_worker(path)
        workbook = load_workbook(job_dir / "review.xlsx")
        self.assertIn(review_workbook.TABLE_THREADED, workbook[review_workbook.THREADED_SHEET].tables)
        self.assertIn(review_workbook.TABLE_LEGACY, workbook[review_workbook.LEGACY_SHEET].tables)


if __name__ == "__main__":
    unittest.main()
