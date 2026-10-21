"""Server-side SVG / Excel / PDF exports."""

from __future__ import annotations

import io
from datetime import datetime
from pathlib import Path
from typing import Iterable, List, Optional, Sequence
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from reportlab.lib.pagesizes import A3, landscape
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

from .layout import Layout, PositionedNode, TITLE_BASELINE, build_routes, edge_points, layout_structure, percentage_label
from .schemas import StructureResult

NODE_FILL = "#FFF8EE"
CORE_FILL = "#FFF2DB"
NODE_STROKE = "#D49A43"
CORE_STROKE = "#C67D18"
TEXT_COLOR = "#17191C"
LINE_COLOR = "#777D83"


def _hex_rgb(value: str, fallback: str = "#777D83") -> tuple[float, float, float]:
    text = str(value or "").strip().lstrip("#")
    if len(text) != 6:
        text = str(fallback or "#777D83").lstrip("#")
    try:
        return tuple(int(text[index:index + 2], 16) / 255 for index in (0, 2, 4))  # type: ignore[return-value]
    except ValueError:
        return (0.47, 0.49, 0.51)


# --------------------------------------------------------------------------- SVG


def build_svg(result: StructureResult, layout_overrides: dict | None = None) -> str:
    layout = layout_structure(result, layout_overrides)
    positions = layout.nodes
    styles = (layout_overrides or {}).get("styles") if isinstance((layout_overrides or {}).get("styles"), dict) else {}
    node_styles = styles.get("nodes") if isinstance(styles.get("nodes"), dict) else {}
    edge_styles = styles.get("edges") if isinstance(styles.get("edges"), dict) else {}
    label_styles = styles.get("labels") if isinstance(styles.get("labels"), dict) else {}
    canvas_style = styles.get("canvas") if isinstance(styles.get("canvas"), dict) else {}
    parts = [
        f'<svg xmlns="http://www.w3.org/2000/svg" width="{layout.width:.0f}" height="{layout.height:.0f}" '
        f'viewBox="0 0 {layout.width:.0f} {layout.height:.0f}">',
        f'<rect width="100%" height="100%" fill="{escape(str(canvas_style.get("background") or "#ffffff"))}"/>',
        f'<text x="{layout.width / 2:.1f}" y="{TITLE_BASELINE}" text-anchor="middle" font-size="30" font-weight="700" '
        f'fill="{TEXT_COLOR}" font-family="PingFang SC, Noto Sans CJK SC, Noto Sans SC, Microsoft YaHei, Arial Unicode MS, WenQuanYi Zen Hei, sans-serif">'
        f'{escape(result.title)}</text>',
    ]
    routes = build_routes(result, layout, layout_overrides)
    for line in routes.lines:
        role = f' data-role="{line.role}"' if line.role != "edge" else ""
        style = edge_styles.get(line.edge_id) or {}
        stroke = str(style.get("stroke") or LINE_COLOR)
        width = float(style.get("strokeWidth", 1.4) or 1.4)
        dash = str(style.get("dash") or "none")
        dash_attr = f' stroke-dasharray="{escape(dash)}"' if dash != "none" else ""
        parts.append(
            f'<path{role} d="M{line.x1:.1f} {line.y1:.1f} L{line.x2:.1f} {line.y2:.1f}" '
            f'fill="none" stroke="{escape(stroke)}" stroke-width="{width}"{dash_attr}/>'
        )
    for x, y in routes.arrows:
        parts.append(f'<polygon points="{x - 4:.1f},{y - 8:.1f} {x + 4:.1f},{y - 8:.1f} {x:.1f},{y:.1f}" fill="{LINE_COLOR}"/>')
    for label in routes.labels:
        style = label_styles.get(label.edge_id) or {}
        bg = escape(str(style.get("background") or "#ffffff"))
        color = escape(str(style.get("color") or "#333333"))
        font_size = float(style.get("fontSize", 11) or 11)
        pad_x = float(style.get("paddingX", 5) or 0)
        pad_y = float(style.get("paddingY", 3) or 0)
        radius = float(style.get("radius", 3) or 0)
        parts.append(f'<rect x="{label.x:.1f}" y="{label.y:.1f}" width="{label.width:.1f}" height="{label.height:.1f}" rx="{radius}" fill="{bg}"/>')
        parts.append(
            f'<text x="{label.x + pad_x:.1f}" y="{label.y + label.height - pad_y - 1:.1f}" font-size="{font_size}" fill="{color}" '
            f'font-family="PingFang SC, Noto Sans CJK SC, Noto Sans SC, Microsoft YaHei, Arial Unicode MS, WenQuanYi Zen Hei, sans-serif">{escape(label.text)}</text>'
        )
    for node in result.nodes:
        if node.status == "excluded_cancelled":
            continue
        position = positions.get(node.id)
        if not position:
            continue
        _svg_node(parts, position, node.name == result.core_company, node_styles.get(node.id) or {})
    parts.append("</svg>")
    return "\n".join(parts)


def _svg_node(parts: List[str], position: PositionedNode, core: bool, style: dict | None = None) -> None:
    style = style or {}
    fill = style.get("fill") or (CORE_FILL if core else NODE_FILL)
    stroke = style.get("stroke") or (CORE_STROKE if core else NODE_STROKE)
    text_color = style.get("textColor") or TEXT_COLOR
    stroke_width = float(style.get("strokeWidth", 2 if core else 1.35) or (2 if core else 1.35))
    radius = float(style.get("radius", 10) or 0)
    font_size = float(style.get("fontSize", 21 if core else 15) or (21 if core else 15))
    font_weight = int(style.get("fontWeight", 700 if core else 500) or (700 if core else 500))
    parts.append(f'<g data-name="{escape(position.node.name)}"><title>{escape(position.node.name)}</title>')
    parts.append(
        f'<rect x="{position.x:.1f}" y="{position.y:.1f}" width="{position.width:.1f}" '
        f'height="{position.height:.1f}" rx="{radius}" fill="{escape(str(fill))}" stroke="{escape(str(stroke))}" '
        f'stroke-width="{stroke_width}"/>'
    )
    line_height = max(18.0, font_size * 1.42)
    total = len(position.lines) * line_height
    start_y = position.y + position.height / 2 - total / 2 + font_size
    for index, line in enumerate(position.lines):
        parts.append(
            f'<text x="{position.x + position.width / 2:.1f}" y="{start_y + index * line_height:.1f}" '
            f'text-anchor="middle" font-size="{font_size}" font-weight="{font_weight}" '
            f'fill="{escape(str(text_color))}" font-family="PingFang SC, Noto Sans CJK SC, Noto Sans SC, Microsoft YaHei, Arial Unicode MS, WenQuanYi Zen Hei, sans-serif">'
            f'{escape(line)}</text>'
        )
    parts.append('</g>')


# ------------------------------------------------------------------------- Excel


def build_excel(result: StructureResult) -> bytes:
    workbook = Workbook()
    header_fill = PatternFill("solid", fgColor="F4E5CB")
    header_font = Font(bold=True)
    thin = Side(style="thin", color="D9E2EC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    node_sheet = workbook.active
    node_sheet.title = "股权结构"
    node_sheet.append(["层级", "实体名称", "实体类型", "状态", "OCR置信度"])
    for node in sorted(result.nodes, key=lambda item: (item.level, item.name)):
        node_sheet.append([node.level, node.name, node.entity_type, node.status, node.confidence])

    idmap = {node.id: node.name for node in result.nodes}
    edge_sheet = workbook.create_sheet("投资关系")
    edge_sheet.append(["出资方/母公司", "被投资方/子公司", "持股比例", "关系", "置信度", "比例识别来源"])
    for edge in result.edges:
        edge_sheet.append([
            idmap.get(edge.source, edge.source),
            idmap.get(edge.target, edge.target),
            edge.percentage,
            edge.relation,
            edge.confidence,
            edge.percentage_source,
        ])

    notes = workbook.create_sheet("识别备注")
    notes.append(["项目", "内容"])
    notes.append(["标题", result.title])
    notes.append(["核心主体", result.core_company])
    notes.append(["生成时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
    notes.append(["警告", "；".join(result.warnings)])
    notes.append(["低置信度项", "；".join(result.quality.low_confidence_items)])
    notes.append(["字典纠错", "；".join(
        f"{item.original} → {item.corrected}（{item.similarity:.1f}）"
        for item in result.dictionary_corrections
    )])
    notes.append(["已排除企业", "；".join(f"{item.name}（{item.reason}）" for item in result.excluded)])
    notes.append(["识别引擎", result.diagnostics.engine])
    notes.append(["图片尺寸", f"{result.diagnostics.image_width}×{result.diagnostics.image_height}"])
    notes.append(["切片数", result.diagnostics.tile_count])
    notes.append(["节点框数", result.diagnostics.detected_node_boxes])
    notes.append(["水平线/垂直线", f"{result.diagnostics.detected_horizontal_lines}/{result.diagnostics.detected_vertical_lines}"])
    notes.append(["耗时(秒)", result.diagnostics.duration_seconds])

    for sheet in (node_sheet, edge_sheet, notes):
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for row in sheet.iter_rows():
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for column in sheet.columns:
            width = max((len(str(cell.value or "")) for cell in column), default=10)
            sheet.column_dimensions[column[0].column_letter].width = min(55, max(12, width + 2))

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


# --------------------------------------------------------------------------- PDF


def _register_font() -> str:
    candidates = [
        "/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
        "/System/Library/Fonts/PingFang.ttc",
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
        "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc",
        "C:/Windows/Fonts/msyh.ttc",
        "C:/Windows/Fonts/simhei.ttf",
    ]
    for candidate in candidates:
        path = Path(candidate)
        if not path.exists():
            continue
        try:
            pdfmetrics.registerFont(TTFont("EquityCN", str(path), subfontIndex=0))
            return "EquityCN"
        except Exception:
            try:
                pdfmetrics.registerFont(TTFont("EquityCN", str(path)))
                return "EquityCN"
            except Exception:
                continue
    return "Helvetica"


def build_pdf(result: StructureResult, layout_overrides: dict | None = None) -> bytes:
    buffer = io.BytesIO()
    page = landscape(A3)
    document = canvas.Canvas(buffer, pagesize=page)
    font = _register_font()
    _draw_pdf_chart(document, page, font, result, layout_overrides)
    document.save()
    return buffer.getvalue()


def _draw_pdf_chart(document: canvas.Canvas, page, font: str, result: StructureResult, layout_overrides: dict | None = None) -> None:
    layout = layout_structure(result, layout_overrides)
    styles = (layout_overrides or {}).get("styles") if isinstance((layout_overrides or {}).get("styles"), dict) else {}
    node_styles = styles.get("nodes") if isinstance(styles.get("nodes"), dict) else {}
    edge_styles = styles.get("edges") if isinstance(styles.get("edges"), dict) else {}
    label_styles = styles.get("labels") if isinstance(styles.get("labels"), dict) else {}
    canvas_style = styles.get("canvas") if isinstance(styles.get("canvas"), dict) else {}
    page_width, page_height = page
    margin = 24
    scale = min(
        (page_width - margin * 2) / max(1, layout.width),
        (page_height - margin * 2) / max(1, layout.height),
        1.0,
    )
    offset_x = (page_width - layout.width * scale) / 2
    offset_y = (page_height - layout.height * scale) / 2

    def x(value: float) -> float:
        return offset_x + value * scale

    def y(value: float) -> float:
        return page_height - offset_y - value * scale

    def node_style(node) -> dict:
        core = node.name == result.core_company
        fallback = {
            "fill": CORE_FILL if core else NODE_FILL,
            "stroke": CORE_STROKE if core else NODE_STROKE,
            "textColor": TEXT_COLOR,
            "fontSize": 21 if core else 15,
            "strokeWidth": 2 if core else 1.35,
            "radius": 10,
        }
        return {**fallback, **(node_styles.get(node.id) or {})}

    def edge_style(edge_id: str) -> dict:
        return {"stroke": LINE_COLOR, "strokeWidth": 1.4, "dash": "none", **((edge_styles.get(edge_id) or {}))}

    def label_style(edge_id: str) -> dict:
        return {"color": "#333333", "background": "#ffffff", "fontSize": 11, "paddingX": 5, "paddingY": 3, "radius": 3, **((label_styles.get(edge_id) or {}))}

    document.setFillColorRGB(*_hex_rgb(canvas_style.get("background"), "#ffffff"))
    document.rect(0, 0, page_width, page_height, fill=1, stroke=0)
    document.setFont(font, max(12, 22 * scale))
    document.setFillColorRGB(*_hex_rgb(TEXT_COLOR))
    document.drawCentredString(x(layout.width / 2), y(TITLE_BASELINE), result.title)

    routes = build_routes(result, layout, layout_overrides)
    for line in routes.lines:
        style = edge_style(line.edge_id)
        document.setStrokeColorRGB(*_hex_rgb(style.get("stroke"), LINE_COLOR))
        document.setLineWidth(max(0.35, float(style.get("strokeWidth", 1.4)) * scale))
        dash = str(style.get("dash") or "none")
        if dash != "none":
            document.setDash([3 * scale, 2 * scale])
        document.line(x(line.x1), y(line.y1), x(line.x2), y(line.y2))
        if dash != "none":
            document.setDash([])
    document.setFillColorRGB(*_hex_rgb(LINE_COLOR))
    document.setFont(font, max(6, 9 * scale))
    for arrow_x, arrow_y in routes.arrows:
        document.drawString(x(arrow_x - 3), y(arrow_y - 4), "▼")

    for label in routes.labels:
        style = label_style(label.edge_id)
        pad_x = float(style.get("paddingX", 5))
        pad_y = float(style.get("paddingY", 3))
        radius = float(style.get("radius", 3))
        background = _hex_rgb(style.get("background"), "#ffffff")
        color = _hex_rgb(style.get("color"), "#333333")
        document.setFillColorRGB(*background)
        document.setStrokeColorRGB(*background)
        document.roundRect(
            x(label.x), y(label.y + label.height),
            label.width * scale, label.height * scale,
            max(0, radius * scale), fill=1, stroke=0,
        )
        document.setFillColorRGB(*color)
        document.setFont(font, max(6, float(style.get("fontSize", 11)) * scale))
        document.drawString(x(label.x + pad_x), y(label.y + label.height - pad_y - 1), label.text)

    for node in result.nodes:
        if node.status == "excluded_cancelled":
            continue
        position = layout.nodes.get(node.id)
        if not position:
            continue
        style = node_style(node)
        document.setFillColorRGB(*_hex_rgb(style.get("fill"), NODE_FILL))
        document.setStrokeColorRGB(*_hex_rgb(style.get("stroke"), NODE_STROKE))
        document.setLineWidth(max(0.2, float(style.get("strokeWidth", 1.35)) * scale))
        document.roundRect(
            x(position.x), y(position.y + position.height),
            position.width * scale, position.height * scale,
            max(0, float(style.get("radius", 10)) * scale), fill=1, stroke=1,
        )
        document.setFillColorRGB(*_hex_rgb(style.get("textColor"), TEXT_COLOR))
        max_chars = max((len(line) for line in position.lines), default=1)
        fitted_font = (position.width * scale - 10) / max(1, max_chars) * 0.92
        font_size = min(max(6, float(style.get("fontSize", 15)) * scale), max(5, fitted_font))
        document.setFont(font, font_size)
        line_height = font_size * 1.42
        total = len(position.lines) * line_height
        start_y = position.y + position.height / 2 - total / 2 + font_size
        for index, line in enumerate(position.lines):
            document.drawCentredString(
                x(position.x + position.width / 2),
                y(start_y + index * line_height),
                line,
            )

    if scale < 0.55:
        document.setFont(font, 7)
        document.setFillColorRGB(0.4, 0.43, 0.47)
        document.drawString(margin, 14, "图像过大已自动缩放；完整数据见 Excel 导出。")
    document.showPage()


def build_pdf_with_tables(result: StructureResult, layout_overrides: dict | None = None) -> bytes:
    """Optional appendix export if callers need the table view as well."""
    buffer = io.BytesIO()
    page = landscape(A3)
    document = canvas.Canvas(buffer, pagesize=page)
    font = _register_font()
    _draw_pdf_chart(document, page, font, result, layout_overrides)
    _draw_pdf_tables(document, page, font, result)
    document.save()
    return buffer.getvalue()


def _draw_pdf_tables(document: canvas.Canvas, page, font: str, result: StructureResult) -> None:
    width, height = page
    margin = 36
    y = height - 40
    document.setFont(font, 15)
    document.drawString(margin, y, "股权结构（要点）")
    y -= 24
    document.setFont(font, 9)
    for node in sorted(result.nodes, key=lambda item: (item.level, item.name)):
        if y < 36:
            document.showPage()
            document.setFont(font, 9)
            y = height - 40
        document.drawString(margin, y, f"L{node.level}  {node.name}  [{node.entity_type}/{node.status}]  {node.confidence:.2f}")
        y -= 14
    y -= 10
    document.setFont(font, 15)
    document.drawString(margin, y, "投资关系")
    y -= 24
    document.setFont(font, 9)
    idmap = {node.id: node.name for node in result.nodes}
    for edge in result.edges:
        if y < 36:
            document.showPage()
            document.setFont(font, 9)
            y = height - 40
        document.drawString(
            margin, y,
            f"{idmap.get(edge.source, edge.source)} -> {idmap.get(edge.target, edge.target)}  "
            f"{edge.percentage or '-'}  {edge.relation}  {edge.confidence:.2f}  {edge.percentage_source}",
        )
        y -= 14
    document.showPage()
