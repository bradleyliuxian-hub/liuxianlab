import json
import sys
import tempfile
import unittest
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

REPO = Path(__file__).resolve().parents[2]
if str(REPO) not in sys.path:
    sys.path.insert(0, str(REPO))

from services.excel_standardizer.processor import process  # noqa: E402


class SelectableFeatureTests(unittest.TestCase):
    def _run_job(self, selection):
        with tempfile.TemporaryDirectory() as tmp:
            job_dir = Path(tmp)
            wb = Workbook()
            ws = wb.active
            ws.title = 'Sheet1'
            ws['A1'] = ' Name '
            ws['B1'] = ' Date '
            ws['C1'] = ' DateTime '
            ws['A2'] = '  Alice  '
            ws['B2'] = date(2026, 10, 16)
            ws['C2'] = datetime(2026, 10, 16, 15, 30)
            ws['A4'] = 'tail'
            wb.save(job_dir / 'input.xlsx')
            (job_dir / 'job.json').write_text(json.dumps({
                'id': 'excel_test',
                'selection': selection,
            }), encoding='utf-8')

            process(job_dir, 'default')
            result = json.loads((job_dir / 'result.json').read_text(encoding='utf-8'))
            out = load_workbook(job_dir / 'output.xlsx', data_only=False)
            try:
                values = {
                    'a2': out['Sheet1']['A2'].value,
                    'b2_format': out['Sheet1']['B2'].number_format,
                    'c2_format': out['Sheet1']['C2'].number_format,
                    'freeze': out['Sheet1'].freeze_panes,
                    'filter': out['Sheet1'].auto_filter.ref,
                    'max_row': out['Sheet1'].max_row,
                    'rules': result['rulesApplied'],
                    'selection': result['selection'],
                }
            finally:
                out.close()
            return values

    def test_only_trim_text_runs(self):
        selection = {
            'trim_text': True,
            'normalize_dates': False,
            'normalize_datetime': False,
            'header_style': False,
            'autofit_columns': False,
            'freeze_header': False,
            'auto_filter': False,
            'remove_empty_rows': False,
        }
        values = self._run_job(selection)
        self.assertEqual(values['a2'], 'Alice')
        self.assertEqual(values['rules'], ['trim_text'])
        self.assertIsNone(values['freeze'])
        self.assertIsNone(values['filter'])

    def test_date_and_datetime_are_independent(self):
        selection = {
            'trim_text': False,
            'normalize_dates': True,
            'normalize_datetime': False,
            'header_style': False,
            'autofit_columns': False,
            'freeze_header': False,
            'auto_filter': False,
            'remove_empty_rows': False,
        }
        values = self._run_job(selection)
        self.assertEqual(values['b2_format'], 'yyyy-mm-dd')
        self.assertNotEqual(values['c2_format'], 'yyyy-mm-dd hh:mm')
        self.assertEqual(values['rules'], ['normalize_dates'])

    def test_freeze_and_filter_are_independent(self):
        selection = {
            'trim_text': False,
            'normalize_dates': False,
            'normalize_datetime': False,
            'header_style': False,
            'autofit_columns': False,
            'freeze_header': True,
            'auto_filter': False,
            'remove_empty_rows': False,
        }
        values = self._run_job(selection)
        self.assertEqual(values['freeze'], 'A2')
        self.assertIsNone(values['filter'])
        self.assertEqual(values['rules'], ['freeze_header'])


if __name__ == '__main__':
    unittest.main()
