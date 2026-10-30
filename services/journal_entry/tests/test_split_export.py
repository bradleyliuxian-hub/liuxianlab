import json
import tempfile
import unittest
import zipfile
from pathlib import Path

from openpyxl import load_workbook

from services.journal_entry.normalize import normalize_file
from services.journal_entry.split_export import (
    SplitExportError,
    build_split_output,
    dimension_options,
    preview_split,
)

FIXTURES = Path(__file__).resolve().parent / "fixtures"


def make_result(tmp: Path) -> Path:
    normalized = normalize_file(
        FIXTURES / "journal_audit_split.xlsx",
        sheet_name="序时账",
        header_row=1,
        mapping={
            "date": 1, "voucher_id": 2, "account_code": 3, "account_name": 4,
            "customer_code": 5, "customer_name": 6,
            "supplier_code": 7, "supplier_name": 8,
            "counterparty_name": 9, "description": 10,
            "debit": 11, "credit": 12,
        },
    )
    payload = {
        "schema_version": 1,
        "generated_at": "2026-01-01T00:00:00+00:00",
        "vouchers": [entry.to_dict() for entry in normalized.entries],
    }
    path = tmp / "result.json"
    path.write_text(json.dumps(payload), encoding="utf-8")
    return path


def request_for(result_path: Path, **overrides):
    base = {
        "schema_version": 1,
        "generated_at": "2026-01-02T00:00:00+00:00",
        "scope": {"type": "job", "id": "job-1"},
        "jobs": [{
            "job_id": "job-1", "original_filename": "journal_audit_split.xlsx",
            "position": 0, "result_path": str(result_path),
        }],
        "filters": {},
        "group_by": "account",
        "group_options": {},
        "row_mode": "matched_only",
        "output_mode": "multi_sheet",
        "limits": {"max_groups": 200, "max_rows_per_sheet": 1000},
    }
    base.update(overrides)
    return base


class SplitPreviewTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.result_path = make_result(self.root)

    def tearDown(self):
        self.tmp.cleanup()

    def test_filter_then_group_and_row_preservation(self):
        preview = preview_split(request_for(
            self.result_path,
            filters={"months": ["2025-01"], "account_codes": ["1122"]},
            group_by="customer",
        ))
        self.assertEqual(preview["matched_lines"], 2)
        self.assertEqual(preview["matched_vouchers"], 2)
        self.assertEqual(preview["group_count"], 2)
        self.assertEqual(
            {group["display_name"] for group in preview["groups"]},
            {"客户A", "客户B"},
        )

        all_preview = preview_split(request_for(self.result_path, group_by="customer"))
        self.assertEqual(all_preview["matched_lines"], 10)
        self.assertEqual(all_preview["group_count"], 3)
        self.assertEqual(
            {group["display_name"]: group["line_count"] for group in all_preview["groups"]},
            {"客户A": 2, "客户B": 1, "未指定客户": 7},
        )
        self.assertEqual(
            sum(group["line_count"] for group in all_preview["groups"]),
            all_preview["matched_lines"],
        )

    def test_supplier_and_month_preview(self):
        supplier = preview_split(request_for(self.result_path, group_by="supplier"))
        self.assertEqual(
            {group["display_name"]: group["line_count"] for group in supplier["groups"]},
            {"供应商A": 1, "供应商B": 1, "未指定供应商": 8},
        )
        month = preview_split(request_for(self.result_path, group_by="month"))
        self.assertEqual(
            {group["display_name"]: group["line_count"] for group in month["groups"]},
            {"2025-01": 4, "2025-02": 4, "2025-03": 2},
        )
        counterparty = preview_split(request_for(self.result_path, group_by="counterparty"))
        self.assertEqual(
            {group["display_name"]: group["line_count"] for group in counterparty["groups"]},
            {"部门A": 1, "往来B": 1, "未指定往来单位": 8},
        )



    def test_multi_dimension_combinations_and_validation(self):
        account_month = preview_split(request_for(
            self.result_path, group_by=["account", "month"]
        ))
        self.assertEqual(account_month["dimensions"], ["account", "month"])
        self.assertEqual(account_month["group_count"], 7)
        self.assertEqual(
            sum(group["line_count"] for group in account_month["groups"]),
            10,
        )
        self.assertTrue(all("values" in group for group in account_month["groups"]))

        customer_month = preview_split(request_for(
            self.result_path, group_by=["customer", "month"]
        ))
        self.assertEqual(customer_month["group_count"], 6)
        self.assertEqual(
            sum(group["line_count"] for group in customer_month["groups"]),
            10,
        )
        self.assertTrue(any(
            group["display_name"].startswith("未指定客户 / 2025-01")
            for group in customer_month["groups"]
        ))

        supplier_month = preview_split(request_for(
            self.result_path, group_by=["supplier", "month"]
        ))
        self.assertEqual(supplier_month["group_count"], 5)
        self.assertEqual(
            sum(group["line_count"] for group in supplier_month["groups"]),
            10,
        )

        account_customer = preview_split(request_for(
            self.result_path, group_by=["account", "customer"]
        ))
        self.assertEqual(account_customer["group_count"], 6)
        self.assertEqual(
            sum(group["line_count"] for group in account_customer["groups"]),
            10,
        )

        reversed_dimensions = preview_split(request_for(
            self.result_path, group_by=["month", "account"]
        ))
        self.assertEqual(reversed_dimensions["dimensions"], ["month", "account"])
        self.assertEqual(
            list(reversed_dimensions["groups"][0]["values"].keys()),
            ["month", "account"],
        )
        self.assertEqual(reversed_dimensions["group_count"], 7)

        three = preview_split(request_for(
            self.result_path, group_by=["entity", "account", "month"]
        ))
        self.assertEqual(three["dimensions"], ["entity", "account", "month"])
        self.assertEqual(three["group_count"], 7)
        self.assertEqual(
            sum(group["line_count"] for group in three["groups"]), 10
        )

        prefix = preview_split(request_for(
            self.result_path, group_by=["account", "month"],
            group_options={"account_prefix_length": 2},
        ))
        self.assertTrue(any(
            group["values"].get("account") == "11" for group in prefix["groups"]
        ))

        with self.assertRaises(SplitExportError):
            preview_split(request_for(self.result_path, group_by=["customer", "customer"]))
        with self.assertRaises(SplitExportError):
            preview_split(request_for(
                self.result_path, group_by=["account", "customer", "month", "entity"]
            ))

    def test_manifest_uses_structured_dimensions(self):
        output = self.root / "multi-dim"
        result = build_split_output(output, request_for(
            self.result_path, group_by=["customer", "month"], output_mode="zip"
        ))
        with zipfile.ZipFile(output / result["output_file"]) as archive:
            manifest = json.loads(archive.read("manifest.json"))
        self.assertEqual(manifest["dimensions"], ["customer", "month"])
        self.assertTrue(all("values" in group for group in manifest["groups"]))
        self.assertTrue(all("month" in group["values"] for group in manifest["groups"]))

        multi = self.root / "multi-dim-sheet"
        multi_result = build_split_output(multi, request_for(
            self.result_path, group_by=["customer", "month"], output_mode="multi_sheet"
        ))
        workbook = load_workbook(multi / multi_result["output_file"], read_only=True, data_only=False)
        try:
            headers = [cell.value for cell in next(workbook["分组索引"].iter_rows(max_row=1))]
            self.assertEqual(headers[:2], ["客户", "月份"])
            self.assertEqual(headers[2:5], ["Sheet 名称", "部分", "行数"])
            self.assertTrue(all(len(name) <= 31 for name in workbook.sheetnames))
            self.assertEqual(len(workbook.sheetnames), len(set(workbook.sheetnames)))
        finally:
            workbook.close()


    def test_group_limit_boundaries(self):
        path = self.root / "many-groups.json"
        vouchers = []
        for index in range(501):
            customer = f"C{index:04d}"
            vouchers.append({
                "voucher_key": f"K{index}", "voucher_id": f"V{index}",
                "date": "2025-01-01", "period": "2025-01",
                "lines": [{
                    "voucher_id": f"V{index}", "sheet_name": "序时账",
                    "source_row": index + 2, "internal_line_key": f"3:序时账:{index+2}",
                    "line_id": str(index), "account_code": "1122", "account_name": "应收账款",
                    "customer_code": customer, "customer_name": f"客户{index}",
                    "debit": "1.00", "credit": "0.00",
                }],
            })
        path.write_text(json.dumps({"schema_version": 1, "vouchers": vouchers}), encoding="utf-8")
        path201 = self.root / "groups-201.json"
        path201.write_text(
            json.dumps({"schema_version": 1, "vouchers": vouchers[:201]}),
            encoding="utf-8",
        )
        base = request_for(path201, group_by=["customer"], output_mode="multi_sheet")
        for limit, exceeded in ((199, True), (200, True), (201, False)):
            request = {**base, "limits": {"max_groups": limit, "max_rows_per_sheet": 1000}}
            preview = preview_split(request)
            self.assertEqual(preview["group_count"], 201)
            self.assertEqual(preview["limit_exceeded"], exceeded, limit)
        base500 = request_for(path, group_by=["customer"], output_mode="zip")
        for limit, exceeded in ((499, True), (500, True), (501, False)):
            request = {**base500, "limits": {"max_groups": limit, "max_rows_per_sheet": 1000}}
            preview = preview_split(request)
            self.assertEqual(preview["group_count"], 501)
            self.assertEqual(preview["limit_exceeded"], exceeded, limit)

    def test_source_file_groups_keep_same_named_jobs_distinct(self):
        request = request_for(self.result_path, group_by="source_file")
        request["jobs"] = [
            {
                "job_id": "job-1", "original_filename": "journal.xlsx",
                "position": 0, "result_path": str(self.result_path),
            },
            {
                "job_id": "job-2", "original_filename": "journal.xlsx",
                "position": 1, "result_path": str(self.result_path),
            },
        ]
        preview = preview_split(request)
        self.assertEqual(preview["matched_lines"], 20)
        self.assertEqual(preview["group_count"], 2)
        self.assertEqual({group["display_name"] for group in preview["groups"]}, {"journal.xlsx"})
        self.assertEqual(
            sum(group["line_count"] for group in preview["groups"]),
            preview["matched_lines"],
        )

    def test_account_and_prefix_groups(self):
        account = preview_split(request_for(self.result_path, group_by="account"))
        self.assertEqual(account["group_count"], 5)
        prefix = preview_split(request_for(
            self.result_path, group_by="account",
            group_options={"account_prefix_length": 2},
        ))
        self.assertEqual(prefix["group_count"], 5)
        self.assertIn("11", {group["display_name"] for group in prefix["groups"]})

    def test_group_limit_and_empty_result(self):
        limited = preview_split(request_for(
            self.result_path, group_by="account",
            limits={"max_groups": 2, "max_rows_per_sheet": 1000},
        ))
        self.assertTrue(limited["limit_exceeded"])
        with self.assertRaises(SplitExportError):
            build_split_output(self.root / "limit", request_for(
                self.result_path, group_by="account",
                limits={"max_groups": 2, "max_rows_per_sheet": 1000},
            ))
        empty = preview_split(request_for(
            self.result_path, filters={"date_from": "2030-01-01"}
        ))
        self.assertEqual(empty["matched_lines"], 0)
        with self.assertRaises(SplitExportError):
            build_split_output(self.root / "empty", request_for(
                self.result_path, filters={"date_from": "2030-01-01"}
            ))

    def test_dimension_options_and_blank_count(self):
        options = dimension_options({
            "dimension": "customer", "q": "客户", "limit": 10,
            "jobs": [{
                "job_id": "job-1", "original_filename": "journal_audit_split.xlsx",
                "position": 0, "result_path": str(self.result_path),
            }],
        })
        self.assertEqual({item["value"] for item in options["options"]}, {"C001", "C002"})
        self.assertGreater(options["blank_count"], 0)


class SplitOutputTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.result_path = make_result(self.root)

    def tearDown(self):
        self.tmp.cleanup()

    def test_multi_sheet_workbook_structure(self):
        output = self.root / "multi"
        result = build_split_output(output, request_for(
            self.result_path, group_by="account", output_mode="multi_sheet"
        ))
        self.assertEqual(result["group_count"], 5)
        path = output / result["output_file"]
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            self.assertIn("拆分说明", workbook.sheetnames)
            self.assertIn("分组索引", workbook.sheetnames)
            self.assertEqual(len(workbook.sheetnames), 7)
            self.assertTrue(all(len(name) <= 31 for name in workbook.sheetnames))
            self.assertEqual(len(set(workbook.sheetnames)), len(workbook.sheetnames))
            headers = [cell.value for cell in next(workbook["分组索引"].iter_rows(max_row=1))]
            self.assertEqual(headers[0:2], ["科目编码", "科目名称"])
            self.assertEqual(headers[2:5], ["Sheet 名称", "部分", "行数"])
            group_sheet = next(name for name in workbook.sheetnames if name not in ("拆分说明", "分组索引"))
            standard_headers = [cell.value for cell in next(workbook[group_sheet].iter_rows(max_row=1))]
            self.assertIn("日期", standard_headers)
            self.assertIn("凭证号", standard_headers)
            self.assertIn("客户名称", standard_headers)
            self.assertIn("供应商名称", standard_headers)
            self.assertNotIn("内部行标识（系统追溯）", standard_headers)
        finally:
            workbook.close()

    def test_zip_manifest_and_group_files(self):
        output = self.root / "zip"
        result = build_split_output(output, request_for(
            self.result_path, group_by="customer", output_mode="zip"
        ))
        self.assertEqual(result["group_count"], 3)
        path = output / result["output_file"]
        with zipfile.ZipFile(path) as archive:
            names = archive.namelist()
            self.assertIn("manifest.json", names)
            reports = [name for name in names if name.endswith(".xlsx")]
            self.assertEqual(len(reports), 3)
            self.assertEqual(len(set(reports)), 3)
            for name in names:
                self.assertFalse(name.startswith("/"))
                self.assertFalse(".." in name.split("/"))
                self.assertNotIn("\\", name)
            manifest = json.loads(archive.read("manifest.json"))
        self.assertEqual(manifest["matched_lines"], 10)
        self.assertEqual(manifest["group_count"], 3)
        self.assertIn("未指定客户", {group["display_name"] for group in manifest["groups"]})
        manifest_without_columns = dict(manifest)
        manifest_without_columns.pop("selected_columns", None)
        self.assertNotIn("debit", json.dumps(manifest_without_columns).lower())

    def test_formula_injection_and_duplicate_group_names(self):
        payload = {
            "schema_version": 1,
            "generated_at": "2026-01-01T00:00:00+00:00",
            "vouchers": [{
                "voucher_key": "K", "voucher_id": "V", "date": "2025-01-01",
                "period": "2025-01", "entity": "E",
                "lines": [
                    {
                        "voucher_id": "V", "sheet_name": "序时账", "source_row": 2,
                        "internal_line_key": "3:序时账:2", "line_id": "001",
                        "account_code": "1122", "account_name": "+CMD",
                        "description": "@cmd", "debit": "1.00", "credit": "0.00",
                        "customer_code": "C1", "customer_name": "A/B",
                    },
                    {
                        "voucher_id": "V", "sheet_name": "序时账", "source_row": 3,
                        "internal_line_key": "3:序时账:3", "line_id": "002",
                        "account_code": "1122", "account_name": "=ACCOUNT",
                        "description": "=HYPERLINK(\"x\")", "debit": "1.00", "credit": "0.00",
                        "customer_code": "C2", "customer_name": "A\\B",
                    },
                ],
            }],
        }
        self.result_path.write_text(json.dumps(payload), encoding="utf-8")
        output = self.root / "safe"
        result = build_split_output(output, request_for(
            self.result_path, group_by="customer", output_mode="zip"
        ))
        with zipfile.ZipFile(output / result["output_file"]) as archive:
            names = archive.namelist()
            reports = [name for name in names if name.endswith(".xlsx")]
            self.assertEqual(len(reports), 2)
            self.assertEqual(len(set(reports)), 2)
        first = load_workbook(output / reports[0], read_only=True, data_only=False)
        try:
            rows = list(first["序时账"].iter_rows(values_only=True))
            headers = list(rows[0])
            row = rows[1]
            self.assertTrue(str(row[headers.index("科目名称")]).startswith("'=") or str(row[headers.index("科目名称")]).startswith("'+"))
            self.assertTrue(str(row[headers.index("摘要")]).startswith("'@") or str(row[headers.index("摘要")]).startswith("'="))
        finally:
            first.close()

        multi = self.root / "safe-multi"
        multi_result = build_split_output(multi, request_for(
            self.result_path, group_by="customer", output_mode="multi_sheet"
        ))
        workbook = load_workbook(multi / multi_result["output_file"], read_only=True, data_only=False)
        try:
            names = workbook.sheetnames
            self.assertEqual(len(names), len(set(names)))
            self.assertTrue(all(len(name) <= 31 for name in names))
        finally:
            workbook.close()


    def test_column_statistics_are_filtered_and_zero_is_not_empty(self):
        payload = {
            "schema_version": 1, "generated_at": "2026-01-01T00:00:00+00:00",
            "vouchers": [
                {"voucher_key": "K1", "voucher_id": "V1", "date": "2025-01-01",
                 "lines": [{"voucher_id": "V1", "sheet_name": "序时账", "source_row": 2,
                            "account_code": "1001", "account_name": "现金", "description": "a",
                            "customer_code": "", "customer_name": "客户A", "currency": "",
                            "original_amount": "", "debit": "0.00", "credit": "0.00",
                            "internal_line_key": "3:序时账:2"}]},
                {"voucher_key": "K2", "voucher_id": "V2", "date": "2025-02-01",
                 "lines": [{"voucher_id": "V2", "sheet_name": "序时账", "source_row": 3,
                            "account_code": "1002", "account_name": "银行", "description": "b",
                            "customer_code": "", "customer_name": "客户B", "currency": "",
                            "original_amount": "", "debit": "0.00", "credit": "0.00",
                            "internal_line_key": "3:序时账:3"}]}
            ],
        }
        self.result_path.write_text(json.dumps(payload), encoding="utf-8")
        request = request_for(self.result_path, filters={"months": ["2025-01"]}, group_by="account")
        preview = preview_split(request)
        stats = preview["column_stats"]
        self.assertTrue(stats["customer_code"]["entirely_empty"])
        self.assertFalse(stats["customer_name"]["entirely_empty"])
        self.assertTrue(stats["currency"]["entirely_empty"])
        self.assertTrue(stats["original_amount"]["entirely_empty"])
        self.assertFalse(stats["debit"]["entirely_empty"])
        self.assertFalse(stats["credit"]["entirely_empty"])
        self.assertNotIn("customer_code", preview["selected_columns"])
        self.assertIn("customer_name", preview["selected_columns"])

    def test_explicit_columns_use_uniform_headers_in_multi_and_zip(self):
        columns = ["date", "voucher_id", "description", "account_code", "account_name",
                   "debit", "credit", "customer_name", "internal_line_key"]
        expected = ["日期", "凭证号", "摘要", "科目编码", "科目名称",
                    "借方金额", "贷方金额", "客户名称", "内部行标识（系统追溯）"]
        output = self.root / "columns-multi"
        result = build_split_output(output, request_for(
            self.result_path, group_by="account", output_mode="multi_sheet",
            columns=columns,
        ))
        workbook = load_workbook(output / result["output_file"], read_only=True, data_only=False)
        try:
            sheets = [name for name in workbook.sheetnames if name not in ("拆分说明", "分组索引")]
            self.assertTrue(sheets)
            for name in sheets:
                headers = [cell.value for cell in next(workbook[name].iter_rows(max_row=1))]
                self.assertEqual(headers, expected)
        finally:
            workbook.close()
        self.assertEqual(result["selected_columns"], columns)

        zip_output = self.root / "columns-zip"
        zip_result = build_split_output(zip_output, request_for(
            self.result_path, group_by="account", output_mode="zip",
            columns=columns,
        ))
        with zipfile.ZipFile(zip_output / zip_result["output_file"]) as archive:
            reports = [name for name in archive.namelist() if name.endswith(".xlsx")]
            self.assertTrue(reports)
            for name in reports:
                with tempfile.NamedTemporaryFile(suffix=".xlsx") as temp:
                    temp.write(archive.read(name))
                    temp.flush()
                    book = load_workbook(temp.name, read_only=True, data_only=False)
                    try:
                        headers = [cell.value for cell in next(book["序时账"].iter_rows(max_row=1))]
                        self.assertEqual(headers, expected)
                    finally:
                        book.close()

    def test_column_selection_validation(self):
        with self.assertRaises(SplitExportError):
            preview_split(request_for(self.result_path, columns=[]))
        with self.assertRaises(SplitExportError):
            preview_split(request_for(self.result_path, columns=["unknown"]))
        with self.assertRaises(SplitExportError):
            preview_split(request_for(self.result_path, columns=["date", "date"]))

    def test_row_limit_splits_sheets_and_zip_parts(self):
        output = self.root / "rows"
        result = build_split_output(output, request_for(
            self.result_path, group_by="account",
            output_mode="multi_sheet",
            limits={"max_groups": 200, "max_rows_per_sheet": 1},
        ))
        workbook = load_workbook(output / result["output_file"], read_only=True, data_only=False)
        try:
            index_rows = list(workbook["分组索引"].iter_rows(values_only=True))
            index_headers = list(index_rows[0])
            part_index = index_headers.index("部分")
            self.assertTrue(any(row[part_index] > 1 for row in index_rows[1:]))
        finally:
            workbook.close()

        zip_output = self.root / "row-zip"
        zip_result = build_split_output(zip_output, request_for(
            self.result_path, group_by="account",
            output_mode="zip",
            limits={"max_groups": 500, "max_rows_per_sheet": 1},
        ))
        with zipfile.ZipFile(zip_output / zip_result["output_file"]) as archive:
            manifest = json.loads(archive.read("manifest.json"))
        self.assertTrue(any(group.get("part_count", 1) > 1 for group in manifest["groups"]))
        self.assertEqual(len({group["output_filename"] for group in manifest["groups"]}), len(manifest["groups"]))


if __name__ == "__main__":
    unittest.main()


class AccountSheetTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def _payload(self, count=201):
        vouchers = []
        for index in range(count):
            code = f"{100000 + index:06d}"
            name = f"科目{index}"
            vouchers.append({
                "voucher_key": f"K{index:05d}",
                "voucher_id": f"V{index:05d}",
                "date": "2025-01-01",
                "period": "2025-01",
                "entity": "E1",
                "lines": [{
                    "voucher_id": f"V{index:05d}",
                    "sheet_name": "序时账",
                    "source_row": index + 2,
                    "internal_line_key": f"3:序时账:{index + 2}",
                    "line_id": str(index),
                    "account_code": code,
                    "account_name": name,
                    "debit": "1.00",
                    "credit": "0.00",
                }],
            })
        path = self.root / "accounts.json"
        path.write_text(json.dumps({"schema_version": 1, "vouchers": vouchers}), encoding="utf-8")
        return path

    def test_one_sheet_per_full_account_code(self):
        path = self.root / "full-codes.json"
        vouchers = []
        rows = [
            ("100201", "招商银行", "V001", 1),
            ("100201", "招商银行", "V002", 2),
            ("112201", "应收账款", "V003", 3),
            ("220201", "应付账款", "V004", 4),
            ("660201", "管理费用", "V005", 5),
            ("600101", "主营业务收入", "V006", 6),
            ("", "", "V007", 7),
        ]
        for code, name, voucher_id, row in rows:
            vouchers.append({
                "voucher_key": f"K{voucher_id}",
                "voucher_id": voucher_id,
                "date": "2025-01-01",
                "period": "2025-01",
                "entity": "E1",
                "lines": [{
                    "voucher_id": voucher_id,
                    "sheet_name": "序时账",
                    "source_row": row,
                    "internal_line_key": f"3:序时账:{row}",
                    "line_id": str(row),
                    "account_code": code,
                    "account_name": name,
                    "debit": "1.00",
                    "credit": "0.00",
                }],
            })
        path.write_text(json.dumps({"schema_version": 1, "vouchers": vouchers}), encoding="utf-8")
        output = self.root / "account-sheets"
        result = build_split_output(output, request_for(
            path,
            group_by="account",
            output_mode="multi_sheet",
            columns=["date", "voucher_id", "internal_line_key", "account_code", "account_name"],
        ))
        self.assertEqual(result["group_count"], 6)
        from services.journal_entry.split_export import preview_rows
        preview = preview_rows(request_for(
            path, group_by="account", output_mode="multi_sheet",
        ))
        workbook = load_workbook(output / result["output_file"], read_only=True, data_only=True)
        try:
            data_sheets = [name for name in workbook.sheetnames if name not in ("拆分说明", "分组索引")]
            self.assertEqual(
                [group["sheet_name"] for group in preview["groups"]],
                data_sheets,
            )
            self.assertEqual(len(data_sheets), 6)
            all_keys = []
            for name in data_sheets:
                headers = [cell.value for cell in next(workbook[name].iter_rows(max_row=1))]
                self.assertEqual(
                    headers,
                    ["日期", "凭证号", "内部行标识（系统追溯）", "科目编码", "科目名称"],
                )
                key_index = headers.index("内部行标识（系统追溯）")
                all_keys.extend(
                    row[key_index]
                    for row in workbook[name].iter_rows(min_row=2, values_only=True)
                )
            self.assertEqual(len(all_keys), 7)
            self.assertEqual(len(set(all_keys)), 7)

            index = list(workbook["分组索引"].iter_rows(values_only=True))
            index_headers = list(index[0])
            self.assertEqual(
                index_headers,
                ["科目编码", "科目名称", "Sheet 名称", "部分", "行数", "凭证数", "借方合计", "贷方合计"],
            )
            code_i = index_headers.index("科目编码")
            name_i = index_headers.index("科目名称")
            sheet_i = index_headers.index("Sheet 名称")
            line_i = index_headers.index("行数")
            voucher_i = index_headers.index("凭证数")
            by_code = {str(row[code_i]): row for row in index[1:]}
            self.assertEqual(len(by_code), 6)
            self.assertEqual(str(by_code["100201"][line_i]), "2")
            self.assertEqual(str(by_code["100201"][voucher_i]), "2")
            self.assertIn("100201", str(by_code["100201"][sheet_i]))
            self.assertEqual(str(by_code["112201"][line_i]), "1")
            self.assertEqual(by_code["未指定"][name_i], "未指定科目")
            self.assertEqual(
                sum(int(row[line_i]) for row in index[1:]),
                7,
            )
            self.assertEqual(
                {str(row[sheet_i]).split("_")[0] for row in index[1:]},
                {str(name).split("_")[0] for name in data_sheets},
            )
        finally:
            workbook.close()

    def test_group_scoped_preview_matches_workbook_sheets(self):
        from services.journal_entry.split_export import preview_rows
        path = self.root / "scoped.json"
        specs = [
            ("100201", "招商银行", "V001", 1),
            ("100201", "招商银行", "V002", 2),
            ("112201", "应收账款", "V003", 3),
            ("220201", "应付账款", "V004", 4),
            ("660201", "管理费用", "V005", 5),
        ]
        vouchers = []
        for code, name, voucher_id, row in specs:
            vouchers.append({
                "voucher_key": f"K{voucher_id}", "voucher_id": voucher_id,
                "date": "2025-01-01", "period": "2025-01", "entity": "E1",
                "lines": [{
                    "voucher_id": voucher_id, "sheet_name": "序时账",
                    "source_row": row, "internal_line_key": f"3:序时账:{row}",
                    "line_id": str(row), "account_code": code, "account_name": name,
                    "debit": "1.00", "credit": "0.00",
                }],
            })
        path.write_text(json.dumps({"schema_version": 1, "vouchers": vouchers}), encoding="utf-8")
        request = request_for(
            path, group_by="account", output_mode="multi_sheet",
            columns=["internal_line_key", "account_code", "account_name"],
        )
        base = preview_rows(request)
        self.assertEqual(base["group_count"], 4)
        self.assertEqual(base["data_sheet_count"], 4)
        self.assertEqual(base["aux_sheet_count"], 2)
        self.assertEqual(
            [group["sheet_name"] for group in base["groups"]],
            [group["sheets"][0]["sheet_name"] for group in base["groups"]],
        )
        for group in base["groups"]:
            self.assertIn("group_ref", group)
            scoped = preview_rows({
                **request,
                "preview_group": group["group_ref"],
                "page": 1,
            })
            self.assertEqual(scoped["active_sheet"]["sheet_name"], group["sheet_name"])
            self.assertEqual(scoped["scope_line_count"], group["line_count"])
            self.assertTrue(scoped["rows"])
            code = group["raw_values"]["account"]
            self.assertTrue(all(
                row["values"]["account_code"] == code for row in scoped["rows"]
            ))
        with self.assertRaises(SplitExportError):
            preview_rows({
                **request,
                "preview_group": {
                    "dimensions": [{"dimension": "account", "value": "不存在"}]
                },
            })
        with self.assertRaises(SplitExportError):
            preview_rows({
                **request,
                "preview_group": {
                    "dimensions": [{"dimension": "month", "value": "2025-01"}]
                },
            })
        output = self.root / "scoped-workbook"
        result = build_split_output(output, request)
        workbook = load_workbook(output / result["output_file"], read_only=True)
        try:
            data_sheets = [
                name for name in workbook.sheetnames
                if name not in ("拆分说明", "分组索引")
            ]
            self.assertEqual(
                data_sheets,
                [group["sheet_name"] for group in base["groups"]],
            )
        finally:
            workbook.close()

    def test_group_scoped_preview_keeps_excluded_rows_visible(self):
        from services.journal_entry.split_export import preview_rows
        path = self.root / "excluded-scope.json"
        vouchers = []
        rows = [
            ("100201", "招商银行", "V001", 1),
            ("100201", "招商银行", "V002", 2),
            ("112201", "应收账款", "V003", 3),
        ]
        for code, name, voucher_id, row in rows:
            vouchers.append({
                "voucher_key": f"K{voucher_id}", "voucher_id": voucher_id,
                "date": "2025-01-01", "period": "2025-01", "entity": "E1",
                "lines": [{
                    "voucher_id": voucher_id, "sheet_name": "序时账",
                    "source_row": row, "internal_line_key": f"3:序时账:{row}",
                    "line_id": str(row), "account_code": code, "account_name": name,
                    "debit": "1.00", "credit": "0.00",
                }],
            })
        path.write_text(json.dumps({"schema_version": 1, "vouchers": vouchers}), encoding="utf-8")
        request = request_for(
            path, group_by="account", output_mode="multi_sheet",
            columns=["internal_line_key", "account_code"],
        )
        group = next(
            item for item in preview_rows(request)["groups"]
            if item["raw_values"]["account"] == "100201"
        )
        excluded_key = "3:序时账:2"
        scoped = preview_rows({
            **request,
            "row_selection_mode": "all_except",
            "excluded_rows": [{"job_id": "job-1", "internal_line_key": excluded_key}],
            "preview_group": group["group_ref"],
        })
        self.assertEqual(scoped["scope_line_count"], 2)
        self.assertEqual(scoped["scope_included_line_count"], 1)
        self.assertEqual(scoped["scope_excluded_line_count"], 1)
        self.assertEqual(scoped["active_sheet"]["line_count"], 1)
        rows_by_key = {row["internal_line_key"]: row for row in scoped["rows"]}
        self.assertIn(excluded_key, rows_by_key)
        self.assertFalse(rows_by_key[excluded_key]["included"])

    def test_group_scoped_preview_respects_sheet_parts(self):
        from services.journal_entry.split_export import preview_rows
        path = self.root / "parts.json"
        vouchers = []
        for index in range(3):
            vouchers.append({
                "voucher_key": f"K{index}", "voucher_id": f"V{index}",
                "date": "2025-01-01", "period": "2025-01", "entity": "E1",
                "lines": [{
                    "voucher_id": f"V{index}", "sheet_name": "序时账",
                    "source_row": index + 2, "internal_line_key": f"3:序时账:{index + 2}",
                    "line_id": str(index), "account_code": "100201", "account_name": "招商银行",
                    "debit": "1.00", "credit": "0.00",
                }],
            })
        path.write_text(json.dumps({"schema_version": 1, "vouchers": vouchers}), encoding="utf-8")
        request = request_for(
            path, group_by="account", output_mode="multi_sheet",
            limits={"max_groups": 200, "max_rows_per_sheet": 1},
            columns=["internal_line_key", "account_code"],
        )
        base = preview_rows(request)
        group = base["groups"][0]
        self.assertEqual(len(group["sheets"]), 3)
        self.assertEqual(base["data_sheet_count"], 3)
        second = preview_rows({
            **request,
            "preview_group": group["group_ref"],
            "preview_part": 2,
        })
        self.assertEqual(second["active_sheet"]["sheet_name"], group["sheets"][1]["sheet_name"])
        self.assertEqual(second["scope_line_count"], 1)
        self.assertEqual(len(second["rows"]), 1)
        with self.assertRaises(SplitExportError):
            preview_rows({
                **request,
                "preview_group": group["group_ref"],
                "preview_part": 4,
            })

    def test_account_sheet_names_are_safe_and_unique(self):
        path = self.root / "names.json"
        long_name = "非常长的科目名称" * 8
        codes = [
            ("1001", "同名科目"),
            ("1002", "同名科目"),
            ("1003", long_name),
            ("1004", "非法/名称:*?"),
            ("1005", long_name),
            ("123456789012345678901234567890123", "超长编码"),
        ]
        vouchers = []
        for index, (code, name) in enumerate(codes):
            vouchers.append({
                "voucher_key": f"K{index}",
                "voucher_id": f"V{index}",
                "date": "2025-01-01",
                "period": "2025-01",
                "entity": "E1",
                "lines": [{
                    "voucher_id": f"V{index}",
                    "sheet_name": "序时账",
                    "source_row": index + 2,
                    "internal_line_key": f"3:序时账:{index + 2}",
                    "line_id": str(index),
                    "account_code": code,
                    "account_name": name,
                    "debit": "1.00",
                    "credit": "0.00",
                }],
            })
        path.write_text(json.dumps({"schema_version": 1, "vouchers": vouchers}), encoding="utf-8")
        output = self.root / "safe-names"
        result = build_split_output(output, request_for(
            path, group_by="account", output_mode="multi_sheet",
        ))
        workbook = load_workbook(output / result["output_file"], read_only=True, data_only=True)
        try:
            data_sheets = [name for name in workbook.sheetnames if name not in ("拆分说明", "分组索引")]
            self.assertEqual(len(data_sheets), 6)
            self.assertEqual(len(data_sheets), len(set(data_sheets)))
            self.assertTrue(all(len(name) <= 31 for name in data_sheets))
            self.assertTrue(all(code in " ".join(sheet for sheet in data_sheets)
                                for code, _ in codes[:5]))
            index = list(workbook["分组索引"].iter_rows(values_only=True))
            headers = list(index[0])
            sheet_i = headers.index("Sheet 名称")
            self.assertEqual(
                {str(row[sheet_i]) for row in index[1:]},
                set(data_sheets),
            )
        finally:
            workbook.close()

    def test_account_prefix_is_opt_in_and_group_limit_is_enforced(self):
        path = self._payload(201)
        base = request_for(path, group_by="account", output_mode="multi_sheet")
        full = preview_split(base)
        self.assertEqual(full["group_count"], 201)
        self.assertTrue(full["limit_exceeded"])
        with self.assertRaises(SplitExportError):
            build_split_output(self.root / "too-many", base)
        prefixed = preview_split({
            **base,
            "group_options": {"account_prefix_length": 3},
            "limits": {"max_groups": 200, "max_rows_per_sheet": 1000},
        })
        self.assertEqual(prefixed["group_count"], 1)
        self.assertEqual(prefixed["groups"][0]["values"]["account"], "100")


class RowSelectionPreviewTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.result_path = make_result(self.root)

    def tearDown(self):
        self.tmp.cleanup()

    def _rows(self, **overrides):
        from services.journal_entry.split_export import preview_rows
        return preview_rows(request_for(self.result_path, group_by="customer", **overrides))

    def test_all_except_row_selection_and_stats(self):
        first = self._rows()
        key = first["rows"][0]["internal_line_key"]
        excluded = [{"job_id": "job-1", "internal_line_key": key}]
        preview = self._rows(
            row_selection_mode="all_except", excluded_rows=excluded
        )
        self.assertEqual(preview["candidate_line_count"], 10)
        self.assertEqual(preview["included_line_count"], 9)
        self.assertEqual(preview["excluded_line_count"], 1)
        self.assertFalse(preview["rows"][0]["included"])
        self.assertEqual(preview["row_selection_mode"], "all_except")
        self.assertEqual(preview["excluded_rows"], excluded)
        groups = {
            group["display_name"]: group["line_count"]
            for group in preview["groups"]
        }
        self.assertEqual(groups["客户A"], 1)

    def test_all_customer_rows_excluded_marks_column_empty(self):
        preview = self._rows()
        customer_keys = [
            {"job_id": "job-1", "internal_line_key": row["internal_line_key"]}
            for row in preview["rows"]
            if row["values"]["customer_name"]
        ]
        self.assertEqual(len(customer_keys), 3)
        filtered = self._rows(
            row_selection_mode="all_except", excluded_rows=customer_keys
        )
        self.assertTrue(filtered["column_stats"]["customer_name"]["entirely_empty"])
        self.assertEqual(filtered["column_stats"]["customer_name"]["non_empty_count"], 0)
        self.assertEqual(filtered["included_line_count"], 7)

    def test_only_selected_mode(self):
        preview = self._rows()
        selected = [
            {"job_id": "job-1", "internal_line_key": preview["rows"][0]["internal_line_key"]},
            {"job_id": "job-1", "internal_line_key": preview["rows"][4]["internal_line_key"]},
        ]
        filtered = self._rows(
            row_selection_mode="only_selected", included_rows=selected
        )
        self.assertEqual(filtered["included_line_count"], 2)
        self.assertEqual(filtered["excluded_line_count"], 8)
        self.assertEqual(
            [row["internal_line_key"] for row in filtered["rows"] if row["included"]],
            [selected[0]["internal_line_key"], selected[1]["internal_line_key"]],
        )
        self.assertEqual(
            filtered["column_stats"]["customer_name"]["non_empty_count"], 1
        )

    def test_row_selection_validation(self):
        from services.journal_entry.split_export import preview_rows
        with self.assertRaises(SplitExportError):
            preview_rows(request_for(
                self.result_path,
                row_selection_mode="only_selected",
                included_rows=[],
            ))
        with self.assertRaises(SplitExportError):
            preview_rows(request_for(
                self.result_path,
                row_selection_mode="all_except",
                excluded_rows=[{"job_id": "other-job", "internal_line_key": "3:序时账:2"}],
            ))
        with self.assertRaises(SplitExportError):
            preview_rows(request_for(
                self.result_path,
                row_selection_mode="all_except",
                excluded_rows=[{"job_id": "job-1", "internal_line_key": "nope"}],
            ))
        with self.assertRaises(SplitExportError):
            preview_rows(request_for(
                self.result_path,
                row_selection_mode="all_except",
                excluded_rows=[{"job_id": "job-1", "internal_line_key": "3:序时账:2"}],
                included_rows=[{"job_id": "job-1", "internal_line_key": "3:序时账:3"}],
            ))

    def test_multi_sheet_and_zip_exclude_rows(self):
        preview = self._rows(
            row_selection_mode="all_except",
            excluded_rows=[{"job_id": "job-1", "internal_line_key": "3:序时账:2"}],
            columns=["date", "voucher_id", "internal_line_key", "debit", "credit"],
        )
        self.assertEqual(preview["included_line_count"], 9)
        for mode, suffix in (("multi_sheet", "xlsx"), ("zip", "zip")):
            out_dir = self.root / f"out-{mode}"
            result = build_split_output(out_dir, request_for(
                self.result_path,
                group_by="customer",
                output_mode=mode,
                row_selection_mode="all_except",
                excluded_rows=[{"job_id": "job-1", "internal_line_key": "3:序时账:2"}],
                columns=["date", "voucher_id", "internal_line_key", "debit", "credit"],
            ))
            self.assertEqual(result["matched_lines"], 9)
            keys = []
            if mode == "multi_sheet":
                workbook = load_workbook(result_path := out_dir / "split-export.xlsx")
                for sheet in workbook.worksheets:
                    if sheet.title in ("拆分说明", "分组索引"):
                        continue
                    headers = [cell.value for cell in sheet[1]]
                    idx = headers.index("内部行标识（系统追溯）")
                    keys.extend(row[idx] for row in sheet.iter_rows(min_row=2, values_only=True))
                workbook.close()
            else:
                with zipfile.ZipFile(out_dir / "split-export.zip") as archive:
                    manifest = json.loads(archive.read("manifest.json"))
                    self.assertEqual(manifest["row_selection_mode"], "all_except")
                    for item in manifest["groups"]:
                        workbook = load_workbook(out_dir / item["output_filename"])
                        sheet = workbook["序时账"]
                        headers = [cell.value for cell in sheet[1]]
                        idx = headers.index("内部行标识（系统追溯）")
                        keys.extend(
                            row[idx]
                            for row in sheet.iter_rows(min_row=2, values_only=True)
                        )
                        workbook.close()
            self.assertEqual(len(keys), 9)
            self.assertNotIn("3:序时账:2", keys)


class StandardExportTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.result_path = make_result(self.root)

    def tearDown(self):
        self.tmp.cleanup()

    def test_standard_preview_and_export(self):
        from services.journal_entry.standard_export import (
            build_standard_output,
            preview_standard,
        )
        request = request_for(self.result_path)
        preview = preview_standard(request)
        self.assertEqual(preview["candidate_line_count"], 10)
        self.assertEqual(preview["included_line_count"], 10)
        self.assertTrue(preview["column_stats"]["customer_name"]["non_empty_count"] == 3)
        out_dir = self.root / "standard"
        result = build_standard_output(out_dir, request_for(
            self.result_path,
            row_selection_mode="all_except",
            excluded_rows=[{"job_id": "job-1", "internal_line_key": "3:序时账:2"}],
            columns=["date", "voucher_id", "internal_line_key", "debit", "credit"],
        ))
        self.assertEqual(result["included_line_count"], 9)
        workbook = load_workbook(out_dir / "standard-journal.xlsx")
        sheet = workbook["标准分录"]
        headers = [cell.value for cell in sheet[1]]
        self.assertEqual(headers, ["日期", "凭证号", "内部行标识（系统追溯）", "借方金额", "贷方金额"])
        keys = [row[2] for row in sheet.iter_rows(min_row=2, values_only=True)]
        self.assertEqual(len(keys), 9)
        self.assertNotIn("3:序时账:2", keys)
        workbook.close()


    def test_custom_full_report_only_changes_standard_sheet(self):
        from services.journal_entry.standard_export import build_custom_report
        from services.journal_entry.export import build_export
        request = {
            **request_for(self.result_path),
            "row_selection_mode": "all_except",
            "excluded_rows": [{"job_id": "job-1", "internal_line_key": "3:序时账:2"}],
            "columns": ["date", "voucher_id", "internal_line_key", "debit", "credit"],
        }
        default_path = self.root / "default-report.xlsx"
        payload = json.loads(self.result_path.read_text(encoding="utf-8"))
        build_export(default_path, payload)
        out_dir = self.root / "custom-report"
        result = build_custom_report(out_dir, request)
        self.assertEqual(result["included_line_count"], 9)
        custom_path = out_dir / "custom-analysis-report.xlsx"

        default_wb = load_workbook(default_path, read_only=True)
        custom_wb = load_workbook(custom_path, read_only=True)
        try:
            self.assertEqual(default_wb.sheetnames, custom_wb.sheetnames)
            custom_sheet = custom_wb["清洗分录"]
            headers = [cell.value for cell in next(custom_sheet.iter_rows(max_row=1))]
            self.assertEqual(
                headers,
                ["日期", "凭证号", "内部行标识（系统追溯）", "借方金额", "贷方金额"],
            )
            keys = [
                row[2]
                for row in custom_sheet.iter_rows(min_row=2, values_only=True)
            ]
            self.assertEqual(len(keys), 9)
            self.assertNotIn("3:序时账:2", keys)
            for name in ("凭证汇总", "业务分类", "对方科目", "异常凭证", "数据质量问题"):
                self.assertEqual(
                    default_wb[name].max_row,
                    custom_wb[name].max_row,
                    f"{name} row count changed",
                )
            self.assertEqual(default_wb["凭证汇总"].max_row, custom_wb["凭证汇总"].max_row)
        finally:
            default_wb.close()
            custom_wb.close()


class RowSelectionPaginationAndBatchTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def _synthetic_result(self, name="synthetic.json", count=60):
        vouchers = []
        for index in range(count):
            line = {
                "entity": "E1", "period": "2025-01",
                "date": "2025-01-01", "voucher_id": f"V{index:05d}",
                "line_id": f"L{index}", "internal_line_key": f"3:Sheet:{index + 2}",
                "account_code": "1002", "account_name": "银行存款",
                "detail_account": "", "customer_code": "", "customer_name": "",
                "supplier_code": "", "supplier_name": "",
                "counterparty_code": "", "counterparty_name": "",
                "debit": "10.00", "credit": "0.00", "original_amount": None,
                "currency": "CNY", "description": "synthetic", "preparer": "tester",
                "sheet_name": "Sheet", "source_row": index + 2,
            }
            vouchers.append({
                "voucher_key": f"K{index:05d}", "voucher_id": f"V{index:05d}",
                "date": "2025-01-01", "period": "2025-01", "entity": "E1",
                "lines": [line],
            })
        path = self.root / name
        path.write_text(json.dumps({"schema_version": 1, "vouchers": vouchers}), encoding="utf-8")
        return path

    def test_preview_rows_pagination_and_exclusion_persistence(self):
        from services.journal_entry.split_export import preview_rows
        result_path = self._synthetic_result()
        request = request_for(result_path, group_by="account")
        first = preview_rows({**request, "page": 1, "page_size": 25})
        self.assertEqual(first["candidate_line_count"], 60)
        self.assertEqual(first["total_pages"], 3)
        self.assertEqual(len(first["rows"]), 25)
        excluded = [{
            "job_id": "job-1",
            "internal_line_key": first["rows"][3]["internal_line_key"],
        }]
        page_one = preview_rows({
            **request, "page": 1, "page_size": 25,
            "row_selection_mode": "all_except", "excluded_rows": excluded,
        })
        self.assertFalse(page_one["rows"][3]["included"])
        self.assertEqual(page_one["included_line_count"], 59)
        page_two = preview_rows({
            **request, "page": 2, "page_size": 25,
            "row_selection_mode": "all_except", "excluded_rows": excluded,
        })
        self.assertEqual(page_two["page"], 2)
        self.assertEqual(page_two["included_line_count"], 59)
        back = preview_rows({
            **request, "page": 1, "page_size": 25,
            "row_selection_mode": "all_except", "excluded_rows": excluded,
        })
        self.assertFalse(back["rows"][3]["included"])
        self.assertEqual(
            back["rows"][3]["internal_line_key"],
            first["rows"][3]["internal_line_key"],
        )

    def test_batch_scope_excludes_only_target_file_row(self):
        from services.journal_entry.split_export import preview_rows, build_split_output
        path_a = self._synthetic_result("a.json", 5)
        path_b = self._synthetic_result("b.json", 5)
        jobs = [
            {"job_id": "job-a", "original_filename": "a.xlsx", "position": 0, "result_path": str(path_a)},
            {"job_id": "job-b", "original_filename": "b.xlsx", "position": 1, "result_path": str(path_b)},
        ]
        request = request_for(path_a, jobs=jobs, group_by="account")
        excluded = [{"job_id": "job-a", "internal_line_key": "3:Sheet:3"}]
        preview = preview_rows({
            **request, "row_selection_mode": "all_except", "excluded_rows": excluded,
        })
        self.assertEqual(preview["candidate_line_count"], 10)
        self.assertEqual(preview["included_line_count"], 9)
        out_dir = self.root / "batch-out"
        result = build_split_output(out_dir, {
            **request, "row_selection_mode": "all_except",
            "excluded_rows": excluded,
            "columns": ["internal_line_key", "original_filename"],
        })
        self.assertEqual(result["matched_lines"], 9)
        workbook = load_workbook(out_dir / "split-export.xlsx")
        pairs = []
        for sheet in workbook.worksheets:
            if sheet.title in ("拆分说明", "分组索引"):
                continue
            headers = [cell.value for cell in sheet[1]]
            key_idx = headers.index("内部行标识（系统追溯）")
            file_idx = headers.index("原始文件")
            pairs.extend(
                (row[file_idx], row[key_idx])
                for row in sheet.iter_rows(min_row=2, values_only=True)
            )
        workbook.close()
        self.assertEqual(len(pairs), 9)
        self.assertNotIn(("a.xlsx", "3:Sheet:3"), pairs)
        self.assertIn(("b.xlsx", "3:Sheet:3"), pairs)

    def test_only_selected_pagination_and_zero_false_are_not_empty(self):
        from services.journal_entry.split_export import preview_rows
        result_path = self._synthetic_result("only-selected.json", 60)
        payload = json.loads(result_path.read_text(encoding="utf-8"))
        payload["vouchers"][0]["lines"][0]["currency"] = False
        payload["vouchers"][0]["lines"][0]["original_amount"] = "0.00"
        payload["vouchers"][1]["lines"][0]["original_amount"] = "0.00"
        result_path.write_text(json.dumps(payload), encoding="utf-8")
        request = request_for(result_path, group_by="account")
        first = preview_rows({**request, "page": 1, "page_size": 25})
        selected = [
            {"job_id": "job-1", "internal_line_key": row["internal_line_key"]}
            for row in first["rows"][:2]
        ]
        page_one = preview_rows({
            **request, "page": 1, "page_size": 25,
            "row_selection_mode": "only_selected", "included_rows": selected,
        })
        self.assertEqual(page_one["included_line_count"], 2)
        self.assertEqual(page_one["excluded_line_count"], 58)
        self.assertEqual(
            [row["internal_line_key"] for row in page_one["rows"] if row["included"]],
            [selected[0]["internal_line_key"], selected[1]["internal_line_key"]],
        )
        page_two = preview_rows({
            **request, "page": 2, "page_size": 25,
            "row_selection_mode": "only_selected", "included_rows": selected,
        })
        self.assertEqual(page_two["included_line_count"], 2)
        self.assertFalse(any(row["included"] for row in page_two["rows"]))
        back = preview_rows({
            **request, "page": 1, "page_size": 25,
            "row_selection_mode": "only_selected", "included_rows": selected,
        })
        self.assertTrue(back["rows"][0]["included"])
        self.assertTrue(back["rows"][1]["included"])
        self.assertTrue(back["column_stats"]["currency"]["non_empty_count"] > 0)
        self.assertTrue(back["column_stats"]["original_amount"]["non_empty_count"] > 0)
        self.assertEqual(back["column_stats"]["currency"]["entirely_empty"], False)
        self.assertEqual(back["column_stats"]["original_amount"]["entirely_empty"], False)


class SummaryStructureExportTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        from services.journal_entry.summary_structure import infer_summary_structure
        vouchers = []
        rows = [
            ("V001", "224103", "员工垫付", "费用报销_陈威涛", 1),
            ("V002", "224103", "员工垫付", "费用报销_丁家蕙", 2),
            ("V003", "100201", "银行存款", "银行手续费_2025.01.07", 3),
            ("V004", "100201", "银行存款", "银行手续费_-", 4),
        ]
        for voucher_id, code, name, description, row in rows:
            vouchers.append({
                "voucher_key": f"K{voucher_id}",
                "voucher_id": voucher_id,
                "date": "2025-01-01",
                "period": "2025-01",
                "entity": "E1",
                "lines": [{
                    "voucher_id": voucher_id,
                    "sheet_name": "序时账",
                    "source_row": row,
                    "internal_line_key": f"3:序时账:{row}",
                    "line_id": str(row),
                    "account_code": code,
                    "account_name": name,
                    "description": description,
                    "debit": "1.00",
                    "credit": "0.00",
                }],
            })
        payload = {"schema_version": 1, "vouchers": vouchers}
        payload["summary_structure"] = infer_summary_structure(vouchers)
        self.result_path = self.root / "result.json"
        self.result_path.write_text(json.dumps(payload), encoding="utf-8")

    def tearDown(self):
        self.tmp.cleanup()

    def _request(self, **overrides):
        return request_for(
            self.result_path,
            group_by="account",
            output_mode="multi_sheet",
            columns=["date", "voucher_id", "description", "account_code", "account_name"],
            **overrides,
        )

    def test_preview_and_excel_derived_fields_match(self):
        from services.journal_entry.split_export import preview_rows
        request = self._request()
        base = preview_rows(request)
        self.assertEqual(base["summary_structure"]["auto_rule_count"], 2)
        self.assertEqual(base["summary_structure"]["covered_line_count"], 4)
        group = next(item for item in base["groups"] if item["account_code"] == "224103")
        scoped = preview_rows({**request, "preview_group": group["group_ref"]})
        self.assertEqual(
            [field["label"] for field in scoped["derived_columns"]],
            ["费用项目", "员工姓名"],
        )
        self.assertEqual(scoped["rows"][0]["values"]["summary:auto:224103:part:1"], "陈威涛")

        output = self.root / "derived"
        result = build_split_output(output, request)
        workbook = load_workbook(output / result["output_file"], read_only=True, data_only=True)
        try:
            sheet = workbook["224103 员工垫付"]
            headers = [cell.value for cell in next(sheet.iter_rows(max_row=1))]
            self.assertEqual(
                headers,
                ["日期", "凭证号", "摘要", "费用项目", "员工姓名", "科目编码", "科目名称"],
            )
            rows = list(sheet.iter_rows(min_row=2, values_only=True))
            self.assertEqual(rows[0][3], "费用报销")
            self.assertEqual(rows[0][4], "陈威涛")
            bank = workbook["100201 银行存款"]
            bank_headers = [cell.value for cell in next(bank.iter_rows(max_row=1))]
            self.assertIn("交易日期", bank_headers)
            bank_rows = list(bank.iter_rows(min_row=2, values_only=True))
            self.assertIsNotNone(bank_rows[0][bank_headers.index("交易日期")])
            self.assertIsNone(bank_rows[1][bank_headers.index("交易日期")])
        finally:
            workbook.close()

    def test_override_hides_and_renames_derived_field(self):
        request = self._request()
        group = next(
            item for item in preview_split(request)["groups"]
            if item["values"].get("account", "").startswith("224103")
        )
        # preview_split groups do not expose group_ref; use preview_rows.
        from services.journal_entry.split_export import preview_rows
        base = preview_rows(request)
        group = next(item for item in base["groups"] if item["account_code"] == "224103")
        override = {
            "224103": {
                "fields": {
                    "summary:auto:224103:part:1": {
                        "selected": False,
                    }
                }
            }
        }
        scoped = preview_rows({
            **request,
            "preview_group": group["group_ref"],
            "summary_overrides": override,
        })
        self.assertEqual(
            [field["label"] for field in scoped["derived_columns"]],
            ["费用项目"],
        )
        output = self.root / "override"
        result = build_split_output(output, {**request, "summary_overrides": override})
        workbook = load_workbook(output / result["output_file"], read_only=True, data_only=True)
        try:
            headers = [cell.value for cell in next(workbook["224103 员工垫付"].iter_rows(max_row=1))]
            self.assertNotIn("员工姓名", headers)
            renamed = {
                "224103": {
                    "fields": {
                        "summary:auto:224103:part:1": {
                            "selected": True,
                            "label": "报销人",
                        }
                    }
                }
            }
        finally:
            workbook.close()
        output2 = self.root / "rename"
        result2 = build_split_output(output2, {**request, "summary_overrides": renamed})
        workbook = load_workbook(output2 / result2["output_file"], read_only=True, data_only=True)
        try:
            headers = [cell.value for cell in next(workbook["224103 员工垫付"].iter_rows(max_row=1))]
            self.assertIn("报销人", headers)
            self.assertNotIn("员工姓名", headers)
        finally:
            workbook.close()

    def test_manifest_and_formula_safe_derived_fields(self):
        from services.journal_entry.summary_structure import infer_summary_structure
        path = self.root / "formula.json"
        vouchers = []
        for index, description in enumerate(("=1+1_张三", "=2+2_李四"), start=1):
            vouchers.append({
                "voucher_key": f"F{index}", "voucher_id": f"V{index}",
                "date": "2025-01-01", "period": "2025-01", "entity": "E1",
                "lines": [{
                    "voucher_id": f"V{index}", "sheet_name": "序时账",
                    "source_row": index + 1, "internal_line_key": f"3:序时账:{index + 1}",
                    "line_id": str(index), "account_code": "999001",
                    "account_name": "测试科目", "description": description,
                    "debit": "1.00", "credit": "0.00",
                }],
            })
        payload = {"schema_version": 1, "vouchers": vouchers}
        payload["summary_structure"] = infer_summary_structure(vouchers)
        path.write_text(json.dumps(payload), encoding="utf-8")
        request = request_for(
            path, group_by="account", output_mode="multi_sheet",
            columns=["date", "voucher_id", "description", "account_code"],
        )
        output = self.root / "formula-out"
        result = build_split_output(output, request)
        self.assertTrue(result["manifest"]["groups"][0]["derived_fields"])
        workbook = load_workbook(output / result["output_file"], read_only=True, data_only=False)
        try:
            sheet_name = next(
                name for name in workbook.sheetnames
                if name not in ("拆分说明", "分组索引")
            )
            headers = [cell.value for cell in next(workbook[sheet_name].iter_rows(max_row=1))]
            derived_index = headers.index("事项")
            values = [
                row[derived_index]
                for row in workbook[sheet_name].iter_rows(min_row=2, values_only=True)
            ]
            self.assertTrue(all(str(value).startswith("'=") for value in values))
        finally:
            workbook.close()
