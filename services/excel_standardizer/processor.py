#!/usr/bin/env python3
import argparse
import copy
import json
import sys
import traceback
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook  # noqa: E402

from services.excel_standardizer.audit import AuditLog  # noqa: E402
from services.excel_standardizer.job_state import read_job, update_job  # noqa: E402
from services.excel_standardizer.preview import build_preview  # noqa: E402
from services.excel_standardizer.rules import (  # noqa: E402
    autofit_columns,
    header_style,
    normalize_dates,
    remove_empty_rows,
    sheet_settings,
    trim_text,
)

FEATURE_LABELS = {
    "trim_text": "\u53bb\u9664\u6587\u672c\u9996\u5c3e\u7a7a\u683c",
    "normalize_dates": "\u7edf\u4e00\u65e5\u671f\u683c\u5f0f",
    "normalize_datetime": "\u7edf\u4e00\u65e5\u671f\u65f6\u95f4\u683c\u5f0f",
    "header_style": "\u7edf\u4e00\u8868\u5934\u6837\u5f0f",
    "autofit_columns": "\u81ea\u52a8\u8c03\u6574\u5217\u5bbd",
    "freeze_header": "\u51bb\u7ed3\u8868\u5934",
    "auto_filter": "\u6dfb\u52a0\u81ea\u52a8\u7b5b\u9009",
    "remove_empty_rows": "\u5220\u9664\u7a7a\u767d\u884c",
}

RECOMMENDED_SELECTION = {
    "trim_text": True,
    "normalize_dates": True,
    "normalize_datetime": False,
    "header_style": True,
    "autofit_columns": True,
    "freeze_header": True,
    "auto_filter": True,
    "remove_empty_rows": False,
}


def load_profile(name: str):
    safe = "".join(ch for ch in name if ch.isalnum() or ch in "_-") or "default"
    path = ROOT / "config" / "excel_profiles" / f"{safe}.json"
    if not path.exists():
        raise FileNotFoundError(f"Excel profile not found: {safe}")
    data = json.loads(path.read_text(encoding="utf-8"))
    return safe, data


def normalize_selection(raw):
    selection = dict(RECOMMENDED_SELECTION)
    if isinstance(raw, dict):
        for key in selection:
            if key in raw:
                selection[key] = raw[key] is True
    if not any(selection.values()):
        raise ValueError("At least one standardization feature must be selected.")
    return selection


def build_feature_plan(profile, selection):
    rules = profile.get("rules", {})
    plan = []

    if selection["trim_text"]:
        config = copy.deepcopy(rules.get("trim_text", {}))
        config["enabled"] = True
        plan.append(("trim_text", trim_text, config))

    if selection["normalize_dates"]:
        config = copy.deepcopy(rules.get("normalize_dates", {}))
        config.update({
            "enabled": True,
            "include_date": True,
            "include_datetime": False,
            "audit_rule": "normalize_dates",
            "label": FEATURE_LABELS["normalize_dates"],
        })
        plan.append(("normalize_dates", normalize_dates, config))

    if selection["normalize_datetime"]:
        config = copy.deepcopy(rules.get("normalize_dates", {}))
        config.update({
            "enabled": True,
            "include_date": False,
            "include_datetime": True,
            "audit_rule": "normalize_datetime",
            "label": FEATURE_LABELS["normalize_datetime"],
        })
        plan.append(("normalize_datetime", normalize_dates, config))

    if selection["header_style"]:
        config = copy.deepcopy(rules.get("header_style", {}))
        config["enabled"] = True
        plan.append(("header_style", header_style, config))

    if selection["autofit_columns"]:
        config = copy.deepcopy(rules.get("autofit_columns", {}))
        config["enabled"] = True
        plan.append(("autofit_columns", autofit_columns, config))

    if selection["freeze_header"]:
        config = copy.deepcopy(rules.get("sheet_settings", {}))
        config.update({
            "enabled": True,
            "freeze_header": True,
            "auto_filter": False,
            "freeze_audit_rule": "freeze_header",
            "label": FEATURE_LABELS["freeze_header"],
        })
        plan.append(("freeze_header", sheet_settings, config))

    if selection["auto_filter"]:
        config = copy.deepcopy(rules.get("sheet_settings", {}))
        config.update({
            "enabled": True,
            "freeze_header": False,
            "auto_filter": True,
            "filter_audit_rule": "auto_filter",
            "label": FEATURE_LABELS["auto_filter"],
        })
        plan.append(("auto_filter", sheet_settings, config))

    if selection["remove_empty_rows"]:
        config = copy.deepcopy(rules.get("remove_empty_rows", {}))
        config["enabled"] = True
        plan.append(("remove_empty_rows", remove_empty_rows, config))

    return plan


def workbook_summary(workbook):
    return {
        "sheetCount": len(workbook.worksheets),
        "sheets": [
            {"name": ws.title, "rows": ws.max_row, "columns": ws.max_column}
            for ws in workbook.worksheets
        ],
    }


def set_recalculation_flags(workbook):
    try:
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        workbook.calculation.calcMode = "auto"
    except Exception:
        pass


def process(job_dir: Path, profile_name: str):
    input_path = job_dir / "input.xlsx"
    output_path = job_dir / "output.xlsx"
    changes_path = job_dir / "changes.jsonl"
    preview_path = job_dir / "preview.json"

    profile_name, profile = load_profile(profile_name)
    job = read_job(job_dir)
    selection = normalize_selection(job.get("selection"))
    feature_plan = build_feature_plan(profile, selection)

    update_job(
        job_dir,
        status="validating",
        stage="validating",
        progress=5,
        message="\u6b63\u5728\u68c0\u67e5 Excel \u5de5\u4f5c\u7c3f\u2026",
        profile=profile_name,
        selection=selection,
        selectedFeatures=[name for name, _module, _config in feature_plan],
    )

    if not input_path.exists():
        raise FileNotFoundError("input.xlsx not found")

    workbook = load_workbook(input_path, data_only=False, read_only=False, keep_links=True)
    initial_summary = workbook_summary(workbook)

    update_job(
        job_dir,
        status="processing",
        stage="processing",
        progress=12,
        message=f"\u5df2\u8bfb\u53d6 {initial_summary['sheetCount']} \u4e2a\u5de5\u4f5c\u8868\uff0c\u5f00\u59cb\u6267\u884c {len(feature_plan)} \u9879\u5df2\u9009\u529f\u80fd\u2026",
        workbook=initial_summary,
    )

    audit = AuditLog(changes_path)
    try:
        total_features = max(1, len(feature_plan))
        for index, (name, module, config) in enumerate(feature_plan, start=1):
            start_progress = 15 + int((index - 1) / total_features * 58)
            update_job(
                job_dir,
                status="processing",
                stage=name,
                progress=start_progress,
                message=config.get("label") or FEATURE_LABELS.get(name) or f"\u6b63\u5728\u6267\u884c\uff1a{name}",
            )
            module.apply(workbook, audit, config)

        set_recalculation_flags(workbook)

        update_job(job_dir, status="checking", stage="saving", progress=78, message="\u6b63\u5728\u4fdd\u5b58\u6807\u51c6\u5316\u540e\u7684\u5de5\u4f5c\u7c3f\u2026")
        workbook.save(output_path)

        update_job(job_dir, status="checking", stage="checking", progress=86, message="\u6b63\u5728\u6821\u9a8c\u8f93\u51fa\u6587\u4ef6\u2026")
        check_book = load_workbook(output_path, data_only=False, read_only=True, keep_links=True)
        check_summary = workbook_summary(check_book)
        check_book.close()

        update_job(job_dir, status="reporting", stage="reporting", progress=92, message="\u6b63\u5728\u751f\u6210\u53d8\u66f4\u62a5\u544a\u548c\u7f51\u9875\u9884\u89c8\u2026")
        preview_config = profile.get("preview", {})
        build_preview(
            workbook,
            preview_path,
            max_rows=max(1, int(preview_config.get("rows", 50))),
            max_cols=max(1, int(preview_config.get("columns", 20))),
        )

        summary = audit.summary()
        summary.update({
            "sheetCount": check_summary["sheetCount"],
            "sheets": check_summary["sheets"],
            "inputBytes": input_path.stat().st_size,
            "outputBytes": output_path.stat().st_size,
            "rulesApplied": [name for name, _module, _config in feature_plan],
            "featuresApplied": [
                {"key": name, "label": FEATURE_LABELS[name], "changes": int(summary.get("byRule", {}).get(name, 0))}
                for name, _module, _config in feature_plan
            ],
            "selection": selection,
        })

        (job_dir / "result.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

        update_job(
            job_dir,
            status="completed",
            stage="completed",
            progress=100,
            message="Excel \u6807\u51c6\u5316\u5b8c\u6210\u3002",
            summary=summary,
            outputBytes=output_path.stat().st_size,
        )
    finally:
        audit.close()
        workbook.close()


def main():
    parser = argparse.ArgumentParser(description="Liuxian Lab Excel Standardizer")
    parser.add_argument("--job-dir", required=True)
    parser.add_argument("--profile", default="default")
    args = parser.parse_args()

    job_dir = Path(args.job_dir).resolve()
    try:
        process(job_dir, args.profile)
    except Exception as exc:
        error_text = f"{type(exc).__name__}: {exc}"
        traceback.print_exc()
        try:
            update_job(
                job_dir,
                status="failed",
                stage="failed",
                progress=100,
                message="Excel \u5904\u7406\u5931\u8d25\uff0c\u8bf7\u67e5\u770b\u9519\u8bef\u4fe1\u606f\u3002",
                error=error_text,
            )
        except Exception:
            pass
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
