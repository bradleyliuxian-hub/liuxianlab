"""Threaded-comment resolved/status tests (V1.3 status fix).

Status is a thread-level property read from OOXML ``done`` on the root
threaded comment and propagated to all replies.
"""

import json
import shutil
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from services.excel_review import parser, review_workbook
from services.excel_review.ooxml_threaded_comments import parse_done
from services.excel_review.status import Reporter
from services.excel_review.worker import run
from services.excel_review.tests.threaded_fixture import (
    build_base_workbook,
    inject_threaded_comments,
    new_guid,
)

FIXTURES = Path(__file__).resolve().parent / "fixtures"
STATUS_REAL = FIXTURES / "threaded_status_real.xlsx"


def _person():
    pid = new_guid()
    return pid, [{"id": pid, "displayName": "审阅人", "email": "reviewer@example.com"}]


def build_threaded(path, threads):
    pid, persons = _person()
    for thread in threads:
        thread.setdefault("personId", pid)
    workbook = build_base_workbook([("S", {"A1": 1, "B2": 2, "F4": 3})])
    workbook.save(path)
    inject_threaded_comments(str(path), "S", threads, persons)
    return pid


def thread(ref, tid, parent=None, done=None, text="t"):
    entry = {"ref": ref, "id": tid, "dT": "2026-10-13T06:00:00", "text": text}
    if parent:
        entry["parentId"] = parent
    if done is not None:
        entry["done"] = done
    return entry


class TestParseDone(unittest.TestCase):
    def test_true_tokens(self):
        for raw in ("1", "true", "TRUE", "True", "yes", "on"):
            self.assertTrue(parse_done(raw)[0], raw)

    def test_false_tokens(self):
        for raw in ("0", "false", "FALSE", "no", "off", "", None, "garbage"):
            self.assertFalse(parse_done(raw)[0], repr(raw))

    def test_raw_preserved(self):
        self.assertEqual(parse_done("1")[1], "1")


class TempBase(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def extract_items(self, threads):
        path = self.tmp / "s.xlsx"
        build_threaded(path, threads)
        return parser.extract(path)["items"]

    def run_worker(self, source_path, name="job"):
        job_dir = self.tmp / name
        job_dir.mkdir(parents=True, exist_ok=True)
        shutil.copy2(str(source_path), str(job_dir / "source.xlsx"))
        run(job_dir, Reporter(job_dir))
        return job_dir


class TestThreadStatus(TempBase):
    def test_root_missing_done_pending(self):
        root = new_guid()
        items = self.extract_items([thread("B2", root, text="no done")])
        self.assertEqual(items[0]["resolved"], False)
        self.assertEqual(items[0]["status"], "pending")

    def test_root_done_1_completed(self):
        root = new_guid()
        items = self.extract_items([thread("F4", root, done="1", text="resolved no reply")])
        self.assertEqual(items[0]["resolved"], True)
        self.assertEqual(items[0]["status"], "completed")

    def test_root_done_true_completed(self):
        root = new_guid()
        items = self.extract_items([thread("F4", root, done="true", text="resolved")])
        self.assertEqual(items[0]["status"], "completed")

    def test_root_done_0_pending(self):
        root = new_guid()
        items = self.extract_items([thread("F4", root, done="0", text="open")])
        self.assertEqual(items[0]["status"], "pending")

    def test_reply_inherits_pending_root(self):
        root, r1, r2 = new_guid(), new_guid(), new_guid()
        items = self.extract_items([
            thread("B2", root, text="root"),
            thread("B2", r1, parent=root, text="reply1"),
            thread("B2", r2, parent=root, text="reply2"),
        ])
        by_id = {i["threadId"]: i for i in items}
        self.assertEqual(by_id[root]["status"], "pending")
        self.assertEqual(by_id[r1]["status"], "pending")
        self.assertEqual(by_id[r2]["status"], "pending")

    def test_reply_inherits_completed_root(self):
        root, r1, r2 = new_guid(), new_guid(), new_guid()
        items = self.extract_items([
            thread("B2", root, done="1", text="root"),
            thread("B2", r1, parent=root, text="reply1"),
            thread("B2", r2, parent=root, text="reply2"),
        ])
        by_id = {i["threadId"]: i for i in items}
        self.assertEqual(by_id[root]["status"], "completed")
        self.assertEqual(by_id[r1]["status"], "completed")
        self.assertEqual(by_id[r2]["status"], "completed")
        self.assertEqual(by_id[r1]["rootThreadId"], root)

    def test_two_threads_do_not_contaminate(self):
        ta, tb = new_guid(), new_guid()
        items = self.extract_items([
            thread("B2", ta, text="open thread"),          # pending
            thread("F4", tb, done="1", text="closed thread"),  # completed
        ])
        by_id = {i["threadId"]: i for i in items}
        self.assertEqual(by_id[ta]["status"], "pending")
        self.assertEqual(by_id[tb]["status"], "completed")

    def test_orphan_reply_pending_no_crash(self):
        orphan = new_guid()
        items = self.extract_items([
            thread("B2", orphan, parent="{MISSING-PARENT}", text="orphan reply"),
        ])
        self.assertEqual(items[0]["status"], "pending")


class TestStatusInReview(TempBase):
    def test_status_column_and_extracted_json(self):
        root, r1 = new_guid(), new_guid()
        path = self.tmp / "s.xlsx"
        build_threaded(path, [
            thread("B2", root, done="1", text="root"),
            thread("B2", r1, parent=root, text="reply"),
        ])
        job_dir = self.run_worker(path)

        data = json.loads((job_dir / "extracted.json").read_text(encoding="utf-8"))
        statuses = {i["threadId"]: i["status"] for i in data["items"]}
        resolved = {i["threadId"]: i["resolved"] for i in data["items"]}
        self.assertEqual(statuses[root], "completed")
        self.assertEqual(statuses[r1], "completed")
        self.assertTrue(all(resolved.values()))

        sheet = load_workbook(job_dir / "review.xlsx")[review_workbook.THREADED_SHEET]
        header = [c.value for c in sheet[1]]
        status_col = header.index("状态") + 1
        for r in range(2, sheet.max_row + 1):
            self.assertEqual(sheet.cell(row=r, column=status_col).value, "completed")

    def test_resolved_thread_not_filtered_out(self):
        root = new_guid()
        path = self.tmp / "s.xlsx"
        build_threaded(path, [thread("F4", root, done="1", text="resolved only")])
        job_dir = self.run_worker(path)
        sheet = load_workbook(job_dir / "review.xlsx")[review_workbook.THREADED_SHEET]
        self.assertGreaterEqual(sheet.max_row, 2)  # still present


class TestStatusRealFixture(TempBase):
    def test_real_resolved_and_unresolved(self):
        items = parser.extract(STATUS_REAL)["items"]
        by_cell = {}
        for item in items:
            if item["commentType"] in ("threaded", "threaded_reply"):
                by_cell.setdefault(item["cellAddress"], []).append(item)

        # F4: root resolved, no replies -> completed
        self.assertTrue(by_cell["F4"])
        self.assertTrue(all(i["status"] == "completed" for i in by_cell["F4"]))
        self.assertEqual(by_cell["F4"][0]["commentType"], "threaded")

        # D4: root + 3 replies, none resolved -> all pending
        self.assertEqual(len(by_cell["D4"]), 4)
        self.assertTrue(all(i["status"] == "pending" for i in by_cell["D4"]))

        # B5: single unresolved root -> pending
        self.assertEqual(len(by_cell["B5"]), 1)
        self.assertEqual(by_cell["B5"][0]["status"], "pending")

    def test_legacy_unaffected(self):
        items = parser.extract(STATUS_REAL)["items"]
        legacy = [i for i in items if i["commentType"] == "legacy"]
        self.assertTrue(legacy)
        for item in legacy:
            self.assertEqual(item["status"], "pending")


class TestThreadSingleRow(TempBase):
    def test_thread_is_one_row_with_merged_replies(self):
        root, r1, r2 = new_guid(), new_guid(), new_guid()
        path = self.tmp / "s.xlsx"
        build_threaded(path, [
            thread("B2", root, text="问题一"),
            thread("B2", r1, parent=root, text="回复一"),
            thread("B2", r2, parent=root, text="回复二"),
        ])
        job_dir = self.run_worker(path)
        sheet = load_workbook(job_dir / "review.xlsx")[review_workbook.THREADED_SHEET]
        header = [c.value for c in sheet[1]]
        # One data row only (root), not three.
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet.cell(row=2, column=header.index("类型") + 1).value, "新版批注")
        self.assertEqual(sheet.cell(row=2, column=header.index("批注内容") + 1).value, "问题一")
        self.assertEqual(
            sheet.cell(row=2, column=header.index("回复") + 1).value,
            "回复一\n回复二",
        )

    def test_real_fixture_one_row_per_thread(self):
        job_dir = self.run_worker(STATUS_REAL)
        sheet = load_workbook(job_dir / "review.xlsx")[review_workbook.THREADED_SHEET]
        header = [c.value for c in sheet[1]]
        ci = header.index("Cell") + 1
        ri = header.index("回复") + 1
        si = header.index("状态") + 1
        rows = {}
        for r in range(2, sheet.max_row + 1):
            rows[sheet.cell(row=r, column=ci).value] = (
                sheet.cell(row=r, column=ri).value,
                sheet.cell(row=r, column=si).value,
            )
        # One row per thread: D4, F4, B5 => 3 rows total.
        self.assertEqual(sheet.max_row, 4)
        self.assertEqual(set(rows.keys()), {"D4", "F4", "B5"})
        # D4 has 3 replies merged into the 回复 cell.
        self.assertEqual(rows["D4"][0].count("\n"), 2)  # 3 lines
        self.assertEqual(rows["D4"][1], "pending")
        self.assertEqual(rows["F4"][1], "completed")


if __name__ == "__main__":
    unittest.main()
