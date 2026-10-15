"""Multi-file Control Log (V1.4) tests.

Run from the project root with a Python that has openpyxl:

    .venv/bin/python -m unittest discover -s services/excel_review/tests -t .
"""

import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from services.excel_review import control_log
from services.excel_review.control_log import (
    CONTROL_SHEET,
    DASHBOARD_SHEET,
    MAPPING_SHEET,
    SOURCE_FILES_SHEET,
)


def make_item(item_id, comment_type, sheet, cell, text, author, review_id="C001",
              thread_id="", parent="", resolved=False, source_file="file.xlsx"):
    return {
        "itemId": item_id,
        "reviewId": review_id,
        "commentType": comment_type,
        "sheetName": sheet,
        "cellAddress": cell,
        "cellValue": 10.5,
        "commentText": text,
        "author": author,
        "commentDatetime": "",
        "threadId": thread_id,
        "parentThreadId": parent,
        "rootThreadId": "",
        "personId": "",
        "resolved": resolved,
        "status": "completed" if resolved else "pending",
        "originalTextHash": "hash-" + item_id,
        "sourceReference": " → ".join([source_file, sheet, cell]),
        "excelAddress": "'%s'!%s" % (sheet.replace("'", "''"), cell),
    }


def make_job(job_id, source_file):
    return {
        "jobId": job_id,
        "sourceFile": source_file,
        "sourceSha256": "sha-" + job_id[:6],
        "createdAt": "2026-10-15T00:00:00Z",
    }


def two_job_fixture():
    jobs = [make_job("11111111-1111-1111-1111-111111111111", "Project_A.xlsx"),
            make_job("22222222-2222-2222-2222-222222222222", "Project_B.xlsx")]
    items_by_job = {
        jobs[0]["jobId"]: [
            make_item("a1", "legacy", "财务复核", "B5", "Allegro 检查这里", "Alice", source_file="Project_A.xlsx"),
            make_item("a2", "threaded", "财务复核", "C7", "阈值确认", "Bob", thread_id="t1", source_file="Project_A.xlsx"),
            make_item("a3", "threaded_reply", "财务复核", "C7", "已确认", "Carol",
                      thread_id="t2", parent="t1", source_file="Project_A.xlsx"),
        ],
        jobs[1]["jobId"]: [
            make_item("b1", "threaded", "利润表", "D10", "口径问题", "Alice", thread_id="u1", source_file="Project_B.xlsx"),
            make_item("b2", "legacy", "利润表", "E12", "四舍五入", "", source_file="Project_B.xlsx"),
        ],
    }
    return jobs, items_by_job


class SnapshotTests(unittest.TestCase):
    def build(self, jobs, items_by_job, name="October Review"):
        return control_log.build_snapshot("cl-1", name, jobs, items_by_job,
                                          created_at="2026-10-15T10:00:00Z")

    def test_question_counts_and_replies(self):
        jobs, items = two_job_fixture()
        snapshot = self.build(jobs, items)
        # 2 legacy + 2 threaded roots = 4 questions; 1 threaded reply excluded.
        self.assertEqual(snapshot["statistics"]["questions_total"], 4)
        self.assertEqual(snapshot["statistics"]["legacy_questions"], 2)
        self.assertEqual(snapshot["statistics"]["threaded_questions"], 2)
        self.assertEqual(snapshot["statistics"]["threaded_replies"], 1)
        self.assertEqual(len(snapshot["questions"]), 4)
        self.assertFalse(any(q["commentType"] == "threaded_reply" for q in snapshot["questions"]))

    def test_control_no_unique_and_source_identity(self):
        jobs, items = two_job_fixture()
        snapshot = self.build(jobs, items)
        numbers = [q["controlNo"] for q in snapshot["questions"]]
        self.assertEqual(numbers, ["CL-0001", "CL-0002", "CL-0003", "CL-0004"])
        self.assertEqual(len(set(numbers)), len(numbers))
        for question in snapshot["questions"]:
            self.assertTrue(question["sourceJobId"])
            self.assertTrue(question["sourceItemId"])
            self.assertTrue(question["controlItemId"])

    def test_same_review_id_across_files_does_not_collide(self):
        jobs, items = two_job_fixture()
        # Both jobs have a C001 review id by default.
        snapshot = self.build(jobs, items)
        ids = {(q["sourceJobId"], q["sourceItemId"]) for q in snapshot["questions"]}
        self.assertEqual(len(ids), 4)

    def test_duplicate_job_and_item_are_deduped(self):
        jobs, items = two_job_fixture()
        jobs_dup = jobs + [jobs[0]]  # duplicate job A
        items_dup = {
            jobs[0]["jobId"]: items[jobs[0]["jobId"]] + [items[jobs[0]["jobId"]][0]],
            jobs[1]["jobId"]: items[jobs[1]["jobId"]],
        }
        snapshot = self.build(jobs_dup, items_dup)
        self.assertEqual(len(snapshot["sourceJobs"]), 2)
        self.assertEqual(len(snapshot["questions"]), 4)

    def test_blank_author_normalised(self):
        jobs, items = two_job_fixture()
        snapshot = self.build(jobs, items)
        unknown = [q for q in snapshot["questions"] if q["author"] == "未知作者"]
        self.assertEqual(len(unknown), 1)
        self.assertEqual(unknown[0]["sourceFile"], "Project_B.xlsx")

    def test_single_quote_sheet_escaping(self):
        jobs = [make_job("33333333-3333-3333-3333-333333333333", "Q.xlsx")]
        items = {jobs[0]["jobId"]: [
            make_item("c1", "legacy", "Manager's Review", "C7", "check", "Alice"),
        ]}
        snapshot = self.build(jobs, items)
        self.assertEqual(snapshot["questions"][0]["excelAddress"], "'Manager''s Review'!C7")

    def test_source_file_and_reference(self):
        jobs, items = two_job_fixture()
        snapshot = self.build(jobs, items)
        files = {q["sourceFile"] for q in snapshot["questions"]}
        self.assertEqual(files, {"Project_A.xlsx", "Project_B.xlsx"})
        for question in snapshot["questions"]:
            self.assertIn(question["sourceFile"], question["sourceReference"])


class WorkbookTests(unittest.TestCase):
    def build_workbook(self, jobs, items_by_job, name="October Review"):
        snapshot = control_log.build_snapshot("cl-1", name, jobs, items_by_job,
                                              created_at="2026-10-15T10:00:00Z")
        tmp = Path(tempfile.mkdtemp())
        path = tmp / "control_log.xlsx"
        control_log.build_control_log_workbook(path, snapshot)
        return snapshot, path, tmp

    def test_sheets_and_mapping_hidden(self):
        jobs, items = two_job_fixture()
        _, path, _ = self.build_workbook(jobs, items)
        workbook = load_workbook(path)
        self.assertEqual(
            workbook.sheetnames,
            [DASHBOARD_SHEET, CONTROL_SHEET, SOURCE_FILES_SHEET, MAPPING_SHEET],
        )
        self.assertEqual(workbook[MAPPING_SHEET].sheet_state, "hidden")
        for name in (DASHBOARD_SHEET, CONTROL_SHEET, SOURCE_FILES_SHEET):
            self.assertEqual(workbook[name].sheet_state, "visible")
        workbook.close()

    def test_visible_sheets_have_no_protection(self):
        jobs, items = two_job_fixture()
        _, path, _ = self.build_workbook(jobs, items)
        workbook = load_workbook(path)
        for name in (DASHBOARD_SHEET, CONTROL_SHEET, SOURCE_FILES_SHEET):
            self.assertFalse(bool(workbook[name].protection.sheet), name)
        workbook.close()

    def test_control_sheet_columns_and_rows(self):
        jobs, items = two_job_fixture()
        snapshot, path, _ = self.build_workbook(jobs, items)
        workbook = load_workbook(path)
        sheet = workbook[CONTROL_SHEET]
        header = [sheet.cell(row=1, column=c).value for c in range(1, 17)]
        self.assertEqual(header, control_log.CONTROL_HEADERS)
        # 4 questions -> rows 2..5
        self.assertEqual(sheet.max_row, 1 + len(snapshot["questions"]))
        self.assertEqual(sheet["A2"].value, "CL-0001")
        self.assertEqual(sheet["K2"].value, "Open")  # Status default
        self.assertIn(sheet["J2"].value, (None, ""))  # Reply default blank
        workbook.close()

    def test_source_files_sheet(self):
        jobs, items = two_job_fixture()
        _, path, _ = self.build_workbook(jobs, items)
        workbook = load_workbook(path)
        sheet = workbook[SOURCE_FILES_SHEET]
        self.assertEqual(sheet["A2"].value, 1)
        self.assertEqual(sheet["B2"].value, "Project_A.xlsx")
        self.assertEqual(sheet["D2"].value, 2)  # questions in file A (reply excluded)
        self.assertEqual(sheet["B3"].value, "Project_B.xlsx")
        workbook.close()

    def test_mapping_identity_preserved(self):
        jobs, items = two_job_fixture()
        snapshot, path, _ = self.build_workbook(jobs, items)
        workbook = load_workbook(path)
        sheet = workbook[MAPPING_SHEET]
        headers = [sheet.cell(row=1, column=c).value for c in range(1, len(control_log.MAPPING_HEADERS) + 1)]
        self.assertEqual(headers, control_log.MAPPING_HEADERS)
        col = {name: i + 1 for i, name in enumerate(control_log.MAPPING_HEADERS)}
        row_job = sheet.cell(row=2, column=col["source_job_id"]).value
        row_item = sheet.cell(row=2, column=col["source_item_id"]).value
        self.assertEqual(row_job, snapshot["questions"][0]["sourceJobId"])
        self.assertEqual(row_item, snapshot["questions"][0]["sourceItemId"])
        self.assertEqual(sheet.cell(row=2, column=col["control_log_id"]).value, "cl-1")
        workbook.close()

    def test_five_jobs(self):
        jobs = [make_job("%08d-0000-0000-0000-000000000000" % i, "F%d.xlsx" % i) for i in range(5)]
        items = {}
        for index, job in enumerate(jobs):
            items[job["jobId"]] = [
                make_item("i%d" % index, "legacy", "S", "A1", "q", "Alice"),
                make_item("t%d" % index, "threaded", "S", "A2", "q", "Bob", thread_id="t%d" % index),
            ]
        snapshot, path, _ = self.build_workbook(jobs, items)
        self.assertEqual(len(snapshot["sourceJobs"]), 5)
        self.assertEqual(snapshot["statistics"]["questions_total"], 10)
        workbook = load_workbook(path)
        self.assertEqual(workbook[SOURCE_FILES_SHEET].max_row, 6)
        workbook.close()

    def test_empty_questions_still_builds(self):
        jobs = [make_job("44444444-4444-4444-4444-444444444444", "Empty.xlsx")]
        items = {jobs[0]["jobId"]: []}
        snapshot, path, _ = self.build_workbook(jobs, items)
        self.assertEqual(snapshot["statistics"]["questions_total"], 0)
        workbook = load_workbook(path)
        self.assertEqual(workbook[CONTROL_SHEET].max_row, 1)  # header only
        workbook.close()


if __name__ == "__main__":
    unittest.main()
