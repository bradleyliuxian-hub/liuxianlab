import json
import tempfile
import unittest
import zipfile
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from services.journal_entry.batch_output import (
    BatchOutputError,
    build_batch_output,
)
from services.journal_entry.export import build_export, safe_excel_text


def result_payload(*, filename, voucher_id="001", internal_key="3:序时账:2", amount="100.00"):
    return {
        "schema_version": 1,
        "generated_at": "2026-01-01T00:00:00+00:00",
        "input_metadata": {"source_filename": filename, "sheet_name": "序时账", "header_row": 1},
        "cleaning_summary": {
            "raw_row_count": 1,
            "line_count": 1,
            "issue_count": 1,
            "unbalanced_voucher_count": 0,
            "voucher_count": 1,
        },
        "business_classification_summary": {
            "voucher_count": 1,
            "classified_count": 1,
            "business_counts": {"JE-B006": 1},
            "business_labels": {"JE-B006": "销售收入"},
        },
        "anomaly_summary": {
            "finding_count": 1,
            "finding_voucher_count": 1,
            "code_counts": {"JE-A001": 1},
            "finding_severity_counts": {"LOW": 1},
        },
        "findings": [
            {
                "code": "JE-A001", "severity": "LOW", "voucher_key": f"{filename}|K",
                "voucher_id": voucher_id, "risk_score": "10.00",
            }
        ],
        "validation_issues": [{"code": "OPTIONAL_AMOUNT_INVALID"}],
        "row_warnings": [],
        "vouchers": [
            {
                "voucher_key": f"{filename}|K",
                "voucher_id": voucher_id,
                "date": "2025-01-01",
                "period": "2025-01",
                "entity": "E",
                "total_debit": amount,
                "total_credit": amount,
                "exactly_balanced": True,
                "lines": [
                    {
                        "voucher_id": voucher_id,
                        "sheet_name": "序时账",
                        "source_row": 2,
                        "line_id": "001",
                        "internal_line_key": internal_key,
                        "account_code": "1002",
                        "account_name": "银行存款",
                        "detail_account": "",
                        "description": "=HYPERLINK(\"x\")",
                        "debit": amount,
                        "credit": "0.00",
                        "currency": "CNY",
                        "original_amount": amount,
                        "preparer": "tester",
                    }
                ],
            }
        ],
    }


class BatchOutputTests(unittest.TestCase):
    def test_summary_merge_zip_and_snapshot(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            output = root / "output"
            output.mkdir()
            jobs = []
            for index in (1, 2):
                job_dir = root / f"job-{index}"
                job_dir.mkdir()
                result = result_payload(
                    filename="journal.xlsx",
                    voucher_id="001",
                    internal_key="3:序时账:2",
                    amount="100.00",
                )
                result_path = job_dir / "result.json"
                result_path.write_text(json.dumps(result), encoding="utf-8")
                export_path = job_dir / "journal-entry-analysis-report.xlsx"
                if index == 2:
                    # Force the generator to rebuild the per-file report.
                    pass
                else:
                    export_path.write_bytes(b"fake-xlsx")
                jobs.append({
                    "position": index - 1,
                    "job_id": f"job-{index}",
                    "original_filename": "journal.xlsx",
                    "status": "completed",
                    "result_path": str(result_path),
                    "export_path": str(export_path),
                    "summary": {
                        "raw_row_count": 1, "line_count": 1, "voucher_count": 10,
                        "issue_count": 1, "unbalanced_voucher_count": 0,
                        "business_classified_count": 5, "business_voucher_count": 10,
                        "finding_count": 1, "finding_voucher_count": 1,
                    },
                    "amount_layout": "debit_credit",
                    "positive_side": None,
                    "sheet_name": "序时账",
                    "header_row": 1,
                    "completed_at": "2026-01-01T00:00:00+00:00",
                })
            request = {
                "schema_version": 1,
                "batch": {
                    "id": "batch-1", "name": "batch", "generated_at": "2026-01-02T00:00:00+00:00",
                    "total_files": 3, "included_files": 2, "excluded_files": 1,
                    "excluded_draft": 1, "excluded_active": 0, "excluded_failed": 0,
                    "summary": {
                        "totalFileCount": 3, "completedFileCount": 2,
                        "rawRowCount": 2, "lineCount": 2, "voucherCount": 20,
                        "issueCount": 2, "unbalancedVoucherCount": 0,
                        "businessClassifiedCount": 10, "businessVoucherCount": 20,
                        "findingCount": 2, "findingVoucherCount": 2,
                        "businessCoverage": "0.500000", "reviewRate": "0.100000",
                    },
                },
                "included": jobs,
                "excluded": [{
                    "position": 2, "job_id": "job-3", "original_filename": "draft.xlsx",
                    "status": "draft", "reason": "未完成配置。"
                }],
            }
            outcome = build_batch_output(output, request)
            self.assertEqual(outcome["merged_rows"], 2)
            self.assertTrue((output / "batch-summary.xlsx").exists())
            self.assertTrue((output / "merged-standardized-journal.xlsx").exists())
            self.assertTrue((output / "manifest.json").exists())
            self.assertTrue((output / "batch-report.zip").exists())
            self.assertTrue((root / "job-2" / "journal-entry-analysis-report.xlsx").exists())

            workbook = load_workbook(output / "batch-summary.xlsx", read_only=True, data_only=False)
            try:
                self.assertEqual(
                    workbook.sheetnames,
                    ["批次概览", "文件清单", "业务分类汇总", "异常汇总", "数据质量汇总"],
                )
                overview = {row[0].value: row[1].value for row in workbook["批次概览"].iter_rows(min_row=2)}
                self.assertEqual(overview["Batch ID"], "batch-1")
                self.assertAlmostEqual(float(overview["业务分类覆盖率"]), 0.5)
                self.assertAlmostEqual(float(overview["Review Rate"]), 0.1)
                files = list(workbook["文件清单"].iter_rows(min_row=2, values_only=True))
                self.assertEqual(len(files), 3)
                excluded = [row for row in files if row[1] == "draft.xlsx"][0]
                self.assertEqual(excluded[4], "否")
                self.assertEqual(excluded[7], "—")
            finally:
                workbook.close()

            merged = load_workbook(output / "merged-standardized-journal.xlsx", read_only=True, data_only=False)
            try:
                self.assertEqual(merged.sheetnames, ["标准分录", "来源文件"])
                all_rows = list(merged["标准分录"].iter_rows(values_only=True))
                headers = list(all_rows[0])
                rows = all_rows[1:]
                self.assertEqual(len(rows), 2)
                self.assertEqual(headers[:3], ["日期", "凭证号", "摘要"])
                self.assertEqual(headers[-1], "批次ID")
                idx = {name: headers.index(name) for name in headers}
                self.assertEqual({row[idx["批次ID"]] for row in rows}, {"batch-1"})
                self.assertEqual({row[idx["任务ID"]] for row in rows}, {"job-1", "job-2"})
                self.assertEqual({row[idx["原始文件"]] for row in rows}, {"journal.xlsx"})
                self.assertEqual({row[idx["凭证内部标识"]] for row in rows}, {"journal.xlsx|K"})
                self.assertEqual({row[idx["内部行标识"]] for row in rows}, {"3:序时账:2"})
                self.assertTrue(all(str(row[idx["摘要"]]).startswith("'=") for row in rows))
                self.assertTrue(all(Decimal(str(row[idx["借方金额"]])) == Decimal("100.00") for row in rows))
            finally:
                merged.close()

            with zipfile.ZipFile(output / "batch-report.zip") as archive:
                names = archive.namelist()
                self.assertIn("batch-summary.xlsx", names)
                self.assertIn("merged-standardized-journal.xlsx", names)
                self.assertIn("manifest.json", names)
                report_names = [name for name in names if name.startswith("reports/")]
                self.assertEqual(len(report_names), 2)
                self.assertEqual(len(set(report_names)), 2)
                self.assertFalse(any(name.startswith("/") or ".." in name.split("/") for name in names))

            manifest = json.loads((output / "manifest.json").read_text(encoding="utf-8"))
            self.assertEqual(manifest["included_files"], 2)
            self.assertEqual(manifest["excluded_files"], 1)
            self.assertEqual(len(manifest["items"]), 3)
            self.assertEqual(manifest["result_schema_version"], "1")
            self.assertNotIn("debit", json.dumps(manifest).lower())
            self.assertNotIn("100.00", json.dumps(manifest))


    def test_zip_entry_names_are_safe_for_malicious_filenames(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            output = root / "output"
            output.mkdir()
            names = ["../../evil.xlsx", "..\\evil.xlsx", "/absolute.xlsx", "a/b.xlsx"]
            included = []
            for index, name in enumerate(names):
                job_dir = root / f"job-{index}"
                job_dir.mkdir()
                result = result_payload(filename=name)
                result_path = job_dir / "result.json"
                result_path.write_text(json.dumps(result), encoding="utf-8")
                export_path = job_dir / "report.xlsx"
                export_path.write_bytes(b"fake")
                included.append({
                    "position": index, "job_id": f"job-{index}", "original_filename": name,
                    "status": "completed", "result_path": str(result_path),
                    "export_path": str(export_path), "summary": {}, "amount_layout": "debit_credit",
                })
            request = {
                "batch": {
                    "id": "b", "name": "b", "generated_at": "2026-01-01T00:00:00+00:00",
                    "total_files": 4, "included_files": 4, "excluded_files": 0,
                    "summary": {"totalFileCount": 4, "completedFileCount": 4},
                },
                "included": included, "excluded": [],
            }
            build_batch_output(output, request)
            with zipfile.ZipFile(output / "batch-report.zip") as archive:
                entries = archive.namelist()
            reports = [name for name in entries if name.startswith("reports/")]
            self.assertEqual(len(reports), 4)
            self.assertEqual(len(set(reports)), 4)
            for name in entries:
                self.assertFalse(name.startswith("/"))
                self.assertFalse(".." in name.split("/"))
                self.assertNotIn("\\", name)
            self.assertTrue(all(name.startswith("reports/") for name in reports))


    def test_per_file_export_also_blocks_formula_injection(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "per-file.xlsx"
            payload = {
                "vouchers": [{
                    "voucher_key": "K", "voucher_id": "V", "date": "2025-01-01",
                    "period": "2025-01", "entity": "E", "total_debit": "1.00",
                    "total_credit": "1.00", "imbalance": "0.00",
                    "exactly_balanced": True, "is_within_tolerance": True,
                    "line_count": 1, "validation_flags": [],
                    "lines": [{
                        "voucher_key": "K", "voucher_id": "V", "date": "2025-01-01",
                        "period": "2025-01", "entity": "E", "account_code": "1002",
                        "account_name": "=ACCOUNT", "detail_account": "",
                        "debit": "1.00", "credit": "0.00", "description": "@cmd",
                        "currency": "CNY", "original_amount": "1.00",
                        "sheet_name": "序时账", "source_row": 2,
                        "internal_line_key": "3:序时账:2", "line_id": "001",
                    }],
                }],
                "findings": [],
                "business_classification_summary": {},
                "anomaly_summary": {},
                "cleaning_summary": {},
                "input_metadata": {},
            }
            build_export(path, payload)
            workbook = load_workbook(path, read_only=True, data_only=False)
            try:
                rows = list(workbook["清洗分录"].iter_rows(values_only=True))
                headers = list(rows[0])
                row = rows[1]
                self.assertTrue(str(row[headers.index("科目名称")]).startswith("'="))
                self.assertTrue(str(row[headers.index("摘要")]).startswith("'@"))
            finally:
                workbook.close()

    def test_formula_injection_helper(self):
        self.assertEqual(safe_excel_text("=1+1"), "'=1+1")
        self.assertEqual(safe_excel_text("normal"), "normal")

    def test_corrupt_result_fails_explicitly(self):
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "output"
            output.mkdir()
            job_dir = Path(tmp) / "job"
            job_dir.mkdir()
            (job_dir / "result.json").write_text("{bad json", encoding="utf-8")
            request = {
                "batch": {"id": "b", "name": "b", "summary": {}},
                "included": [{
                    "position": 0, "job_id": "j", "original_filename": "x.xlsx",
                    "result_path": str(job_dir / "result.json"),
                    "export_path": str(job_dir / "report.xlsx"),
                }],
                "excluded": [],
            }
            with self.assertRaises(BatchOutputError):
                build_batch_output(output, request)


if __name__ == "__main__":
    unittest.main()
