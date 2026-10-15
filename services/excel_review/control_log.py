"""Multi-file Review Control Log (Excel Review Assistant V1.4).

A Control Log is a *snapshot aggregation* of several single-file review
jobs. It never re-parses the source workbooks: it consumes the already
extracted, normalized items of each job and produces

- ``control_log.json``  - the authoritative, self-contained snapshot
- ``control_log.xlsx``  - Dashboard + Control Log + Source Files + _Mapping

Question semantics match V1.3 exactly (see :mod:`services.excel_review.statistics`):
a question is a legacy note OR a Microsoft 365 threaded root; a threaded
reply is not a question. Author = comment author (blank -> 未知作者).

Identity: a Control Log question is identified by
``control_log_id + control_item_id``, and traced back through
``source_job_id + source_item_id``. Display numbers (``Control No.``) are
CL-0001, CL-0002, ... and are for humans only.
"""

import json
import uuid
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from .parser import (
    COMMENT_TYPE_LEGACY,
    COMMENT_TYPE_THREADED,
    format_excel_address,
)
from .review_workbook import set_literal_cell, verify_openable
from .statistics import build_statistics, normalize_author

CONTROL_MAPPING_VERSION = "excel-review-control-log-v1"

DASHBOARD_SHEET = "Dashboard"
CONTROL_SHEET = "Control Log"
SOURCE_FILES_SHEET = "Source Files"
MAPPING_SHEET = "_Mapping"
CONTROL_ITEM_ID_COLUMN = "_CONTROL_ITEM_ID"

TABLE_CONTROL = "ControlLogTable"
TABLE_SOURCES = "SourceFilesTable"

QUESTION_TYPES = (COMMENT_TYPE_LEGACY, COMMENT_TYPE_THREADED)

TYPE_LABELS = {
    COMMENT_TYPE_LEGACY: "旧版批注",
    COMMENT_TYPE_THREADED: "新版批注",
}

DEFAULT_STATUS = "Open"

# Control Log sheet columns. Source Reference stays last visible; the hidden
# identity column comes after it.
CONTROL_HEADERS = [
    "Control No.",
    "Source File",
    "Sheet",
    "Cell",
    "Excel Address",
    "Type",
    "Original Value",
    "Review Issue",
    "Author",
    "Reply",
    "Status",
    "Owner",
    "Priority",
    "Due Date",
    "Notes",
    "Source Reference",
]
_CONTROL_WRAP_COLUMNS = (8, 16)  # Review Issue / Source Reference

SOURCE_HEADERS = [
    "#",
    "Source File",
    "Source Job ID",
    "Questions",
    "Threaded",
    "Legacy",
    "Created At",
]

MAPPING_HEADERS = [
    "mapping_version",
    "control_log_id",
    "control_item_id",
    "control_no",
    "source_job_id",
    "source_item_id",
    "source_filename",
    "source_sha256",
    "comment_type",
    "review_id",
    "sheet_name",
    "cell_address",
    "excel_address",
    "source_reference",
    "original_value",
    "original_text_hash",
    "author",
    "comment_datetime",
    "thread_id",
    "parent_thread_id",
    "person_id",
    "resolved",
]


def _now_iso():
    return datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


def _question_from_item(job, item, control_no):
    """Build one Control Log question snapshot record."""
    return {
        "controlItemId": uuid.uuid4().hex,
        "controlNo": control_no,
        "sourceJobId": job["jobId"],
        "sourceItemId": item.get("itemId", ""),
        "sourceFile": job.get("sourceFile") or item.get("sourceFilename") or "source.xlsx",
        "commentType": item.get("commentType", ""),
        "reviewId": item.get("reviewId", ""),
        "sheetName": item.get("sheetName", ""),
        "cellAddress": item.get("cellAddress", ""),
        "excelAddress": item.get("excelAddress")
        or format_excel_address(item.get("sheetName", ""), item.get("cellAddress", "")),
        "sourceReference": item.get("sourceReference", ""),
        "originalValue": item.get("cellValue", ""),
        "commentText": item.get("commentText", ""),
        "author": normalize_author(item.get("author")),
        "commentDatetime": item.get("commentDatetime", ""),
        "threadId": item.get("threadId", ""),
        "parentThreadId": item.get("parentThreadId", ""),
        "personId": item.get("personId", ""),
        "originalTextHash": item.get("originalTextHash", ""),
        "resolved": bool(item.get("resolved", False)),
    }


def build_questions(jobs, items_by_job):
    """Flatten every job's questions into ordered, de-duplicated records.

    Defensive dedup: (source_job_id, source_item_id) appears at most once,
    and a job id is only processed once even if submitted twice.
    """
    questions = []
    seen_items = set()
    seen_jobs = set()
    counter = 0
    for job in jobs:
        job_id = str(job.get("jobId", ""))
        if not job_id or job_id in seen_jobs:
            continue
        seen_jobs.add(job_id)
        for item in items_by_job.get(job_id, []):
            if item.get("commentType") not in QUESTION_TYPES:
                continue  # threaded replies are not questions
            key = (job_id, str(item.get("itemId", "")))
            if key in seen_items:
                continue
            seen_items.add(key)
            counter += 1
            questions.append(_question_from_item(job, item, "CL-%04d" % counter))
    return questions


def _file_counts(job, questions):
    threaded = sum(
        1 for q in questions
        if q["sourceJobId"] == job["jobId"] and q["commentType"] == COMMENT_TYPE_THREADED
    )
    legacy = sum(
        1 for q in questions
        if q["sourceJobId"] == job["jobId"] and q["commentType"] == COMMENT_TYPE_LEGACY
    )
    return threaded, legacy


def build_snapshot(control_log_id, name, jobs, items_by_job, created_at=None):
    """Build the authoritative, self-contained control_log.json payload."""
    # Defensive dedup: a job id is only counted once even if submitted twice.
    deduped_jobs = []
    seen_jobs = set()
    for job in jobs or []:
        job_id = str(job.get("jobId", ""))
        if not job_id or job_id in seen_jobs:
            continue
        seen_jobs.add(job_id)
        deduped_jobs.append(job)
    jobs = deduped_jobs

    questions = build_questions(jobs, items_by_job)

    all_items = []
    for job in jobs:
        all_items.extend(items_by_job.get(str(job.get("jobId", "")), []))
    statistics = build_statistics(all_items)

    source_jobs = []
    for job in jobs:
        items = items_by_job.get(str(job.get("jobId", "")), [])
        threaded, legacy = _file_counts(job, questions)
        source_jobs.append(
            {
                "jobId": job.get("jobId", ""),
                "sourceFile": job.get("sourceFile") or job.get("sourceFilename") or "source.xlsx",
                "sourceSha256": job.get("sourceSha256", ""),
                "createdAt": job.get("createdAt", ""),
                "questionCount": threaded + legacy,
                "threadedCount": threaded,
                "legacyCount": legacy,
                "statistics": build_statistics(items),
            }
        )

    return {
        "mappingVersion": CONTROL_MAPPING_VERSION,
        "controlLogId": control_log_id,
        "name": name,
        "createdAt": created_at or _now_iso(),
        "sourceJobs": source_jobs,
        "statistics": statistics,
        "questions": questions,
    }


def write_snapshot(path, snapshot):
    with open(path, "w", encoding="utf-8") as handle:
        json.dump(snapshot, handle, ensure_ascii=False, indent=2)
        handle.write("\n")


def _add_table(worksheet, name, header_count, data_rows):
    if data_rows <= 0:
        return False
    ref = "A1:%s%d" % (get_column_letter(header_count), data_rows + 1)
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False
    )
    worksheet.add_table(table)
    return True


def _write_header(worksheet, headers):
    for column, name in enumerate(headers, start=1):
        cell = worksheet.cell(row=1, column=column, value=name)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E4F3CB")
        cell.protection = Protection(locked=False)


def _write_control_sheet(worksheet, questions):
    _write_header(worksheet, CONTROL_HEADERS)
    id_column = len(CONTROL_HEADERS) + 1
    id_header = worksheet.cell(row=1, column=id_column, value=CONTROL_ITEM_ID_COLUMN)
    id_header.protection = Protection(locked=False)
    worksheet.column_dimensions[get_column_letter(id_column)].hidden = True
    worksheet.freeze_panes = "A2"

    for offset, question in enumerate(questions):
        row = offset + 2
        values = [
            question["controlNo"],
            question["sourceFile"],
            question["sheetName"],
            question["cellAddress"],
            question["excelAddress"],
            TYPE_LABELS.get(question["commentType"], question["commentType"]),
            question["originalValue"],
            question["commentText"],
            question["author"],
            "",              # Reply (user-managed)
            DEFAULT_STATUS,  # Status
            "",              # Owner
            "",              # Priority
            "",              # Due Date
            "",              # Notes
            question["sourceReference"],
        ]
        for column, value in enumerate(values, start=1):
            cell = set_literal_cell(worksheet, row, column, value)
            cell.protection = Protection(locked=False)
            if column in _CONTROL_WRAP_COLUMNS:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        id_cell = set_literal_cell(worksheet, row, id_column, question["controlItemId"])
        id_cell.protection = Protection(locked=False)

    widths = {"A": 12, "B": 26, "C": 18, "D": 10, "E": 28, "F": 12, "G": 20,
              "H": 52, "I": 14, "J": 26, "K": 10, "L": 12, "M": 10, "N": 12,
              "O": 26, "P": 40}
    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width


def _write_source_sheet(worksheet, source_jobs):
    _write_header(worksheet, SOURCE_HEADERS)
    worksheet.freeze_panes = "A2"
    for offset, job in enumerate(source_jobs):
        row = offset + 2
        values = [
            offset + 1,
            job["sourceFile"],
            job["jobId"],
            job["questionCount"],
            job["threadedCount"],
            job["legacyCount"],
            job["createdAt"],
        ]
        for column, value in enumerate(values, start=1):
            cell = set_literal_cell(worksheet, row, column, value)
            cell.protection = Protection(locked=False)
    widths = {"A": 6, "B": 30, "C": 38, "D": 12, "E": 12, "F": 12, "G": 22}
    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width


def _build_dashboard(worksheet, snapshot):
    stats = snapshot["statistics"]
    source_jobs = snapshot["sourceJobs"]
    authored = stats.get("authors", [])

    bold = Font(bold=True)
    worksheet["A1"] = "Excel Review Control Log"
    worksheet["A1"].font = Font(bold=True, size=14)
    worksheet["A2"] = "Name"
    worksheet["A2"].font = bold
    set_literal_cell(worksheet, 2, 2, snapshot["name"])
    worksheet["A3"] = "Generated At"
    worksheet["A3"].font = bold
    set_literal_cell(worksheet, 3, 2, snapshot["createdAt"])

    worksheet["A5"] = "Files"
    worksheet["A5"].font = bold
    worksheet["B5"] = "=IFERROR(COUNTA(%s[Source File]),0)" % TABLE_SOURCES
    worksheet["A6"] = "Total Questions"
    worksheet["A6"].font = bold
    worksheet["B6"] = "=IFERROR(COUNTA(%s[Control No.]),0)" % TABLE_CONTROL
    worksheet["A7"] = "Threaded Questions"
    worksheet["A7"].font = bold
    worksheet["B7"] = '=IFERROR(COUNTIF(%s[Type],"%s"),0)' % (
        TABLE_CONTROL, TYPE_LABELS[COMMENT_TYPE_THREADED])
    worksheet["A8"] = "Legacy Questions"
    worksheet["A8"].font = bold
    worksheet["B8"] = '=IFERROR(COUNTIF(%s[Type],"%s"),0)' % (
        TABLE_CONTROL, TYPE_LABELS[COMMENT_TYPE_LEGACY])

    worksheet["A10"] = "Answered"
    worksheet["A10"].font = bold
    worksheet["B10"] = '=IFERROR(COUNTIF(%s[Reply],"<>"),0)' % TABLE_CONTROL
    worksheet["A11"] = "Unanswered"
    worksheet["A11"].font = bold
    worksheet["B11"] = "=B6-B10"

    worksheet["A13"] = "Open"
    worksheet["A13"].font = bold
    worksheet["B13"] = '=IFERROR(COUNTIF(%s[Status],"Open"),0)' % TABLE_CONTROL
    worksheet["A14"] = "In Progress"
    worksheet["A14"].font = bold
    worksheet["B14"] = '=IFERROR(COUNTIF(%s[Status],"In Progress"),0)' % TABLE_CONTROL
    worksheet["A15"] = "Closed"
    worksheet["A15"].font = bold
    worksheet["B15"] = '=IFERROR(COUNTIF(%s[Status],"Closed"),0)' % TABLE_CONTROL

    # Questions by Source File.
    row = 17
    worksheet.cell(row=row, column=1, value="Questions by Source File").font = bold
    row += 1
    for col, name in enumerate(["Source File", "Total", "Threaded", "Legacy"], start=1):
        worksheet.cell(row=row, column=col, value=name).font = bold
    for job in source_jobs:
        row += 1
        set_literal_cell(worksheet, row, 1, job["sourceFile"])
        worksheet.cell(row=row, column=2,
                       value='=IFERROR(COUNTIF(%s[Source File],$A%d),0)' % (TABLE_CONTROL, row))
        worksheet.cell(row=row, column=3,
                       value='=IFERROR(COUNTIFS(%s[Source File],$A%d,%s[Type],"%s"),0)' % (
                           TABLE_CONTROL, row, TABLE_CONTROL, TYPE_LABELS[COMMENT_TYPE_THREADED]))
        worksheet.cell(row=row, column=4,
                       value='=IFERROR(COUNTIFS(%s[Source File],$A%d,%s[Type],"%s"),0)' % (
                           TABLE_CONTROL, row, TABLE_CONTROL, TYPE_LABELS[COMMENT_TYPE_LEGACY]))

    # Questions by Author.
    row += 2
    worksheet.cell(row=row, column=1, value="Questions by Author").font = bold
    row += 1
    for col, name in enumerate(["Author", "Total"], start=1):
        worksheet.cell(row=row, column=col, value=name).font = bold
    for entry in authored:
        row += 1
        set_literal_cell(worksheet, row, 1, entry["author"])
        worksheet.cell(row=row, column=2,
                       value='=IFERROR(COUNTIF(%s[Author],$A%d),0)' % (TABLE_CONTROL, row))

    worksheet.column_dimensions["A"].width = 34
    worksheet.column_dimensions["B"].width = 18
    worksheet.column_dimensions["C"].width = 12
    worksheet.column_dimensions["D"].width = 12


def _write_mapping(worksheet, snapshot):
    mapping_version = snapshot["mappingVersion"]
    control_log_id = snapshot["controlLogId"]
    sha_by_job = {
        str(job["jobId"]): job.get("sourceSha256", "")
        for job in snapshot["sourceJobs"]
    }
    source_by_job = {
        str(job["jobId"]): job.get("sourceFile", "")
        for job in snapshot["sourceJobs"]
    }
    for column, name in enumerate(MAPPING_HEADERS, start=1):
        cell = worksheet.cell(row=1, column=column, value=name)
        cell.protection = Protection(locked=False)
    for row_index, question in enumerate(snapshot["questions"], start=2):
        job_id = str(question["sourceJobId"])
        values = {
            "mapping_version": mapping_version,
            "control_log_id": control_log_id,
            "control_item_id": question["controlItemId"],
            "control_no": question["controlNo"],
            "source_job_id": job_id,
            "source_item_id": question["sourceItemId"],
            "source_filename": source_by_job.get(job_id, question["sourceFile"]),
            "source_sha256": sha_by_job.get(job_id, ""),
            "comment_type": question["commentType"],
            "review_id": question["reviewId"],
            "sheet_name": question["sheetName"],
            "cell_address": question["cellAddress"],
            "excel_address": question["excelAddress"],
            "source_reference": question["sourceReference"],
            "original_value": question["originalValue"],
            "original_text_hash": question["originalTextHash"],
            "author": question["author"],
            "comment_datetime": question["commentDatetime"],
            "thread_id": question["threadId"],
            "parent_thread_id": question["parentThreadId"],
            "person_id": question["personId"],
            "resolved": question["resolved"],
        }
        for column, name in enumerate(MAPPING_HEADERS, start=1):
            cell = set_literal_cell(worksheet, row_index, column, values[name])
            cell.protection = Protection(locked=False)
    worksheet.sheet_state = "hidden"


def build_control_log_workbook(path, snapshot):
    """Write control_log.xlsx (Dashboard + Control Log + Source Files + _Mapping)."""
    workbook = Workbook()

    dashboard = workbook.active
    dashboard.title = DASHBOARD_SHEET
    _build_dashboard(dashboard, snapshot)

    control_sheet = workbook.create_sheet(CONTROL_SHEET)
    _write_control_sheet(control_sheet, snapshot["questions"])
    _add_table(control_sheet, TABLE_CONTROL, len(CONTROL_HEADERS), len(snapshot["questions"]))

    source_sheet = workbook.create_sheet(SOURCE_FILES_SHEET)
    _write_source_sheet(source_sheet, snapshot["sourceJobs"])
    _add_table(source_sheet, TABLE_SOURCES, len(SOURCE_HEADERS), len(snapshot["sourceJobs"]))

    mapping = workbook.create_sheet(MAPPING_SHEET)
    _write_mapping(mapping, snapshot)

    try:
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        workbook.calculation.calcMode = "auto"
    except Exception:  # noqa: BLE001
        pass

    workbook.save(path)
    verify_openable(path)
    return path


def snapshot_counts(snapshot):
    stats = snapshot["statistics"]
    return {
        "fileCount": len(snapshot["sourceJobs"]),
        "questionCount": stats.get("questions_total", 0),
        "threadedCount": stats.get("threaded_questions", 0),
        "legacyCount": stats.get("legacy_questions", 0),
        "authorCount": stats.get("authors_total", 0),
    }
