from openpyxl.utils import get_column_letter


def apply(workbook, audit, config):
    if not config.get("enabled", True):
        return
    freeze_header = bool(config.get("freeze_header", True))
    auto_filter = bool(config.get("auto_filter", True))
    freeze_rule = config.get("freeze_audit_rule", "freeze_header")
    filter_rule = config.get("filter_audit_rule", "auto_filter")

    for ws in workbook.worksheets:
        if ws.max_row < 1 or ws.max_column < 1:
            continue
        has_header = any(cell.value not in (None, "") for cell in ws[1])
        if not has_header:
            continue

        if freeze_header:
            before = ws.freeze_panes.coordinate if getattr(ws.freeze_panes, "coordinate", None) else ws.freeze_panes
            after = "A2"
            if before != after:
                ws.freeze_panes = after
                audit.add(
                    rule=freeze_rule,
                    sheet=ws.title,
                    target="worksheet",
                    change_type="sheet",
                    before={"freezePanes": before},
                    after={"freezePanes": after},
                    reason="\u51bb\u7ed3\u7b2c\u4e00\u884c\u8868\u5934",
                )

        if auto_filter and ws.max_row >= 2:
            end_col = get_column_letter(ws.max_column)
            after = f"A1:{end_col}{ws.max_row}"
            before = ws.auto_filter.ref
            if before != after:
                ws.auto_filter.ref = after
                audit.add(
                    rule=filter_rule,
                    sheet=ws.title,
                    target="worksheet",
                    change_type="sheet",
                    before={"autoFilter": before},
                    after={"autoFilter": after},
                    reason="\u4e3a\u6570\u636e\u533a\u57df\u542f\u7528\u81ea\u52a8\u7b5b\u9009",
                )
