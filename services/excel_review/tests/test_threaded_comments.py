"""Microsoft 365 threaded-comment extraction + split + Excel Address tests (V1.2)."""

import json
import shutil
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.comments import Comment

from services.excel_review import ooxml_threaded_comments as tc
from services.excel_review import parser, review_workbook
from services.excel_review.status import Reporter
from services.excel_review.worker import run
from services.excel_review.tests.threaded_fixture import (
    build_base_workbook,
    inject_threaded_comments,
    new_guid,
)

FIXTURES = Path(__file__).resolve().parent / "fixtures"
REAL_BASELINE = FIXTURES / "threaded_real_baseline.xlsx"


def reporter(job_dir):
    return Reporter(job_dir)


class TempBase(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def extract(self, source_path, name="job"):
        job_dir = self.tmp / name
        job_dir.mkdir(parents=True, exist_ok=True)
        shutil.copy2(str(source_path), str(job_dir / "source.xlsx"))
        run(job_dir, reporter(job_dir))
        return job_dir


class TestRealBaseline(TempBase):
    def test_real_file_detected(self):
        self.assertTrue(tc.has_threaded_comments(REAL_BASELINE))

    def test_parse_threaded_nodes(self):
        items = tc.extract_threaded_comments(REAL_BASELINE)
        self.assertEqual(len(items), 2)
        by_cell = {item["cellAddress"]: item for item in items}
        self.assertIn("B5", by_cell)
        self.assertEqual(by_cell["B5"]["sheetName"], "财务复核")
        self.assertEqual(by_cell["B5"]["author"], "审阅人")
        self.assertEqual(by_cell["B5"]["commentType"], "threaded")

    def test_full_extract_no_fallback_duplicates(self):
        job_dir = self.extract(REAL_BASELINE)
        data = json.loads((job_dir / "extracted.json").read_text(encoding="utf-8"))
        self.assertTrue(data["hasThreaded"])
        threaded = [i for i in data["items"] if i["commentType"].startswith("threaded")]
        self.assertEqual(len(threaded), 2)
        for item in data["items"]:
            if item["commentType"] == "legacy":
                self.assertFalse(item["author"].startswith("tc={"))
                self.assertNotIn("[线程批注]", item["commentText"])

    def test_review_workbook_thread_sheet_columns(self):
        job_dir = self.extract(REAL_BASELINE)
        workbook = load_workbook(job_dir / "review.xlsx")
        threaded = workbook[review_workbook.THREADED_SHEET]
        header = [c.value for c in threaded[1]]
        for name in ("类型", "原文件", "原始值", "批注内容", "Excel Address", "作者", "回复", "Source Reference"):
            self.assertIn(name, header)
        self.assertNotIn("Thread ID", header)
        self.assertNotIn("Parent Thread ID", header)
        self.assertNotIn("Person ID", header)
        self.assertNotIn("Comment Time", header)


def _build_threaded_source(path, with_legacy_reply=False):
    workbook = build_base_workbook(
        [
            ("财务复核", {"A1": "预算", "B2": 480000, "D4": "说明"}),
            ("Sheet With Space", {"C3": "x"}),
        ]
    )
    if with_legacy_reply:
        workbook["财务复核"]["H9"].comment = Comment("传统备注", "会计")
    workbook.save(path)

    alice = new_guid()
    root = new_guid()
    reply1 = new_guid()
    reply2 = new_guid()
    root2 = new_guid()
    threads = [
        {"ref": "B2", "id": root, "personId": alice, "dT": "2026-10-13T06:00:00", "text": "主批注"},
        {"ref": "B2", "id": reply1, "personId": alice, "parentId": root, "dT": "2026-10-13T06:01:00", "text": "回复1"},
        {"ref": "B2", "id": reply2, "personId": alice, "parentId": root, "dT": "2026-10-13T06:02:00", "text": "回复2"},
        {"ref": "D4", "id": root2, "personId": alice, "dT": "2026-10-13T06:03:00", "text": "第二条线程"},
    ]
    persons = [{"id": alice, "displayName": "j36740", "email": "j36740@example.com"}]
    inject_threaded_comments(str(path), "财务复核", threads, persons)
    return root, reply1, reply2, root2


class TestThreadedExtraction(TempBase):
    def test_reply_chain(self):
        path = self.tmp / "chain.xlsx"
        root, reply1, reply2, _ = _build_threaded_source(path)
        job_dir = self.extract(path)
        data = json.loads((job_dir / "extracted.json").read_text(encoding="utf-8"))
        by_id = {i["threadId"]: i for i in data["items"]}
        self.assertEqual(by_id[root]["commentType"], "threaded")
        self.assertEqual(by_id[root]["parentThreadId"], "")
        self.assertEqual(by_id[reply1]["commentType"], "threaded_reply")
        self.assertEqual(by_id[reply1]["parentThreadId"], root)
        self.assertEqual(by_id[reply2]["parentThreadId"], root)
        self.assertEqual(by_id[root]["author"], "j36740")

    def test_reply_shares_parent_excel_address(self):
        path = self.tmp / "chain.xlsx"
        root, reply1, reply2, root2 = _build_threaded_source(path)
        job_dir = self.extract(path)
        data = json.loads((job_dir / "extracted.json").read_text(encoding="utf-8"))
        by_id = {i["threadId"]: i for i in data["items"]}
        self.assertEqual(by_id[root]["excelAddress"], "'财务复核'!B2")
        self.assertEqual(by_id[reply1]["excelAddress"], by_id[root]["excelAddress"])
        self.assertEqual(by_id[reply2]["excelAddress"], by_id[root]["excelAddress"])
        self.assertEqual(by_id[root2]["excelAddress"], "'财务复核'!D4")

    def test_source_reference_format(self):
        path = self.tmp / "chain.xlsx"
        _build_threaded_source(path)
        job_dir = self.extract(path)
        data = json.loads((job_dir / "extracted.json").read_text(encoding="utf-8"))
        root = [i for i in data["items"] if i["cellAddress"] == "B2" and i["commentType"] == "threaded"][0]
        self.assertEqual(root["sourceReference"], "source.xlsx → 财务复核 → B2")

    def test_legacy_and_threaded_mixed(self):
        path = self.tmp / "mixed.xlsx"
        _build_threaded_source(path, with_legacy_reply=True)
        job_dir = self.extract(path)
        data = json.loads((job_dir / "extracted.json").read_text(encoding="utf-8"))
        types = {i["commentType"] for i in data["items"]}
        self.assertIn("legacy", types)
        self.assertIn("threaded", types)
        self.assertIn("threaded_reply", types)


class TestTwoSheetSplit(TempBase):
    def _mixed_source(self, path):
        workbook = build_base_workbook([("S", {"A1": 1, "A2": 2, "B1": 10, "B2": 20})])
        workbook["S"]["A1"].comment = Comment("legacy one", "甲")
        workbook["S"]["A2"].comment = Comment("legacy two", "乙")
        workbook.save(path)
        pid = new_guid()
        inject_threaded_comments(
            str(path),
            "S",
            [
                {"ref": "B1", "id": new_guid(), "personId": pid, "dT": "2026-10-13T06:00:00", "text": "thread one"},
                {"ref": "B2", "id": new_guid(), "personId": pid, "dT": "2026-10-13T06:01:00", "text": "thread two"},
            ],
            [{"id": pid, "displayName": "u"}],
        )

    def test_split_between_sheets(self):
        path = self.tmp / "mix.xlsx"
        self._mixed_source(path)
        job_dir = self.extract(path)
        workbook = load_workbook(job_dir / "review.xlsx")
        threaded = workbook[review_workbook.THREADED_SHEET]
        legacy = workbook[review_workbook.LEGACY_SHEET]
        t_types = [threaded.cell(row=r, column=1).value for r in range(2, threaded.max_row + 1)]
        l_types = [legacy.cell(row=r, column=1).value for r in range(2, legacy.max_row + 1)]
        self.assertEqual(sorted(t_types), ["新版批注", "新版批注"])
        self.assertEqual(sorted(l_types), ["旧版批注", "旧版批注"])

    def test_excel_address_in_both_sheets(self):
        path = self.tmp / "mix.xlsx"
        self._mixed_source(path)
        job_dir = self.extract(path)
        workbook = load_workbook(job_dir / "review.xlsx")
        for sheet_name in review_workbook.VISIBLE_SHEETS:
            sheet = workbook[sheet_name]
            header = [c.value for c in sheet[1]]
            addr_col = header.index("Excel Address") + 1
            for r in range(2, sheet.max_row + 1):
                address = sheet.cell(row=r, column=addr_col).value
                self.assertTrue(address)
                self.assertIn("!", address)


if __name__ == "__main__":
    unittest.main()
