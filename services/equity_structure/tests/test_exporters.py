from __future__ import annotations

import io
import unittest

from openpyxl import load_workbook

from services.equity_structure.exporters import build_excel, build_pdf, build_svg
from services.equity_structure.layout import build_routes, layout_structure
from services.equity_structure.schemas import (
    BBox, Diagnostics, DictionaryCorrection, Edge, ExcludedItem, Node, Quality, StructureResult,
)


def sample_result():
    nodes = [
        Node("n1", "股东甲", "person", 1, "active", 0.9, BBox(0, 0, 100, 40)),
        Node("n2", "上海示例股份有限公司", "company", 2, "active", 0.95, BBox(0, 200, 300, 250)),
        Node("n3", "已注销子公司", "company", 3, "excluded_cancelled", 0.8, BBox(0, 400, 100, 440)),
        Node("n4", "子公司", "company", 3, "active", 0.88, BBox(300, 400, 420, 450)),
    ]
    edges = [
        Edge("n1", "n2", "54%", "shareholder_to_core", 0.94, "secondary_crop_ocr"),
        Edge("n2", "n4", "30.4498%", "investment", 0.86, "global_ocr"),
    ]
    return StructureResult(
        title="上海示例股份有限公司股权结构图",
        core_company="上海示例股份有限公司",
        all_visible_text=["股东甲", "54%", "上海示例股份有限公司"],
        nodes=nodes,
        edges=edges,
        excluded=[ExcludedItem("已注销子公司", "已注销", "n3")],
        dictionary_corrections=[DictionaryCorrection("上海示例股份公司", "上海示例股份有限公司", 95.0)],
        warnings=["存在 1 个低置信度节点。"],
        quality=Quality(0.9, ["已注销子公司"], ["人工复核"]),
        diagnostics=Diagnostics(image_width=2000, image_height=1000, tile_count=2),
    )


class ExportersTest(unittest.TestCase):
    def test_excel_has_required_sheets_and_rows(self):
        data = build_excel(sample_result())
        workbook = load_workbook(io.BytesIO(data))
        self.assertEqual(workbook.sheetnames, ["股权结构", "投资关系", "识别备注"])
        self.assertEqual(workbook["股权结构"].max_row, 5)
        self.assertEqual(workbook["投资关系"].max_row, 3)
        self.assertEqual(workbook["投资关系"]["C2"].value, "54%")
        self.assertIn("已注销", str(workbook["识别备注"]["B8"].value))

    def test_svg_contains_nodes_and_percentage(self):
        svg = build_svg(sample_result())
        self.assertTrue(svg.startswith("<svg"))
        self.assertIn("上海示例股份有限公司", svg)
        self.assertIn("30.4498%", svg)
        self.assertNotIn("已注销子公司", svg)

    def test_svg_uses_one_bus_per_level(self):
        svg = build_svg(sample_result())
        self.assertEqual(svg.count('data-role="shareholder-bus"'), 1)
        self.assertEqual(svg.count('data-role="direct-bus"'), 1)

    def test_layout_overrides_are_used(self):
        layout = {
            "nodes": {"n2": {"x": 900, "y": 300, "width": 300, "height": 80}},
            "labels": {"e1": {"x": 120, "y": 210, "manual": True}},
            "buses": {"shareholder": {"y": 190, "manual": True}},
        }
        svg = build_svg(sample_result(), layout)
        self.assertIn('x="120.0"', svg)
        self.assertIn("data-role=\"shareholder-bus\"", svg)


    def test_pdf_uses_style_overrides(self):
        layout = {
            "styles": {
                "nodes": {"n1": {"fill": "#eef4ff", "stroke": "#84adff", "fontSize": 18, "strokeWidth": 2, "radius": 16}},
                "edges": {"e_stable": {"stroke": "#ff0000", "strokeWidth": 3, "dash": "5 3"}},
                "labels": {"e_stable": {"color": "#ff0000", "background": "#ffffcc", "fontSize": 14}},
                "canvas": {"background": "#f8fafc"},
            }
        }
        result = sample_result()
        result.edges[0].client_id = "e_stable"
        data = build_pdf(result, layout)
        self.assertTrue(data.startswith(b"%PDF"))
        self.assertGreater(len(data), 1000)
        svg = build_svg(result, layout)
        self.assertIn("#ff0000", svg)
        self.assertIn("#ffffcc", svg)


    def test_hidden_label_and_client_id_are_honored(self):
        result = sample_result()
        result.edges[0].client_id = "e_stable"
        layout = layout_structure(result, {"styles": {"labels": {"e_stable": {"hidden": True}}}})
        routes = build_routes(result, layout, {"styles": {"labels": {"e_stable": {"hidden": True}}}})
        self.assertFalse(any(label.text == "54%" for label in routes.labels))
        self.assertTrue(any(line.edge_id == "e_stable" for line in routes.lines))


    def test_layout_reserves_title_header_and_footer(self):
        result = sample_result()
        layout = layout_structure(result)
        self.assertTrue(all(position.y >= 120 for position in layout.nodes.values()))
        svg = build_svg(result)
        self.assertIn('y="54"', svg)
        self.assertGreater(layout.height, max(position.y + position.height for position in layout.nodes.values()))

    def test_pdf_is_generated(self):
        data = build_pdf(sample_result())
        self.assertTrue(data.startswith(b"%PDF"))
        self.assertGreater(len(data), 1000)


if __name__ == "__main__":
    unittest.main()
