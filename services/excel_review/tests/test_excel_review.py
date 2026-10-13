"""Legacy extraction + review-workbook + value-normalisation + Excel Address (V1.2).

Run from the project root with a Python that has openpyxl:

    .venv/bin/python -m unittest discover -s services/excel_review/tests -t .
"""

import json
import re
import tempfile
import unittest
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment

from services.excel_review import parser, review_workbook
from services.excel_review.ooxml_package import (
    OoxmlPackage,
    workbook_sheets,
    write_patched_package,
)
from services.excel_review.status import Reporter, ReviewError
from services.excel_review.worker import run


def build_workbook(spec):
    """spec: list of (sheet_title, {cell: value}, {cell: (comment, author)})."""
    workbook = Workbook()
    default = workbook.active
    first = True
    for title, values, comments in spec:
        if first:
            sheet = default
            sheet.title = title
            first = False
        else:
            sheet = workbook.create_sheet(title)
        for cell, value in values.items():
            sheet[cell] = value
        for cell, (text, author) in comments.items():
            sheet[cell].comment = Comment(text, author)
    return workbook


def inject_cached_value(path, sheet_name, ref, value, error=False):
    """Simulate Excel's cached formula result in the sheet XML."""
    with OoxmlPackage(path) as package:
        part = None
        for sheet in workbook_sheets(package):
            if sheet["name"] == sheet_name:
                part = sheet["part"]
                break
        if part is None:
            raise AssertionError("sheet not found: %s" % sheet_name)
        data = package.read(part).decode("utf-8")

    pattern = re.compile(r'(<c r="%s"[^>]*>)(.*?)(</c>)' % re.escape(ref), re.DOTALL)
    match = pattern.search(data)
    if not match:
        raise AssertionError("cell %s not found" % ref)
    open_tag, inner = match.group(1), match.group(2)
    if error and 't="' not in open_tag:
        open_tag = open_tag[:-1] + ' t="e">'
    new_v = "<v>%s</v>" % value
    if re.search(r"<v\s*/>", inner):
        inner = re.sub(r"<v\s*/>", new_v, inner, count=1)
    elif "<v>" in inner:
        inner = re.sub(r"<v>.*?</v>", new_v, inner, count=1, flags=re.DOTALL)
    else:
        inner = inner + new_v
    patched = data[: match.start()] + open_tag + inner + match.group(3) + data[match.end():]
    write_patched_package(path, path, {part: patched.encode("utf-8")})


def reporter(job_dir):
    return Reporter(job_dir)


class Base(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def extract(self, workbook, name="job"):
        job_dir = self.tmp / name
        job_dir.mkdir(parents=True, exist_ok=True)
        workbook.save(job_dir / "source.xlsx")
        run(job_dir, reporter(job_dir))
        return job_dir

    def extract_path(self, source_path, name="job"):
        job_dir = self.tmp / name
        job_dir.mkdir(parents=True, exist_ok=True)
        (job_dir / "source.xlsx").write_bytes(Path(source_path).read_bytes())
        run(job_dir, reporter(job_dir))
        return job_dir


class TestExcelAddressHelper(unittest.TestCase):
    def test_english(self):
        self.assertEqual(parser.format_excel_address("Sheet1", "A1"), "'Sheet1'!A1")

    def test_chinese(self):
        self.assertEqual(parser.format_excel_address("财务复核", "B5"), "'财务复核'!B5")

    def test_space(self):
        self.assertEqual(
            parser.format_excel_address("Sheet With Space", "D5"),
            "'Sheet With Space'!D5",
        )

    def test_single_quote_escaped(self):
        self.assertEqual(
            parser.format_excel_address("Manager's Review", "A10"),
            "'Manager''s Review'!A10",
        )

    def test_hidden_sheet(self):
        self.assertEqual(parser.format_excel_address("隐藏数据", "B2"), "'隐藏数据'!B2")

    def test_empty(self):
        self.assertEqual(parser.format_excel_address("", "B2"), "")
        self.assertEqual(parser.format_excel_address("S", ""), "")


class TestWorkbookLayout(Base):
    def _legacy_job(self):
        wb = build_workbook([("S", {"A1": 1}, {"A1": ("c", "A")})])
        return self.extract(wb)

    def test_sheet_order_dashboard_first(self):
        job_dir = self._legacy_job()
        workbook = load_workbook(job_dir / "review.xlsx")
        self.assertEqual(
            workbook.sheetnames,
            [
                review_workbook.DASHBOARD_SHEET,
                review_workbook.THREADED_SHEET,
                review_workbook.LEGACY_SHEET,
                review_workbook.MAPPING_SHEET,
            ],
        )

    def test_visible_headers_and_source_reference_last(self):
        job_dir = self._legacy_job()
        workbook = load_workbook(job_dir / "review.xlsx")
        expected = [
            "类型",
            "原文件",
            "Sheet",
            "Cell",
            "原始值",
            "批注内容",
            "Excel Address",
            "作者",
            "回复",
            "状态",
            "Source Reference",
        ]
        self.assertEqual(review_workbook.VISIBLE_HEADERS, expected)
        self.assertEqual(expected[-1], "Source Reference")
        self.assertEqual(expected[6], "Excel Address")
        for sheet_name in review_workbook.VISIBLE_SHEETS:
            header = [c.value for c in workbook[sheet_name][1]][: len(expected)]
            self.assertEqual(header, expected)

    def test_item_id_column_hidden(self):
        job_dir = self._legacy_job()
        workbook = load_workbook(job_dir / "review.xlsx")
        for sheet_name in review_workbook.VISIBLE_SHEETS:
            sheet = workbook[sheet_name]
            id_col = len(review_workbook.VISIBLE_HEADERS) + 1
            self.assertEqual(sheet.cell(row=1, column=id_col).value, "_ITEM_ID")
            self.assertTrue(sheet.column_dimensions[sheet.cell(row=1, column=id_col).column_letter].hidden)

    def test_technical_fields_not_visible(self):
        job_dir = self._legacy_job()
        workbook = load_workbook(job_dir / "review.xlsx")
        forbidden = {"Review ID", "Item ID", "Thread ID", "Parent Thread ID", "Person ID", "Comment Time", "mapping_token"}
        for sheet_name in review_workbook.VISIBLE_SHEETS:
            header = {str(c.value) for c in workbook[sheet_name][1]}
            self.assertFalse(header & forbidden, sheet_name)

    def test_mapping_hidden_and_v3(self):
        job_dir = self._legacy_job()
        workbook = load_workbook(job_dir / "review.xlsx")
        self.assertEqual(workbook["_Mapping"].sheet_state, "hidden")
        self.assertEqual(
            [c.value for c in workbook["_Mapping"][1]],
            review_workbook.MAPPING_HEADERS,
        )
        self.assertEqual(workbook["_Mapping"]["A2"].value, "excel-review-mapping-v3")
        self.assertIn("excel_address", review_workbook.MAPPING_HEADERS)

    def test_no_worksheet_protection(self):
        job_dir = self._legacy_job()
        workbook = load_workbook(job_dir / "review.xlsx")
        for sheet in workbook.worksheets:
            self.assertFalse(sheet.protection.sheet, sheet.title)

    def test_source_reference_and_excel_address_present(self):
        job_dir = self._legacy_job()
        sheet = load_workbook(job_dir / "review.xlsx")[review_workbook.LEGACY_SHEET]
        header = [c.value for c in sheet[1]]
        self.assertEqual(sheet.cell(row=2, column=header.index("Source Reference") + 1).value, "source.xlsx → S → A1")
        self.assertEqual(sheet.cell(row=2, column=header.index("Excel Address") + 1).value, "'S'!A1")


class TestOriginalValueNormalization(unittest.TestCase):
    def test_helpers(self):
        n = parser.normalize_display_value
        self.assertEqual(n(None), "")
        self.assertEqual(n("文本"), "文本")
        self.assertEqual(n(92000), 92000)
        self.assertEqual(n(50000.25), 50000.25)
        self.assertEqual(n(True), "TRUE")
        self.assertEqual(n(date(2026, 10, 13)), "2026-10-13")
        for err in ("#VALUE!", "#REF!", "#DIV/0!", "#NAME?", "#N/A"):
            self.assertEqual(n(err), "", err)


class TestOriginalValueInWorkbook(Base):
    def test_text_number_date_blank(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "V"
        ws["A1"] = "合同编号 HT-2026-001"
        ws["B2"] = 92000
        ws["C3"] = 50000.25
        ws["D4"] = date(2026, 10, 13)
        ws["E5"] = None
        for cell in ("A1", "B2", "C3", "D4", "E5"):
            ws[cell].comment = Comment("c-" + cell, "a")
        job_dir = self.extract(wb)

        data = json.loads((job_dir / "extracted.json").read_text(encoding="utf-8"))
        values = {item["cellAddress"]: item["cellValue"] for item in data["items"]}
        self.assertEqual(values["A1"], "合同编号 HT-2026-001")
        self.assertEqual(values["B2"], 92000)
        self.assertEqual(values["C3"], 50000.25)
        self.assertEqual(values["D4"], "2026-10-13")
        self.assertEqual(values["E5"], "")

    def test_formula_cached_values(self):
        path = self.tmp / "source.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "Calc"
        ws["B6"] = 100
        ws["C6"] = 200
        ws["D6"] = "=B6+C6"
        ws["E6"] = "=1+1"
        ws["F6"] = "=B6+C6"
        ws["D6"].comment = Comment("calc", "a")
        ws["E6"].comment = Comment("err", "a")
        ws["F6"].comment = Comment("nocache", "a")
        wb.save(path)
        inject_cached_value(path, "Calc", "D6", "300")
        inject_cached_value(path, "Calc", "E6", "#VALUE!", error=True)

        job_dir = self.extract_path(path)
        data = json.loads((job_dir / "extracted.json").read_text(encoding="utf-8"))
        values = {item["cellAddress"]: item["cellValue"] for item in data["items"]}
        self.assertEqual(values["D6"], 300)
        self.assertEqual(values["E6"], "")
        self.assertEqual(values["F6"], "")

        sheet = load_workbook(job_dir / "review.xlsx")[review_workbook.LEGACY_SHEET]
        header = [c.value for c in sheet[1]]
        val_col = header.index("原始值") + 1
        for r in range(2, sheet.max_row + 1):
            value = str(sheet.cell(row=r, column=val_col).value)
            self.assertNotIn("=", value)
            self.assertNotIn("#VALUE!", value)
            self.assertNotIn("#REF!", value)


class TestSplitByType(Base):
    def test_legacy_goes_to_legacy_sheet(self):
        wb = build_workbook([("S", {"A1": 1}, {"A1": ("legacy note", "A")})])
        job_dir = self.extract(wb)
        workbook = load_workbook(job_dir / "review.xlsx")
        self.assertEqual(workbook[review_workbook.LEGACY_SHEET]["A2"].value, "旧版批注")
        self.assertIsNone(workbook[review_workbook.THREADED_SHEET]["A2"].value)


class TestDashboard(Base):
    def _legacy_job(self):
        wb = build_workbook([
            ("财务", {"A1": 1, "B2": 2}, {"A1": ("c1", "Alice"), "B2": ("c2", "Bob")}),
        ])
        return self.extract(wb)

    def test_dashboard_first_and_unprotected(self):
        job_dir = self._legacy_job()
        workbook = load_workbook(job_dir / "review.xlsx")
        self.assertEqual(workbook.sheetnames[0], review_workbook.DASHBOARD_SHEET)
        for sheet in workbook.worksheets:
            self.assertFalse(sheet.protection.sheet, sheet.title)

    def test_legacy_table_exists(self):
        job_dir = self._legacy_job()
        workbook = load_workbook(job_dir / "review.xlsx")
        tables = workbook[review_workbook.LEGACY_SHEET].tables
        self.assertIn(review_workbook.TABLE_LEGACY, tables)
        self.assertEqual(tables[review_workbook.TABLE_LEGACY].ref, "A1:K3")

    def test_no_threaded_table_when_empty(self):
        job_dir = self._legacy_job()
        workbook = load_workbook(job_dir / "review.xlsx")
        self.assertNotIn(review_workbook.TABLE_THREADED, workbook[review_workbook.THREADED_SHEET].tables)

    def test_dashboard_uses_formulas(self):
        job_dir = self._legacy_job()
        dashboard = load_workbook(job_dir / "review.xlsx")[review_workbook.DASHBOARD_SHEET]
        total = str(dashboard["B5"].value)
        answered = str(dashboard["B6"].value)
        self.assertTrue(total.startswith("="), total)
        self.assertIn(review_workbook.TABLE_LEGACY, total)
        self.assertIn(review_workbook.TABLE_THREADED, total)
        self.assertTrue(answered.startswith("="), answered)
        self.assertIn("回复", answered)
        # Unanswered and rate are formulas too.
        self.assertTrue(str(dashboard["B7"].value).startswith("="))
        self.assertTrue(str(dashboard["B8"].value).startswith("="))

    def test_dashboard_author_rows(self):
        job_dir = self._legacy_job()
        dashboard = load_workbook(job_dir / "review.xlsx")[review_workbook.DASHBOARD_SHEET]
        self.assertTrue(str(dashboard["B14"].value).startswith("="))  # authors_total via formula
        # Row 18/19 authors are Alice/Bob (sorted), counts via formulas.
        authors = {dashboard.cell(row=r, column=1).value for r in (18, 19)}
        self.assertEqual(authors, {"Alice", "Bob"})
        for r in (18, 19):
            self.assertTrue(str(dashboard.cell(row=r, column=2).value).startswith("="))
            self.assertTrue(str(dashboard.cell(row=r, column=5).value).startswith("="))

    def test_calc_settings(self):
        job_dir = self._legacy_job()
        workbook = load_workbook(job_dir / "review.xlsx")
        self.assertTrue(workbook.calculation.fullCalcOnLoad)


class TestEveryRowHasExcelAddress(Base):
    def test_all_visible_rows_have_address(self):
        wb = build_workbook([
            ("财务 复核", {"A1": 1, "B2": 2}, {"A1": ("c1", "甲"), "B2": ("c2", "乙")}),
            ("Sheet With Space", {"C3": 3}, {"C3": ("c3", "丙")}),
        ])
        job_dir = self.extract(wb)
        workbook = load_workbook(job_dir / "review.xlsx")
        for sheet_name in review_workbook.VISIBLE_SHEETS:
            sheet = workbook[sheet_name]
            header = [c.value for c in sheet[1]]
            addr_col = header.index("Excel Address") + 1
            for r in range(2, sheet.max_row + 1):
                if sheet.cell(row=r, column=1).value is None:
                    continue
                address = sheet.cell(row=r, column=addr_col).value
                self.assertTrue(address, "missing address row %d on %s" % (r, sheet_name))
                self.assertIn("!", address)
                self.assertTrue(address.startswith("'"))


if __name__ == "__main__":
    unittest.main()
