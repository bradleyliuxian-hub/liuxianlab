"""Read-only .xlsx / .xls intake and header-row inspection.

The reader never performs accounting logic. It only opens a workbook, lists
sheets, samples rows, reads a selected sheet, and proposes header-row candidates
for the mapping layer.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, Optional, Sequence, Set, Tuple

from openpyxl import load_workbook

from .errors import WorkbookValidationError

SUPPORTED_SUFFIXES = (".xlsx", ".xls")
DEFAULT_SAMPLE_ROWS = 30
_XLS_SHEET_STATES = {0: "visible", 1: "hidden", 2: "veryHidden"}


def _display(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _is_non_empty(value: Any) -> bool:
    return value is not None and _display(value) != ""


@dataclass(frozen=True)
class SheetInfo:
    name: str
    index: int
    max_row: int
    max_column: int
    sheet_state: str = "visible"

    def to_dict(self) -> Dict[str, Any]:
        return {
            "name": self.name,
            "index": self.index,
            "max_row": self.max_row,
            "max_column": self.max_column,
            "sheet_state": self.sheet_state,
        }


@dataclass(frozen=True)
class SheetBounds:
    min_row: int
    max_row: int
    min_col: int
    max_col: int
    non_empty_count: int

    def to_dict(self) -> Dict[str, Any]:
        return {
            "min_row": self.min_row,
            "max_row": self.max_row,
            "min_col": self.min_col,
            "max_col": self.max_col,
            "non_empty_count": self.non_empty_count,
        }


@dataclass(frozen=True)
class HeaderCandidate:
    row_number: int
    values: Tuple[str, ...]
    non_empty_count: int
    text_count: int
    unique_count: int
    score: float

    def to_dict(self) -> Dict[str, Any]:
        return {
            "row_number": self.row_number,
            "values": list(self.values),
            "non_empty_count": self.non_empty_count,
            "text_count": self.text_count,
            "unique_count": self.unique_count,
            "score": round(self.score, 4),
        }


@dataclass(frozen=True)
class SheetData:
    sheet_name: str
    rows: Tuple[Tuple[Any, ...], ...]
    max_row: int
    max_column: int
    header_candidates: Tuple[HeaderCandidate, ...] = ()
    bounds: SheetBounds = SheetBounds(0, 0, 0, 0, 0)

    def to_dict(self, include_rows: bool = True) -> Dict[str, Any]:
        payload = {
            "sheet_name": self.sheet_name,
            "max_row": self.max_row,
            "max_column": self.max_column,
            "bounds": self.bounds.to_dict(),
            "header_candidates": [
                candidate.to_dict() for candidate in self.header_candidates
            ],
        }
        if include_rows:
            payload["rows"] = [list(row) for row in self.rows]
        return payload


@dataclass(frozen=True)
class SheetInspection:
    sheet: SheetInfo
    sample_rows: Tuple[Tuple[str, ...], ...]
    header_candidates: Tuple[HeaderCandidate, ...]
    bounds: SheetBounds = SheetBounds(0, 0, 0, 0, 0)

    def to_dict(self) -> Dict[str, Any]:
        return {
            "sheet": self.sheet.to_dict(),
            "bounds": self.bounds.to_dict(),
            "sample_rows": [list(row) for row in self.sample_rows],
            "header_candidates": [
                candidate.to_dict() for candidate in self.header_candidates
            ],
        }


def detect_header_candidates(
    rows: Sequence[Sequence[Any]],
    *,
    limit: int = 3,
) -> Tuple[HeaderCandidate, ...]:
    """Score likely header rows without AI or business-specific guesses.

    The score rewards multiple non-empty cells, text cells, and unique labels,
    and lightly penalises rows further down the sheet. It returns several
    candidates rather than pretending one is always certain.
    """

    candidates = []
    for offset, row in enumerate(rows, start=1):
        values = tuple(_display(value) for value in row)
        non_empty_values = [value for value in values if value]
        non_empty = len(non_empty_values)
        if non_empty < 2:
            continue

        text_count = sum(
            1
            for value in values
            if value and not _looks_numeric(value)
        )
        unique_count = len(set(non_empty_values))

        text_ratio = text_count / float(non_empty)
        unique_ratio = unique_count / float(non_empty)
        score = (
            non_empty * 2.0
            + text_ratio * 10.0
            + unique_ratio * 5.0
            - offset * 0.01
        )

        candidates.append(
            HeaderCandidate(
                row_number=offset,
                values=values,
                non_empty_count=non_empty,
                text_count=text_count,
                unique_count=unique_count,
                score=score,
            )
        )

    candidates.sort(key=lambda item: (-item.score, item.row_number))
    result = tuple(candidates[: max(1, int(limit))])
    return result


def detect_header_row(
    rows: Sequence[Sequence[Any]],
) -> Optional[int]:
    candidates = detect_header_candidates(rows, limit=1)
    return candidates[0].row_number if candidates else None


def _non_empty_bounds(
    rows: Sequence[Sequence[Any]],
) -> SheetBounds:
    min_row = None
    max_row = 0
    min_col = None
    max_col = 0
    count = 0
    for row_index, row in enumerate(rows, start=1):
        for col_index, value in enumerate(row, start=1):
            if not _is_non_empty(value):
                continue
            count += 1
            min_row = row_index if min_row is None else min(min_row, row_index)
            max_row = max(max_row, row_index)
            min_col = col_index if min_col is None else min(min_col, col_index)
            max_col = max(max_col, col_index)
    if min_row is None:
        return SheetBounds(0, 0, 0, 0, 0)
    return SheetBounds(
        min_row=min_row,
        max_row=max_row,
        min_col=min_col or 0,
        max_col=max_col,
        non_empty_count=count,
    )


def _looks_numeric(value: str) -> bool:
    try:
        float(value.replace(",", ""))
    except ValueError:
        return False
    return True


class WorkbookReader:
    """Read-only workbook access for .xlsx (openpyxl) and .xls (xlrd)."""

    def __init__(self, path: Any) -> None:
        self.path = Path(path)
        suffix = self.path.suffix.lower()
        if suffix not in SUPPORTED_SUFFIXES:
            raise WorkbookValidationError(
                "仅支持 .xlsx / .xls 文件。",
                details={"suffix": self.path.suffix},
            )
        self.suffix = suffix
        if not self.path.exists():
            raise WorkbookValidationError(
                "文件不存在。",
                details={"path": str(self.path)},
            )
        if not self.path.is_file():
            raise WorkbookValidationError(
                "路径不是文件。",
                details={"path": str(self.path)},
            )

    def _open(self):
        try:
            return load_workbook(
                str(self.path),
                read_only=True,
                data_only=True,
                keep_links=True,
            )
        except Exception as exc:  # noqa: BLE001 - normalise openpyxl errors
            raise WorkbookValidationError(
                "无法读取该 Excel 工作簿，文件可能已损坏或受密码保护。",
                details={"path": str(self.path)},
            ) from exc

    def list_sheets(self) -> Tuple[SheetInfo, ...]:
        if self.suffix == ".xls":
            return self._list_xls_sheets()
        workbook = self._open()
        try:
            sheets = []
            for index, name in enumerate(workbook.sheetnames):
                worksheet = workbook[name]
                sheets.append(
                    SheetInfo(
                        name=name,
                        index=index,
                        max_row=int(worksheet.max_row or 0),
                        max_column=int(worksheet.max_column or 0),
                        sheet_state=worksheet.sheet_state,
                    )
                )
            return tuple(sheets)
        finally:
            workbook.close()

    def read_sheet(
        self,
        sheet_name: str,
        *,
        max_rows: Optional[int] = None,
    ) -> SheetData:
        if self.suffix == ".xls":
            return self._read_xls_sheet(sheet_name, max_rows=max_rows)
        workbook = self._open()
        try:
            if sheet_name not in workbook.sheetnames:
                raise WorkbookValidationError(
                    "工作表不存在。",
                    details={
                        "sheet_name": sheet_name,
                        "available_sheets": list(workbook.sheetnames),
                    },
                )
            worksheet = workbook[sheet_name]
            rows = []
            for index, row in enumerate(
                worksheet.iter_rows(values_only=True),
                start=1,
            ):
                if max_rows is not None and index > max_rows:
                    break
                rows.append(tuple(row))
            if not rows:
                return SheetData(
                    sheet_name=sheet_name,
                    rows=(),
                    max_row=0,
                    max_column=0,
                    header_candidates=(),
                )
            max_column = max(len(row) for row in rows)
            normalized_rows = tuple(
                tuple(row) + (None,) * (max_column - len(row))
                for row in rows
            )
            bounds = _non_empty_bounds(normalized_rows)
            return SheetData(
                sheet_name=sheet_name,
                rows=normalized_rows,
                max_row=len(normalized_rows),
                max_column=max_column,
                header_candidates=detect_header_candidates(
                    normalized_rows[:DEFAULT_SAMPLE_ROWS]
                ),
                bounds=bounds,
            )
        finally:
            workbook.close()

    def formula_cells(self, sheet_name: str) -> Set[Tuple[int, int]]:
        """Return (row, zero-based column) formula cells.

        XLSX: detected via a second formula view.
        XLS: xlrd exposes cached cell values, not formula metadata; returns an
        empty set. This is documented as an XLS limitation.
        """

        if self.suffix == ".xls":
            return set()
        workbook = self._open_formula_view()
        try:
            if sheet_name not in workbook.sheetnames:
                raise WorkbookValidationError(
                    "工作表不存在。",
                    details={"sheet_name": sheet_name},
                )
            worksheet = workbook[sheet_name]
            cells = set()
            for row in worksheet.iter_rows():
                for cell in row:
                    value = getattr(cell, "value", None)
                    if getattr(cell, "data_type", None) == "f" or (
                        isinstance(value, str) and value.startswith("=")
                    ):
                        cells.add((cell.row, cell.column - 1))
            return cells
        finally:
            workbook.close()

    def _open_formula_view(self):
        try:
            return load_workbook(
                str(self.path),
                read_only=True,
                data_only=False,
                keep_links=True,
            )
        except Exception as exc:  # noqa: BLE001
            raise WorkbookValidationError(
                "无法读取该 Excel 工作簿，文件可能已损坏或受密码保护。",
                details={"path": str(self.path)},
            ) from exc

    def _open_xls(self):
        try:
            import xlrd
        except Exception as exc:  # noqa: BLE001
            raise WorkbookValidationError(
                "读取 .xls 需要 xlrd 依赖。",
                details={"path": str(self.path), "dependency": "xlrd"},
            ) from exc
        try:
            return xlrd.open_workbook(str(self.path), on_demand=True)
        except Exception as exc:  # noqa: BLE001 - normalise xlrd errors
            raise WorkbookValidationError(
                "无法读取该 Excel 工作簿，文件可能已损坏或受密码保护。",
                details={"path": str(self.path)},
            ) from exc

    def _list_xls_sheets(self) -> Tuple[SheetInfo, ...]:
        workbook = self._open_xls()
        try:
            sheets = []
            for index, name in enumerate(workbook.sheet_names()):
                sheet = workbook.sheet_by_name(name)
                visibility = int(getattr(sheet, "visibility", 0) or 0)
                sheets.append(
                    SheetInfo(
                        name=name,
                        index=index,
                        max_row=int(sheet.nrows),
                        max_column=int(sheet.ncols),
                        sheet_state=_XLS_SHEET_STATES.get(
                            visibility,
                            "visible",
                        ),
                    )
                )
            return tuple(sheets)
        finally:
            workbook.release_resources()

    def _read_xls_sheet(
        self,
        sheet_name: str,
        *,
        max_rows: Optional[int] = None,
    ) -> SheetData:
        workbook = self._open_xls()
        try:
            names = workbook.sheet_names()
            if sheet_name not in names:
                raise WorkbookValidationError(
                    "工作表不存在。",
                    details={
                        "sheet_name": sheet_name,
                        "available_sheets": list(names),
                    },
                )
            sheet = workbook.sheet_by_name(sheet_name)
            rows = []
            for row_index in range(sheet.nrows):
                if max_rows is not None and row_index >= max_rows:
                    break
                rows.append(
                    tuple(
                        self._xls_cell_value(workbook, sheet.cell(row_index, col_index))
                        for col_index in range(sheet.ncols)
                    )
                )
            if not rows:
                return SheetData(
                    sheet_name=sheet_name,
                    rows=(),
                    max_row=0,
                    max_column=0,
                    header_candidates=(),
                )
            max_column = max(len(row) for row in rows)
            normalized_rows = tuple(
                tuple(row) + (None,) * (max_column - len(row))
                for row in rows
            )
            bounds = _non_empty_bounds(normalized_rows)
            return SheetData(
                sheet_name=sheet_name,
                rows=normalized_rows,
                max_row=len(normalized_rows),
                max_column=max_column,
                header_candidates=detect_header_candidates(
                    normalized_rows[:DEFAULT_SAMPLE_ROWS]
                ),
                bounds=bounds,
            )
        finally:
            workbook.release_resources()

    @staticmethod
    def _xls_cell_value(workbook, cell):
        import xlrd

        if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
            return None
        if cell.ctype == xlrd.XL_CELL_DATE:
            try:
                return xlrd.xldate_as_datetime(cell.value, workbook.datemode)
            except Exception:  # noqa: BLE001
                return cell.value
        if cell.ctype == xlrd.XL_CELL_BOOLEAN:
            return bool(cell.value)
        if cell.ctype == xlrd.XL_CELL_ERROR:
            return xlrd.error_text_from_code.get(
                int(cell.value),
                "#VALUE!",
            )
        if cell.ctype == xlrd.XL_CELL_NUMBER:
            value = float(cell.value)
            return int(value) if value.is_integer() else value
        return cell.value

    def inspect_sheet(
        self,
        sheet_name: str,
        *,
        sample_rows: int = DEFAULT_SAMPLE_ROWS,
    ) -> SheetInspection:
        info = next(
            (sheet for sheet in self.list_sheets() if sheet.name == sheet_name),
            None,
        )
        if info is None:
            raise WorkbookValidationError(
                "工作表不存在。",
                details={"sheet_name": sheet_name},
            )
        data = self.read_sheet(sheet_name, max_rows=sample_rows)
        return SheetInspection(
            sheet=info,
            sample_rows=tuple(
                tuple(_display(value) for value in row)
                for row in data.rows
            ),
            header_candidates=data.header_candidates,
            bounds=data.bounds,
        )

    def inspect_workbook(
        self,
        *,
        sample_rows: int = DEFAULT_SAMPLE_ROWS,
    ) -> Tuple[SheetInspection, ...]:
        return tuple(
            self.inspect_sheet(sheet.name, sample_rows=sample_rows)
            for sheet in self.list_sheets()
        )

    def read_table(
        self,
        sheet_name: str,
        header_row: int,
        *,
        max_data_rows: Optional[int] = None,
    ) -> Tuple[Tuple[str, ...], Tuple[Tuple[Any, ...], ...], int]:
        """Return headers, data rows, and the last Excel row number read.

        ``header_row`` is 1-based. Data rows start immediately after it.
        Empty trailing rows are retained only if they are inside the read range;
        the normalization layer decides whether they are meaningful.
        """

        header_row = int(header_row)
        if header_row < 1:
            raise WorkbookValidationError(
                "header_row 必须 >= 1。",
                details={"header_row": header_row},
            )
        data = self.read_sheet(sheet_name)
        if header_row > len(data.rows):
            raise WorkbookValidationError(
                "header_row 超出工作表范围。",
                details={
                    "header_row": header_row,
                    "max_row": len(data.rows),
                },
            )

        headers = tuple(
            _display(value) for value in data.rows[header_row - 1]
        )
        data_rows = data.rows[header_row:]
        if max_data_rows is not None:
            data_rows = data_rows[: max(0, int(max_data_rows))]

        return         headers, data_rows, header_row + len(data_rows)
