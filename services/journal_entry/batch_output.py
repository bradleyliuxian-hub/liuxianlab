"""Batch output generator.

Consumes completed child result.json files and produces:
- batch-summary.xlsx
- merged-standardized-journal.xlsx
- manifest.json
- batch-report.zip

It never opens source workbooks and never runs Reader / Mapping / Normalize /
Business / Contra / Anomaly. All accounting facts come from existing child
results.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
import zipfile
from collections import Counter, defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

from openpyxl import Workbook
from openpyxl.styles import Font

from .export import HEADER_FILL, _write_sheet, build_export, safe_excel_text

BATCH_OUTPUT_SCHEMA_VERSION = 1
ANOMALY_SEVERITIES = ("HIGH", "MEDIUM", "LOW", "INFO")
QUALITY_NOTE = "explicit skipped line count: NOT AVAILABLE; derived raw-minus-line is shown above"
ILLEGAL_FILENAME = re.compile(r'[\u0000-\u001f\u007f/\\:*?"<>|]+')


class BatchOutputError(Exception):
    """Safe, user-facing batch output failure."""


def _amount(value: Any) -> Optional[Decimal]:
    if value in (None, ""):
        return None
    try:
        return Decimal(str(value))
    except (InvalidOperation, ValueError):
        return None


def _rate(numerator: Any, denominator: Any) -> Optional[Decimal]:
    try:
        n = Decimal(str(numerator))
        d = Decimal(str(denominator))
    except (InvalidOperation, ValueError):
        return None
    if d <= 0:
        return None
    return n / d


def _text(value: Any) -> Any:
    return safe_excel_text(value)


def _rate_display(numerator: Any, denominator: Any) -> Any:
    value = _rate(numerator, denominator)
    return "—" if value is None else value


def _provided_rate(value: Any) -> Any:
    if value is None:
        return None
    try:
        return Decimal(str(value))
    except (InvalidOperation, ValueError):
        return None


def _provided_or_computed(provided: Any, numerator: Any, denominator: Any) -> Any:
    value = _provided_rate(provided)
    return value if value is not None else _rate_display(numerator, denominator)


def _display(value: Any) -> Any:
    return "—" if value is None else _text(value)


def _load_result(path: Path) -> Dict[str, Any]:
    if not path.exists():
        raise BatchOutputError(f"结果文件不存在：{path.name}")
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError) as exc:
        raise BatchOutputError(f"结果文件无法读取：{path.name}") from exc
    if not isinstance(payload, dict) or not isinstance(payload.get("vouchers"), list):
        raise BatchOutputError(f"结果文件结构不完整：{path.name}")
    return payload


def _summary_numbers(summary: Any) -> Dict[str, Any]:
    if not isinstance(summary, dict):
        return {}
    return summary


def _file_rows(
    included: Sequence[Dict[str, Any]],
    excluded: Sequence[Dict[str, Any]],
    results: Dict[int, Dict[str, Any]],
) -> List[List[Any]]:
    rows: List[List[Any]] = []
    included_positions = {item.get("position") for item in included}
    for item in sorted(included + excluded, key=lambda value: value.get("position", 0)):
        position = int(item.get("position") or 0)
        is_included = position in included_positions and item in included
        summary = _summary_numbers(item.get("summary"))
        result = results.get(position, {})
        generated_at = result.get("generated_at") if is_included else None
        if is_included:
            vouchers = summary.get("voucher_count")
            classified = summary.get("business_classified_count")
            findings = summary.get("finding_voucher_count")
            rows.append([
                position + 1,
                _text(item.get("original_filename") or "—"),
                item.get("job_id") or "—",
                item.get("status") or "—",
                "是" if is_included else "否",
                item.get("amount_layout") or "—",
                item.get("positive_side") or "—",
                summary.get("raw_row_count") if summary.get("raw_row_count") is not None else "—",
                summary.get("line_count") if summary.get("line_count") is not None else "—",
                vouchers if vouchers is not None else "—",
                classified if classified is not None else "—",
                _rate_display(classified, vouchers) if classified is not None and vouchers is not None else "—",
                summary.get("finding_count") if summary.get("finding_count") is not None else "—",
                findings if findings is not None else "—",
                _rate_display(findings, vouchers) if findings is not None and vouchers is not None else "—",
                summary.get("issue_count") if summary.get("issue_count") is not None else "—",
                summary.get("unbalanced_voucher_count") if summary.get("unbalanced_voucher_count") is not None else "—",
                generated_at or item.get("completed_at") or "—",
            ])
        else:
            rows.append([
                position + 1,
                _text(item.get("original_filename") or "—"),
                item.get("job_id") or "—",
                item.get("status") or "—",
                "否",
                item.get("amount_layout") or "—",
                item.get("positive_side") or "—",
                "—", "—", "—", "—", "—", "—", "—", "—", "—", "—",
                "—",
            ])
    return rows


def _business_rows(results: Iterable[Dict[str, Any]], voucher_total: Any) -> List[List[Any]]:
    counts: Counter = Counter()
    labels: Dict[str, str] = {}
    for result in results:
        summary = result.get("business_classification_summary") or {}
        for code, count in (summary.get("business_counts") or {}).items():
            counts[code] += int(count or 0)
        for code, label in (summary.get("business_labels") or {}).items():
            labels.setdefault(code, str(label))
    if not labels:
        codes = sorted(counts)
    else:
        codes = sorted(set(labels) | set(counts))
    denominator = _amount(voucher_total)
    rows = []
    for code in codes:
        count = counts.get(code, 0)
        rate = _rate(count, denominator) if denominator is not None else None
        rows.append([code, labels.get(code, code), count, rate if rate is not None else "—"])
    return rows


def _anomaly_data(results: Iterable[Dict[str, Any]]) -> Tuple[List[List[Any]], List[List[Any]]]:
    severity_counts: Counter = Counter()
    severity_vouchers: Dict[str, set] = defaultdict(set)
    code_counts: Counter = Counter()
    code_vouchers: Dict[str, set] = defaultdict(set)
    for result in results:
        findings = result.get("findings") or []
        if not isinstance(findings, list):
            continue
        for finding in findings:
            if not isinstance(finding, dict):
                continue
            severity = str(finding.get("severity") or "INFO")
            code = str(finding.get("code") or "")
            voucher = finding.get("voucher_key") or finding.get("voucher_id") or ""
            severity_counts[severity] += 1
            code_counts[code] += 1
            if voucher:
                severity_vouchers[severity].add(voucher)
                code_vouchers[code].add(voucher)
    severity_rows = [
        [severity, severity_counts.get(severity, 0), len(severity_vouchers[severity])]
        for severity in ANOMALY_SEVERITIES
    ]
    codes = sorted(code for code in code_counts if code)
    code_rows = [
        [code, code_counts[code], len(code_vouchers[code])]
        for code in codes
    ]
    return severity_rows, code_rows


def _quality_data(results: Iterable[Dict[str, Any]]) -> Tuple[Dict[str, int], List[List[Any]]]:
    totals = {
        "validation_issues": 0,
        "unbalanced_vouchers": 0,
        "raw_rows": 0,
        "journal_lines": 0,
    }
    issue_counts: Counter = Counter()
    for result in results:
        cleaning = result.get("cleaning_summary") or {}
        totals["validation_issues"] += int(cleaning.get("issue_count") or 0)
        totals["unbalanced_vouchers"] += int(cleaning.get("unbalanced_voucher_count") or 0)
        totals["raw_rows"] += int(cleaning.get("raw_row_count") or 0)
        totals["journal_lines"] += int(cleaning.get("line_count") or 0)
        for issue in result.get("validation_issues") or []:
            if isinstance(issue, dict):
                code = str(issue.get("code") or "UNKNOWN")
                issue_counts[code] += 1
        for warning in result.get("row_warnings") or []:
            if isinstance(warning, dict):
                code = str(warning.get("code") or "UNKNOWN")
                issue_counts[code] += 1
    issue_rows = [[code, count] for code, count in sorted(issue_counts.items())]
    return totals, issue_rows


def _write_overview(worksheet, batch: Dict[str, Any], summary: Dict[str, Any]) -> None:
    rows = [
        ("Batch ID", batch.get("id") or "—"),
        ("Batch 名称", batch.get("name") or "—"),
        ("生成时间", batch.get("generated_at") or "—"),
        ("文件总数", batch.get("total_files", summary.get("totalFileCount", 0))),
        ("Included completed files", batch.get("included_files", summary.get("completedFileCount", 0))),
        ("Excluded draft files", batch.get("excluded_draft", 0)),
        ("Excluded queued/running files", batch.get("excluded_active", 0)),
        ("Excluded failed files", batch.get("excluded_failed", 0)),
        ("原始总行数", _display(summary.get("rawRowCount"))),
        ("有效分录总数", _display(summary.get("lineCount"))),
        ("凭证总数", _display(summary.get("voucherCount"))),
        ("分类凭证数", _display(summary.get("businessClassifiedCount"))),
        ("业务分类覆盖率", _provided_or_computed(summary.get("businessCoverage"), summary.get("businessClassifiedCount"), summary.get("voucherCount"))),
        ("finding_count", _display(summary.get("findingCount"))),
        ("finding_voucher_count", _display(summary.get("findingVoucherCount"))),
        ("Review Rate", _provided_or_computed(summary.get("reviewRate"), summary.get("findingVoucherCount"), summary.get("voucherCount"))),
        ("validation issue count", _display(summary.get("issueCount"))),
        ("unbalanced voucher count", _display(summary.get("unbalancedVoucherCount"))),
    ]
    worksheet.append(["指标", "值"])
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
    for label, value in rows:
        worksheet.append([label, _text(value)])
    worksheet.freeze_panes = "A2"
    worksheet.column_dimensions["A"].width = 34
    worksheet.column_dimensions["B"].width = 52
    for row in worksheet.iter_rows(min_row=2):
        if row[0].value in ("业务分类覆盖率", "Review Rate"):
            row[1].number_format = "0.00%"


def _write_two_tables(worksheet, first_title, first_headers, first_rows, second_title, second_headers, second_rows):
    worksheet.append([first_title])
    worksheet.cell(row=worksheet.max_row, column=1).font = Font(bold=True)
    start = worksheet.max_row + 1
    for col, header in enumerate(first_headers, start=1):
        cell = worksheet.cell(row=start, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
    for row_index, row in enumerate(first_rows, start=start + 1):
        for col, value in enumerate(row, start=1):
            worksheet.cell(row=row_index, column=col, value=_text(value))
    worksheet.append([])
    worksheet.append([second_title])
    worksheet.cell(row=worksheet.max_row, column=1).font = Font(bold=True)
    start = worksheet.max_row + 1
    for col, header in enumerate(second_headers, start=1):
        cell = worksheet.cell(row=start, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
    for row_index, row in enumerate(second_rows, start=start + 1):
        for col, value in enumerate(row, start=1):
            worksheet.cell(row=row_index, column=col, value=_text(value))
    worksheet.column_dimensions["A"].width = 18
    worksheet.column_dimensions["B"].width = 22
    worksheet.column_dimensions["C"].width = 22


def _build_summary_workbook(
    path: Path,
    batch: Dict[str, Any],
    summary: Dict[str, Any],
    included: Sequence[Dict[str, Any]],
    excluded: Sequence[Dict[str, Any]],
    results: Dict[int, Dict[str, Any]],
) -> None:
    workbook = Workbook()
    default = workbook.active
    workbook.remove(default)

    overview = workbook.create_sheet("批次概览")
    _write_overview(overview, batch, summary)

    files = workbook.create_sheet("文件清单")
    _write_sheet(
        files,
        [
            "Position", "Original Filename", "Job ID", "Status", "Included in Export",
            "Amount Layout", "Positive Side", "Raw Rows", "Journal Lines",
            "Voucher Count", "Classified Count", "Business Coverage",
            "Finding Count", "Finding Voucher Count", "Review Rate",
            "Validation Issues", "Unbalanced Vouchers", "Result Generated At",
        ],
        _file_rows(included, excluded, results),
    )
    for row in files.iter_rows(min_row=2, min_col=12, max_col=12):
        for cell in row:
            if isinstance(cell.value, Decimal):
                cell.number_format = "0.00%"
    for row in files.iter_rows(min_row=2, min_col=15, max_col=15):
        for cell in row:
            if isinstance(cell.value, Decimal):
                cell.number_format = "0.00%"

    business = workbook.create_sheet("业务分类汇总")
    _write_sheet(
        business,
        ["Code", "Name", "Voucher Count", "Percentage of Included Vouchers"],
        _business_rows([results[pos] for pos in results], summary.get("voucherCount")),
    )
    for row in business.iter_rows(min_row=2, min_col=4, max_col=4):
        for cell in row:
            if isinstance(cell.value, Decimal):
                cell.number_format = "0.00%"

    anomaly = workbook.create_sheet("异常汇总")
    severity_rows, code_rows = _anomaly_data(results.values())
    _write_two_tables(
        anomaly,
        "风险等级",
        ["Risk Level", "Finding Count", "Finding Voucher Count"],
        severity_rows,
        "规则",
        ["Rule Code", "Finding Count", "Finding Voucher Count"],
        code_rows,
    )

    quality = workbook.create_sheet("数据质量汇总")
    totals, issue_rows = _quality_data(results.values())
    quality.append(["指标", "值"])
    for cell in quality[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
    for label, value in (
        ("validation issue count", totals["validation_issues"]),
        ("unbalanced voucher count", totals["unbalanced_vouchers"]),
        ("未形成标准分录行数（原始−有效，派生）", max(0, totals["raw_rows"] - totals["journal_lines"])),
        ("备注", QUALITY_NOTE),
    ):
        quality.append([label, _text(value)])
    quality.append([])
    quality.append(["Issue Code", "Count"])
    for col in range(1, 3):
        quality.cell(row=quality.max_row, column=col).font = Font(bold=True)
        quality.cell(row=quality.max_row, column=col).fill = HEADER_FILL
    for code, count in issue_rows:
        quality.append([_text(code), count])
    quality.column_dimensions["A"].width = 48
    quality.column_dimensions["B"].width = 24

    _safe_save(workbook, path)


def _safe_save(workbook, path: Path) -> None:
    workbook.save(path)
    try:
        path.chmod(0o600)
    except OSError:
        pass


def _voucher_lines(
    item: Dict[str, Any],
    result: Dict[str, Any],
    batch_id: str,
) -> List[List[Any]]:
    rows = []
    for voucher in result.get("vouchers") or []:
        if not isinstance(voucher, dict):
            continue
        for line in voucher.get("lines") or []:
            if not isinstance(line, dict):
                continue
            rows.append([
                batch_id,
                item.get("position", 0) + 1,
                item.get("job_id") or "—",
                _text(item.get("original_filename") or "—"),
                voucher.get("voucher_key") or "—",
                voucher.get("voucher_id") or line.get("voucher_id") or "—",
                voucher.get("date") or line.get("date") or "—",
                voucher.get("period") or line.get("period") or "—",
                voucher.get("entity") or line.get("entity") or "—",
                line.get("sheet_name") or "—",
                line.get("source_row") if line.get("source_row") is not None else "—",
                _text(line.get("line_id") or "—"),
                line.get("internal_line_key") or "—",
                line.get("account_code") or "—",
                line.get("account_name") or "—",
                _text(line.get("detail_account") or "—"),
                _text(line.get("description") or "—"),
                _amount(line.get("debit")),
                _amount(line.get("credit")),
                _text(line.get("currency") or "—"),
                _amount(line.get("original_amount")),
                _text(line.get("preparer") or "—"),
                _text(line.get("customer_code") or "—"),
                _text(line.get("customer_name") or "—"),
                _text(line.get("supplier_code") or "—"),
                _text(line.get("supplier_name") or "—"),
                _text(line.get("counterparty_code") or "—"),
                _text(line.get("counterparty_name") or "—"),
            ])
    rows.sort(key=lambda row: (
        row[1],
        str(row[9]),
        int(row[10]) if isinstance(row[10], int) else 0,
        str(row[12]),
    ))
    return rows


def _build_merged_workbook(
    path: Path,
    included: Sequence[Dict[str, Any]],
    results: Dict[int, Dict[str, Any]],
    batch_id: str,
) -> int:
    workbook = Workbook()
    default = workbook.active
    workbook.remove(default)

    sheet = workbook.create_sheet("标准分录")
    headers = [
        "Batch ID", "Batch Position", "Job ID", "Original Filename", "Voucher Key", "Voucher ID",
        "Date", "Period", "Entity", "Source Sheet", "Source Row",
        "External Line ID", "Internal Line Key", "Account Code", "Account Name",
        "Detail Account", "Description", "Debit", "Credit", "Currency",
        "Original Currency Amount", "Preparer",
        "Customer Code", "Customer Name", "Supplier Code", "Supplier Name",
        "Counterparty Code", "Counterparty Name",
    ]
    rows: List[List[Any]] = []
    for item in sorted(included, key=lambda value: value.get("position", 0)):
        result = results.get(int(item.get("position") or 0))
        if not result:
            continue
        rows.extend(_voucher_lines(item, result, batch_id))
    _write_sheet(sheet, headers, rows)

    source = workbook.create_sheet("来源文件")
    source_rows = []
    for item in sorted(included, key=lambda value: value.get("position", 0)):
        result = results.get(int(item.get("position") or 0)) or {}
        vouchers = result.get("vouchers") or []
        source_rows.append([
            item.get("position", 0) + 1,
            _text(item.get("original_filename") or "—"),
            item.get("job_id") or "—",
            item.get("sheet_name") or "—",
            item.get("header_row") if item.get("header_row") is not None else "—",
            item.get("amount_layout") or "—",
            item.get("positive_side") or "—",
            sum(len(voucher.get("lines") or []) for voucher in vouchers if isinstance(voucher, dict)),
            len(vouchers),
            "completed",
            result.get("generated_at") or item.get("completed_at") or "—",
        ])
    _write_sheet(
        source,
        [
            "Position", "Original Filename", "Job ID", "Sheet", "Header Row",
            "Amount Layout", "Positive Side", "Journal Lines", "Voucher Count",
            "Status", "Result Generated At",
        ],
        source_rows,
    )
    _safe_save(workbook, path)
    return len(rows)


def _safe_filename(value: Any, fallback: str = "workbook") -> str:
    name = str(value or "").replace("\\", "/").split("/")[-1]
    name = ILLEGAL_FILENAME.sub("_", name)
    name = name.strip(" .")
    if not name or name in (".", ".."):
        name = fallback
    return name[:80]


def _report_entry_name(position: int, original_filename: Any) -> str:
    stem = _safe_filename(original_filename, "workbook")
    if stem.lower().endswith(".xlsx") or stem.lower().endswith(".xls"):
        stem = stem.rsplit(".", 1)[0]
    return f"{position + 1:03d}-{stem}-分析报告.xlsx"


def _build_manifest(
    batch: Dict[str, Any],
    included: Sequence[Dict[str, Any]],
    excluded: Sequence[Dict[str, Any]],
    results: Dict[int, Dict[str, Any]],
    report_names: Dict[int, str],
    outputs: Dict[str, str],
) -> Dict[str, Any]:
    items = []
    included_positions = {int(item.get("position") or 0) for item in included}
    for item in sorted(included + excluded, key=lambda value: value.get("position", 0)):
        position = int(item.get("position") or 0)
        is_included = position in included_positions and item in included
        result = results.get(position) or {}
        items.append({
            "position": position,
            "job_id": item.get("job_id"),
            "original_filename": item.get("original_filename"),
            "status": item.get("status"),
            "included": is_included,
            "exclusion_reason": None if is_included else item.get("reason"),
            "result_generated_at": result.get("generated_at") if is_included else None,
            "export_filename": report_names.get(position) if is_included else None,
        })
    versions = sorted({
        str(result.get("schema_version"))
        for result in results.values()
        if result.get("schema_version") is not None
    })
    return {
        "schema_version": BATCH_OUTPUT_SCHEMA_VERSION,
        "batch_output_schema_version": BATCH_OUTPUT_SCHEMA_VERSION,
        **({"result_schema_version": versions[0]} if len(versions) == 1 else {}),
        **({"result_schema_versions": versions} if len(versions) > 1 else {}),
        "batch_id": batch.get("id"),
        "batch_name": batch.get("name"),
        "generated_at": batch.get("generated_at"),
        "total_files": batch.get("total_files", len(items)),
        "included_files": len(included),
        "excluded_files": len(excluded),
        "outputs": outputs,
        "items": items,
    }


def build_batch_output(output_dir: Path, request: Dict[str, Any]) -> Dict[str, Any]:
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    batch = request.get("batch") or {}
    summary = batch.get("summary") or {}
    included = list(request.get("included") or [])
    excluded = list(request.get("excluded") or [])
    if not included:
        raise BatchOutputError("没有已完成文件可导出。")

    results: Dict[int, Dict[str, Any]] = {}
    for item in included:
        position = int(item.get("position") or 0)
        result_path = Path(str(item.get("result_path") or ""))
        result = _load_result(result_path)
        export_path = Path(str(item.get("export_path") or ""))
        if not str(item.get("export_path") or "").strip() or export_path.is_dir():
            export_path = result_path.parent / "journal-entry-analysis-report.xlsx"
        if not export_path.exists():
            build_export(export_path, result)
            try:
                export_path.chmod(0o600)
            except OSError:
                pass
        results[position] = result

    summary_path = output_dir / "batch-summary.xlsx"
    merged_path = output_dir / "merged-standardized-journal.xlsx"
    manifest_path = output_dir / "manifest.json"
    zip_path = output_dir / "batch-report.zip"

    _build_summary_workbook(summary_path, batch, summary, included, excluded, results)
    merged_rows = _build_merged_workbook(
        merged_path, included, results, str(batch.get("id") or "")
    )

    report_names = {}
    for item in included:
        position = int(item.get("position") or 0)
        report_names[position] = _report_entry_name(position, item.get("original_filename"))
    outputs = {
        "summary": summary_path.name,
        "merged": merged_path.name,
        "manifest": manifest_path.name,
        "zip": zip_path.name,
    }
    manifest = _build_manifest(
        batch, included, excluded, results, report_names, outputs
    )
    manifest_path.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )

    with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED, allowZip64=True) as archive:
        archive.write(summary_path, arcname=summary_path.name)
        archive.write(merged_path, arcname=merged_path.name)
        archive.writestr(
            "manifest.json",
            json.dumps(manifest, ensure_ascii=False, indent=2) + "\n",
        )
        for item in sorted(included, key=lambda value: value.get("position", 0)):
            position = int(item.get("position") or 0)
            export_path = Path(str(item.get("export_path") or ""))
            archive.write(export_path, arcname=f"reports/{report_names[position]}")
    for path in (summary_path, merged_path, manifest_path, zip_path):
        try:
            path.chmod(0o600)
        except OSError:
            pass

    return {
        "summary_file": summary_path.name,
        "merged_file": merged_path.name,
        "zip_file": zip_path.name,
        "manifest_file": manifest_path.name,
        "merged_rows": merged_rows,
        "included_files": len(included),
        "excluded_files": len(excluded),
    }


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description="Liuxian Lab batch output generator")
    parser.add_argument("--output-dir", required=True)
    args = parser.parse_args(argv)
    output_dir = Path(args.output_dir).resolve()
    request_path = output_dir / "batch_request.json"
    try:
        request = json.loads(request_path.read_text(encoding="utf-8"))
        result = build_batch_output(output_dir, request)
        print(json.dumps(result, ensure_ascii=False))
        return 0
    except BatchOutputError as exc:
        print(str(exc), file=sys.stderr)
        return 1
    except Exception as exc:  # noqa: BLE001
        print(f"批量导出失败：{type(exc).__name__}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
