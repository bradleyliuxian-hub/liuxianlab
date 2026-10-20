import tempfile
import unittest
from openpyxl import Workbook
from pathlib import Path

from services.journal_entry.errors import WorkbookValidationError
from services.journal_entry.reader import (
    WorkbookReader,
    detect_header_candidates,
    detect_header_row,
)
from services.journal_entry.tests.fixture_factory import write_workbook

FIXTURES = Path(__file__).resolve().parent / "fixtures"


class ReaderTests(unittest.TestCase):
    def test_lists_sheets_and_dimensions(self):
        reader = WorkbookReader(FIXTURES / "journal_basic.xlsx")
        sheets = reader.list_sheets()
        self.assertEqual(len(sheets), 1)
        self.assertEqual(sheets[0].name, "序时账")
        self.assertEqual(sheets[0].max_row, 7)
        self.assertEqual(sheets[0].max_column, 11)

    def test_non_empty_bounds(self):
        reader = WorkbookReader(FIXTURES / "journal_basic.xlsx")
        data = reader.read_sheet("序时账")
        self.assertEqual(data.bounds.min_row, 1)
        self.assertEqual(data.bounds.max_row, 7)
        self.assertEqual(data.bounds.min_col, 1)
        self.assertEqual(data.bounds.max_col, 11)
        self.assertGreater(data.bounds.non_empty_count, 0)
        payload = reader.inspect_sheet("序时账").to_dict()
        self.assertEqual(payload["bounds"]["max_row"], 7)

    def test_reads_sheet_with_row_limit(self):
        reader = WorkbookReader(FIXTURES / "journal_basic.xlsx")
        data = reader.read_sheet("序时账", max_rows=3)
        self.assertEqual(len(data.rows), 3)
        self.assertEqual(data.rows[0][0], "日期")
        self.assertGreaterEqual(len(data.header_candidates), 1)

    def test_header_detection_on_offset_header(self):
        reader = WorkbookReader(FIXTURES / "journal_headers_offset.xlsx")
        inspection = reader.inspect_sheet("序时账", sample_rows=6)
        self.assertEqual(inspection.header_candidates[0].row_number, 4)
        self.assertEqual(
            inspection.header_candidates[0].values[:3],
            ("日期", "凭证号", "科目名称"),
        )

    def test_read_table_returns_data_after_header(self):
        reader = WorkbookReader(FIXTURES / "journal_headers_offset.xlsx")
        headers, rows, last_row = reader.read_table("序时账", 4)
        self.assertEqual(headers[:3], ("日期", "凭证号", "科目名称"))
        self.assertEqual(len(rows), 2)
        self.assertEqual(last_row, 6)
        self.assertEqual(rows[0][1], "记-101")

    def test_missing_sheet_is_structured(self):
        reader = WorkbookReader(FIXTURES / "journal_basic.xlsx")
        with self.assertRaises(WorkbookValidationError) as context:
            reader.read_sheet("不存在")
        self.assertEqual(context.exception.code, "WORKBOOK_VALIDATION_ERROR")
        self.assertIn("available_sheets", context.exception.details)

    def test_invalid_extension_is_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "book.csv"
            path.write_bytes(b"not an excel workbook")
            with self.assertRaises(WorkbookValidationError):
                WorkbookReader(path)

    def test_corrupt_xls_is_rejected_on_read(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "book.xls"
            path.write_bytes(b"not an xls workbook")
            reader = WorkbookReader(path)
            with self.assertRaises(WorkbookValidationError):
                reader.list_sheets()

    def test_missing_file_is_rejected(self):
        with self.assertRaises(WorkbookValidationError):
            WorkbookReader(FIXTURES / "does-not-exist.xlsx")

    def test_corrupt_xlsx_is_normalized_to_validation_error(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "corrupt.xlsx"
            path.write_bytes(b"PK\x03\x04broken")
            reader = WorkbookReader(path)
            with self.assertRaises(WorkbookValidationError):
                reader.list_sheets()

    def test_detect_header_helpers(self):
        rows = [
            ["Title"],
            [],
            ["日期", "凭证号", "借方", "贷方"],
            ["2026-01-01", "记-1", 100, 0],
        ]
        candidates = detect_header_candidates(rows)
        self.assertEqual(candidates[0].row_number, 3)
        self.assertEqual(detect_header_row(rows), 3)

    def test_formula_cells_are_detected(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "formula.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "S"
            sheet.append(["a", "b"])
            sheet["B2"] = "=SUM(1,2)"
            workbook.save(path)
            workbook.close()
            reader = WorkbookReader(path)
            self.assertEqual(reader.formula_cells("S"), {(2, 1)})

    def test_workbook_without_table_has_no_header_candidate(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = write_workbook(Path(tmp) / "empty.xlsx", [["only a title"]])
            reader = WorkbookReader(path)
            inspection = reader.inspect_sheet("序时账")
            self.assertEqual(inspection.header_candidates, ())


if __name__ == "__main__":
    unittest.main()
