from openpyxl.utils import get_column_letter


def _display_length(value):
    if value is None:
        return 0
    text = str(value)
    # Approximate CJK characters as double width.
    return sum(2 if ord(ch) > 255 else 1 for ch in text)


def apply(workbook, audit, config):
    if not config.get("enabled", True):
        return
    min_width = float(config.get("min_width", 8))
    max_width = float(config.get("max_width", 40))
    padding = float(config.get("padding", 2))
    scan_rows = max(1, int(config.get("scan_rows", 5000)))

    for ws in workbook.worksheets:
        max_col = ws.max_column or 0
        max_row = min(ws.max_row or 0, scan_rows)
        for col_idx in range(1, max_col + 1):
            letter = get_column_letter(col_idx)
            width = 0
            for row_idx in range(1, max_row + 1):
                width = max(width, _display_length(ws.cell(row_idx, col_idx).value))
            if width == 0:
                continue
            target_width = max(min_width, min(max_width, width + padding))
            before_width = ws.column_dimensions[letter].width
            if before_width != target_width:
                ws.column_dimensions[letter].width = target_width
                audit.add(
                    rule="autofit_columns",
                    sheet=ws.title,
                    target=f"{letter}:{letter}",
                    change_type="dimension",
                    before={"width": before_width},
                    after={"width": target_width},
                    reason=f"根据前 {max_row} 行内容自动调整列宽",
                )
