import json
import shutil
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from services.journal_entry.job_io import write_json
from services.journal_entry.preview import build_inspection, build_preview
from services.journal_entry.status import Reporter
from services.journal_entry.worker import run

FIXTURES = Path(__file__).resolve().parent / "fixtures"


class WorkerIntegrationTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.job_dir = Path(self.tmp.name) / "job"
        self.job_dir.mkdir()
        shutil.copyfile(
            FIXTURES / "journal_direction_amount.xlsx",
            self.job_dir / "source.xlsx",
        )
        write_json(
            self.job_dir / "meta.json",
            {"originalFilename": "synthetic-direction-amount.xlsx"},
        )
        self.request = {
            "sheetName": "序时账",
            "headerRow": 1,
            "amountLayout": "direction_amount",
            "mapping": {
                "date": 2,
                "voucher_id": 3,
                "account_code": 4,
                "account_name": 5,
                "description": 6,
                "currency": 7,
                "direction": 8,
                "original_amount": 10,
                "amount": 11,
            },
        }
        write_json(self.job_dir / "request.json", self.request)

    def tearDown(self):
        self.tmp.cleanup()

    def test_inspection_and_preview(self):
        inspection = build_inspection(self.job_dir)
        self.assertTrue(inspection["ok"])
        self.assertEqual(inspection["selectedSheet"], "序时账")
        self.assertEqual(inspection["headerRow"], 1)
        self.assertEqual(inspection["amountLayout"]["layout"], "direction_amount")

        preview = build_preview(self.job_dir)
        self.assertTrue(preview["ok"])
        self.assertTrue(preview["canStart"])
        self.assertEqual(preview["cleaningSummary"]["voucherCount"], 3)
        self.assertEqual(preview["cleaningSummary"]["lineCount"], 8)
        self.assertEqual(len(preview["sampleVouchers"]), 3)

    def test_signed_amount_inspection_worker_and_metadata(self):
        shutil.copyfile(
            FIXTURES / "journal_signed_amount.xlsx",
            self.job_dir / "source.xlsx",
        )
        self.request = {
            "sheetName": "序时账",
            "headerRow": 1,
            "amountLayout": "signed_amount",
            "positiveSide": "debit",
            "mapping": {
                "date": 1,
                "voucher_id": 2,
                "account_code": 3,
                "account_name": 4,
                "description": 5,
                "signed_amount": 6,
            },
        }
        write_json(self.job_dir / "request.json", self.request)

        inspection = build_inspection(self.job_dir)
        self.assertTrue(inspection["ok"])
        self.assertEqual(inspection["amountLayout"]["layout"], "signed_amount")
        self.assertEqual(inspection["amountLayout"]["positive_side"], "debit")
        self.assertEqual(inspection["explicitMapping"]["signed_amount"], 6)

        reporter = Reporter(self.job_dir)
        run(self.job_dir, reporter)
        payload = json.loads(
            (self.job_dir / "result.json").read_text(encoding="utf-8")
        )
        self.assertEqual(
            payload["input_metadata"]["amount_layout"], "signed_amount"
        )
        self.assertEqual(payload["input_metadata"]["positive_side"], "debit")
        voucher = next(
            item for item in payload["vouchers"] if item["voucher_id"] == "SA-001"
        )
        bank = next(
            line for line in voucher["lines"] if line["account_code"] == "1002"
        )
        self.assertEqual(bank["debit"], "1000.00")
        self.assertEqual(bank["credit"], "0.00")

        export_path = self.job_dir / "journal-entry-analysis-report.xlsx"
        self.assertTrue(export_path.exists())
        workbook = load_workbook(export_path, read_only=True, data_only=False)
        try:
            self.assertEqual(
                workbook.sheetnames,
                ["清洗分录", "凭证汇总", "业务分类", "对方科目", "异常凭证", "数据质量问题"],
            )
        finally:
            workbook.close()

    def test_worker_writes_result_and_export(self):
        reporter = Reporter(self.job_dir)
        run(self.job_dir, reporter)

        payload = json.loads((self.job_dir / "result.json").read_text(encoding="utf-8"))
        self.assertEqual(payload["status"], "completed")
        self.assertEqual(payload["cleaning_summary"]["voucher_count"], 3)
        self.assertEqual(payload["cleaning_summary"]["line_count"], 8)
        self.assertEqual(payload["voucher_summary"]["exact_balanced_voucher_count"], 3)
        self.assertEqual(payload["business_classification_summary"]["voucher_count"], 3)
        self.assertEqual(payload["contra_summary"]["analyzed_voucher_count"], 3)
        self.assertEqual(payload["anomaly_summary"]["voucher_count"], 3)
        self.assertEqual(len(payload["vouchers"]), 3)

        export_path = self.job_dir / "journal-entry-analysis-report.xlsx"
        self.assertTrue(export_path.exists())
        workbook = load_workbook(export_path, read_only=True, data_only=False)
        try:
            self.assertEqual(
                workbook.sheetnames,
                ["清洗分录", "凭证汇总", "业务分类", "对方科目", "异常凭证", "数据质量问题"],
            )
            self.assertGreaterEqual(workbook["清洗分录"].max_row, 9)
        finally:
            workbook.close()

        ui_dir = self.job_dir / "ui"
        for name in ("summary.json", "vouchers.ndjson", "findings.json", "details.bin", "details.idx.json"):
            self.assertTrue((ui_dir / name).exists(), name)
        ui_summary = json.loads((ui_dir / "summary.json").read_text(encoding="utf-8"))
        self.assertEqual(ui_summary["voucher_summary"], payload["voucher_summary"])
        self.assertEqual(ui_summary["anomaly_summary"], payload["anomaly_summary"])
        ui_rows = [
            json.loads(line)
            for line in (ui_dir / "vouchers.ndjson").read_text(encoding="utf-8").splitlines()
            if line
        ]
        self.assertEqual(len(ui_rows), len(payload["vouchers"]))
        self.assertEqual(
            {row["k"] for row in ui_rows},
            {voucher["voucher_key"] for voucher in payload["vouchers"]},
        )

        status = json.loads((self.job_dir / "status.json").read_text(encoding="utf-8"))
        self.assertEqual(status["status"], "completed")
        self.assertEqual(status["progress"], 100)
        self.assertEqual(status["result"]["voucher_count"], 3)
        self.assertEqual(status["result"]["raw_row_count"], 8)
        self.assertEqual(status["result"]["line_count"], 8)
        self.assertIn("business_classified_count", status["result"])
        self.assertIn("finding_voucher_count", status["result"])


if __name__ == "__main__":
    unittest.main()
