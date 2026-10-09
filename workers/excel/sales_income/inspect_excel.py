import sys
import json
from datetime import date, datetime, time

from openpyxl import load_workbook


MAX_ROWS = 30
MAX_COLS = 80


def display_value(value):
    """
    把 Excel 单元格转换成适合 JSON / 网页显示的字符串。
    """

    if value is None:
        return ""

    if isinstance(value, (datetime, date, time)):
        return value.isoformat()

    return str(value).strip()


def inspect_row(values, row_number):
    """
    分析某一行是否可能是表头。
    """

    visible_values = [
        display_value(value)
        for value in values[:MAX_COLS]
    ]

    # 去掉尾部完全为空的单元格
    while (
        visible_values
        and visible_values[-1] == ""
    ):
        visible_values.pop()

    non_empty = [
        value
        for value in visible_values
        if value != ""
    ]

    if not non_empty:
        return None

    string_count = sum(
        1
        for value in values[:len(visible_values)]
        if isinstance(value, str)
        and value.strip()
    )

    string_ratio = (
        string_count / len(non_empty)
        if non_empty
        else 0
    )

    # 表头一般：
    # 1. 非空字段较多
    # 2. 文本字段比例较高
    # 3. 通常出现在文件较前方
    score = (
        string_ratio * 100
        + len(non_empty) * 3
        - row_number * 0.05
    )

    return {
        "rowNumber": row_number,
        "values": visible_values,
        "nonEmpty": len(non_empty),
        "score": round(score, 3),
    }


def inspect_sheet(ws):
    """
    读取工作表前 MAX_ROWS 行。
    """

    rows = []

    max_row = min(
        ws.max_row or MAX_ROWS,
        MAX_ROWS,
    )

    for row_number, row in enumerate(
        ws.iter_rows(
            min_row=1,
            max_row=max_row,
            values_only=True,
        ),
        start=1,
    ):
        inspected = inspect_row(
            list(row),
            row_number,
        )

        if inspected:
            rows.append(inspected)

    candidates = [
        row
        for row in rows
        if row["nonEmpty"] >= 2
    ]

    if candidates:
        detected = max(
            candidates,
            key=lambda item: (
                item["score"],
                -item["rowNumber"],
            ),
        )

        detected_header_row = (
            detected["rowNumber"]
        )
    else:
        detected_header_row = 1

    return {
        "name": ws.title,
        "detectedHeaderRow": detected_header_row,
        "rows": rows,
    }


def inspect_workbook(input_file):
    workbook = load_workbook(
        input_file,
        read_only=True,
        data_only=True,
    )

    sheets = []

    for ws in workbook.worksheets:
        sheets.append(
            inspect_sheet(ws)
        )

    return {
        "ok": True,
        "sheets": sheets,
    }


def main():
    if len(sys.argv) != 2:
        raise SystemExit(
            "Usage: inspect_excel.py input.xlsx"
        )

    input_file = sys.argv[1]

    result = inspect_workbook(
        input_file
    )

    print(
        json.dumps(
            result,
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
