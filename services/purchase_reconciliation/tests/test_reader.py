from __future__ import annotations

import tempfile
import unittest
from datetime import date
from pathlib import Path

from openpyxl import Workbook

from services.purchase_reconciliation.column_mapper import build_records, detect_mapping
from services.purchase_reconciliation.excel_reader import inspect_workbook, read_rows
from services.purchase_reconciliation.schema import FIELD_LABELS


def write_workbook(path: Path, rows, header_row: int = 1, headers=None, prefix_rows=None) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "DATABASE"
    if headers is None:
        headers = ["送货日期", "供应商", "采购订单", "采购订单项目", "物料编码", "订单单位", "数量", "单价", "每", "总金额", "移动类型"]
    for index, row in enumerate(prefix_rows or [], start=1):
        for column, value in enumerate(row, start=1):
            sheet.cell(row=index, column=column, value=value)
    for column, value in enumerate(headers, start=1):
        sheet.cell(row=header_row, column=column, value=value)
    for index, row in enumerate(rows, start=header_row + 1):
        for column, value in enumerate(row, start=1):
            sheet.cell(row=index, column=column, value=value)
    workbook.save(path)


class ExcelReaderTest(unittest.TestCase):
    def test_reads_populated_rows_and_converts_dates(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "sample.xlsx"
            write_workbook(path, [
                ["2026-01-02", "V1", "PO1", "10", "M1", "KG", 10, 10, 1, 100, "101"],
                ["2026-01-03", "V1", "PO1", "10", "M1", "KG", -4, 10, 1, -40, "122"],
            ])
            headers, rows, sheet = read_rows(path, header_row=1)
            self.assertEqual(sheet, "DATABASE")
            self.assertEqual(len(rows), 2)
            self.assertEqual(rows[0]["_source_row"], 2)
            self.assertEqual(rows[0]["送货日期"], "2026-01-02")
            self.assertEqual(rows[1]["数量"], -4)
            self.assertEqual(headers[0], "送货日期")

    def test_date_header_accepts_excel_date_objects(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "date-object.xlsx"
            write_workbook(path, [
                [date(2026, 3, 1), "V1", "PO1", "10", "M1", "KG", 10, 10, 1, 100, "101"],
            ])
            _headers, rows, _sheet = read_rows(path, header_row=1)
            self.assertEqual(rows[0]["送货日期"], "2026-03-01")

    def test_header_row_detection_and_count(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "offset-header.xlsx"
            write_workbook(
                path,
                [
                    ["2026-01-02", "V1", "PO1", "10", "M1", "KG", 10, 10, 1, 100, "101"],
                    ["2026-01-03", "V1", "PO1", "10", "M1", "KG", -10, 10, 1, -100, "122"],
                ],
                header_row=3,
                prefix_rows=[["SAP 导出"], ["导出日期", "2026-01-01"]],
            )
            inspection = inspect_workbook(path)
            self.assertEqual(inspection["header_row"], 3)
            self.assertEqual(inspection["data_row_count"], 2)
            self.assertEqual(inspection["headers"][0], "送货日期")

    def test_blank_rows_before_trailing_end_are_skipped(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "gaps.xlsx"
            write_workbook(path, [
                ["2026-01-02", "V1", "PO1", "10", "M1", "KG", 10, 10, 1, 100, "101"],
                [None] * 11,
                [None] * 11,
                ["2026-01-03", "V1", "PO1", "10", "M1", "KG", -10, 10, 1, -100, "122"],
            ])
            _headers, rows, _sheet = read_rows(path, header_row=1)
            self.assertEqual(len(rows), 2)

    def test_alias_detection_in_english_export(self):
        headers = ["Delivery Date", "Vendor", "PO", "Item", "Material", "Order Unit", "Qty", "Unit Price", "Per", "Amount", "Movement Type"]
        detection = detect_mapping(headers)
        self.assertEqual(detection["missing"], [])
        self.assertEqual(detection["mapping"]["quantity"], "Qty")
        self.assertEqual(detection["mapping"]["purchase_order_item"], "Item")
        self.assertEqual(detection["mapping"]["amount"], "Amount")

    def test_build_records_reports_invalid_quantity_row(self):
        headers = ["送货日期", "供应商", "采购订单", "采购订单项目", "物料编码", "订单单位", "数量", "单价", "每", "总金额", "移动类型"]
        mapping = detect_mapping(headers)["mapping"]
        rows = [dict(zip(headers, ["2026-01-02", "V1", "PO1", "10", "M1", "KG", "abc", 10, 1, 100, "101"]), _source_row=2)]
        with self.assertRaises(ValueError) as context:
            build_records(headers, rows, mapping)
        self.assertIn("第 2 行", str(context.exception))

    def test_missing_field_detection_uses_labels(self):
        detection = detect_mapping(["供应商", "数量"])
        self.assertIn("purchase_order", detection["missing"])
        self.assertIn("采购订单", detection["missing_labels"])
        self.assertIn("order_unit", FIELD_LABELS)


if __name__ == "__main__":
    unittest.main()
