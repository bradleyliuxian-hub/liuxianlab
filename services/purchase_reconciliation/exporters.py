"""CSV ZIP and XLSX exporters for reconciliation results."""

from __future__ import annotations

import csv
import io
import json
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any, Dict, Iterable, List, Sequence
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .engine import ReconciliationResult, to_decimal

SEGMENT_COLUMNS = [
    ("match_id", "匹配ID"),
    ("business_key", "业务键"),
    ("positive_source_row", "正向原始行号"),
    ("negative_source_row", "负向原始行号"),
    ("positive_movement_type", "正向移动类型"),
    ("negative_movement_type", "负向移动类型"),
    ("positive_date", "正向日期"),
    ("negative_date", "负向日期"),
    ("positive_quantity", "正向数量"),
    ("negative_quantity", "负向数量"),
    ("matched_quantity", "本次对抵数量"),
    ("matched_amount", "本次对抵金额"),
    ("negative_amount_abs", "负向对应金额绝对值"),
]

AZURE = "1F4E78"

METRIC_LABELS = [
    ("original_record_count", "原始记录数"),
    ("match_segment_count", "对抵配对段数"),
    ("matched_record_count", "涉及对抵原始记录数"),
    ("remaining_record_count", "对抵后剩余记录数"),
    ("original_positive_amount", "原始正向金额"),
    ("original_negative_amount_abs", "原始负向金额绝对值"),
    ("matched_amount", "成功对抵金额"),
    ("unmatched_negative_amount_abs", "未匹配负向金额绝对值"),
    ("unmatched_negative_quantity_abs", "未匹配负向数量绝对值"),
    ("final_net_amount", "最终净金额"),
    ("conservation_difference", "金额守恒差异"),
    ("positive_record_count", "正向记录数"),
    ("negative_record_count", "负向记录数"),
    ("zero_quantity_record_count", "零数量记录数"),
    ("fully_matched_record_count", "完全对抵记录数"),
    ("business_key_count", "业务键数量"),
]

MOVEMENT_COLUMNS = [
    ("movement_type", "移动类型"),
    ("record_count", "记录数"),
    ("original_quantity", "原始数量"),
    ("original_amount", "原始金额"),
    ("matched_quantity", "已对抵数量"),
    ("remaining_quantity", "剩余数量"),
    ("remaining_amount", "剩余金额"),
]

LOGIC_LINES = [
    ("方向判断", "数量>0 为正向；数量<0 为负向；数量=0 不参与对抵。方向只由数量决定，移动类型仅用于审计。"),
    ("匹配键", "默认：供应商 + 采购订单 + 采购订单项目 + 物料编码 + 订单单位 + 单价 + 每。"),
    ("匹配限制", "只在同一业务键内互抵；不跨订单、行项目、物料、单位或价格凑金额。"),
    ("排序", "同一业务键内按送货日期、原始行号升序执行 FIFO。"),
    ("对抵规则", "支持一对一、一对多、多对一与部分对抵。"),
    ("剩余数量", "正向剩余 = 原数量 - 已对抵数量；负向剩余 = 原数量 + 已对抵数量。"),
    ("剩余金额", "剩余金额 = 原始金额 × 剩余数量 ÷ 原始数量，不重新猜测价格。"),
    ("审计锚点", "每个配对段保留匹配ID与正负原始行号；每条结果保留原始字段与原始行号。"),
    ("金额校验", "原始净金额 = 原始正向金额 - 原始负向金额绝对值；应与最终净金额一致。"),
]


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, Decimal):
        return format(value, "f")
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, (list, tuple)):
        return "、".join(_text(item) for item in value)
    return str(value)


def _json_scalar(value: Any) -> Any:
    if isinstance(value, Decimal):
        return format(value, "f")
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return value


def segment_payload(segment: Dict[str, Any]) -> Dict[str, Any]:
    return {
        "matchId": segment["match_id"],
        "businessKey": segment["business_key"],
        "positiveSourceRow": segment["positive_source_row"],
        "negativeSourceRow": segment["negative_source_row"],
        "positiveMovementType": segment["positive_movement_type"],
        "negativeMovementType": segment["negative_movement_type"],
        "positiveDate": _json_scalar(segment["positive_date"]),
        "negativeDate": _json_scalar(segment["negative_date"]),
        "positiveQuantity": _json_scalar(segment["positive_quantity"]),
        "negativeQuantity": _json_scalar(segment["negative_quantity"]),
        "matchedQuantity": _json_scalar(segment["matched_quantity"]),
        "matchedAmount": _json_scalar(segment["matched_amount"]),
        "negativeAmountAbs": _json_scalar(segment["negative_amount_abs"]),
    }


def matched_payload(row: Dict[str, Any]) -> Dict[str, Any]:
    return {
        "sourceRow": row["source_row"],
        "direction": row["direction"],
        "originalQuantity": _json_scalar(row["original_quantity"]),
        "matchedQuantity": _json_scalar(row["matched_quantity"]),
        "remainingQuantity": _json_scalar(row["remaining_quantity"]),
        "originalAmount": _json_scalar(row["original_amount"]),
        "remainingAmount": _json_scalar(row["remaining_amount"]),
        "matchedAmount": _json_scalar(row["matched_amount"]),
        "businessKey": row["business_key"],
        "matchIds": list(row["match_ids"]),
        "raw": {key: _json_scalar(value) for key, value in (row.get("raw") or {}).items()},
    }


def remaining_payload(row: Dict[str, Any]) -> Dict[str, Any]:
    return {
        "sourceRow": row["source_row"],
        "direction": row["direction"],
        "originalQuantity": _json_scalar(row["original_quantity"]),
        "matchedQuantity": _json_scalar(row["matched_quantity"]),
        "remainingQuantity": _json_scalar(row["remaining_quantity"]),
        "originalAmount": _json_scalar(row["original_amount"]),
        "remainingAmount": _json_scalar(row["remaining_amount"]),
        "businessKey": row["business_key"],
        "matchIds": list(row["match_ids"]),
        "raw": {key: _json_scalar(value) for key, value in (row.get("raw") or {}).items()},
    }


def write_ndjson(path: str | Path, rows: Iterable[Dict[str, Any]]) -> int:
    count = 0
    with open(path, "w", encoding="utf-8") as handle:
        for row in rows:
            handle.write(json.dumps(row, ensure_ascii=False, separators=(",", ":")))
            handle.write("\n")
            count += 1
    return count


def write_csv_zip(path: str | Path, result: ReconciliationResult) -> None:
    headers = result.headers

    def write_csv(archive: ZipFile, name: str, header: Sequence[str], rows: Iterable[Sequence[Any]]) -> None:
        buffer = io.StringIO(newline="")
        writer = csv.writer(buffer)
        writer.writerow(header)
        for row in rows:
            writer.writerow([_text(value) for value in row])
        # UTF-8 BOM keeps Chinese headers readable in Excel on Windows.
        archive.writestr(name, "\ufeff" + buffer.getvalue())

    with ZipFile(path, "w", ZIP_DEFLATED, compresslevel=6) as archive:
        write_csv(
            archive,
            "对抵分段.csv",
            [label for _key, label in SEGMENT_COLUMNS],
            (
                [segment.get(key) for key, _label in SEGMENT_COLUMNS]
                for segment in result.segments
            ),
        )

        detail_header = ["匹配ID", "方向", "已对抵数量", "已对抵金额", "原始行号"] + headers
        write_csv(
            archive,
            "已对抵明细.csv",
            detail_header,
            (
                [
                    "、".join(record["match_ids"]),
                    record["direction"],
                    record["matched_quantity"],
                    record["matched_amount"],
                    record["source_row"],
                ] + [record.get("raw", {}).get(header) for header in headers]
                for record in result.matched_records
            ),
        )

        remaining_header = headers + [
            "原始行号", "原始数量", "原始总金额", "已对抵数量", "剩余数量", "剩余金额", "匹配ID",
        ]
        write_csv(
            archive,
            "对抵后剩余数据.csv",
            remaining_header,
            (
                [record.get("raw", {}).get(header) for header in headers]
                + [
                    record["source_row"],
                    record["original_quantity"],
                    record["original_amount"],
                    record["matched_quantity"],
                    record["remaining_quantity"],
                    record["remaining_amount"],
                    "、".join(record["match_ids"]),
                ]
                for record in result.remaining_rows
            ),
        )

        write_csv(
            archive,
            "Dashboard_指标.csv",
            ["指标", "结果"],
            ([label, result.metrics.get(key)] for key, label in METRIC_LABELS),
        )
        write_csv(
            archive,
            "Dashboard_移动类型.csv",
            [label for _key, label in MOVEMENT_COLUMNS],
            (
                [item.get(key) for key, _label in MOVEMENT_COLUMNS]
                for item in result.movement_summary
            ),
        )
        archive.writestr(
            "逻辑说明.txt",
            "\n".join(f"{title}：{body}" for title, body in LOGIC_LINES) + "\n",
        )


def write_xlsx(path: str | Path, result: ReconciliationResult, source_name: str = "") -> None:
    workbook = Workbook(write_only=True)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor=AZURE)

    def append_header(sheet, values: Sequence[Any]) -> None:
        cells: List[WriteOnlyCell] = []
        for value in values:
            cell = WriteOnlyCell(sheet, value=_excel_value(value))
            cell.font = header_font
            cell.fill = header_fill
            cells.append(cell)
        sheet.append(cells)

    logic_sheet = workbook.create_sheet("逻辑说明")
    append_header(logic_sheet, ["项目", "说明"])
    for title, body in LOGIC_LINES:
        logic_sheet.append([title, body])
    if source_name:
        logic_sheet.append(["来源文件", source_name])
    _set_widths(logic_sheet, [22, 110])
    logic_sheet.freeze_panes = "A2"

    dashboard = workbook.create_sheet("Dashboard")
    append_header(dashboard, ["指标", "结果"])
    for key, label in METRIC_LABELS:
        dashboard.append([label, _excel_value(result.metrics.get(key))])
    dashboard.append([])
    append_header(dashboard, [label for _key, label in MOVEMENT_COLUMNS])
    for item in result.movement_summary:
        dashboard.append([_excel_value(item.get(key)) for key, _label in MOVEMENT_COLUMNS])
    _set_widths(dashboard, [34, 20, 18, 18, 18, 18, 18])
    dashboard.freeze_panes = "A2"

    segment_sheet = workbook.create_sheet("对抵分段")
    append_header(segment_sheet, [label for _key, label in SEGMENT_COLUMNS])
    for segment in result.segments:
        segment_sheet.append([_excel_value(segment.get(key)) for key, _label in SEGMENT_COLUMNS])
    _set_widths(segment_sheet, [16, 46, 16, 16, 16, 16, 14, 14, 14, 14, 14, 16, 18])
    segment_sheet.freeze_panes = "A2"

    matched_sheet = workbook.create_sheet("已对抵明细")
    matched_header = ["匹配ID", "方向", "已对抵数量", "已对抵金额", "原始行号"] + result.headers
    append_header(matched_sheet, matched_header)
    for record in result.matched_records:
        matched_sheet.append(
            [
                "、".join(record["match_ids"]),
                record["direction"],
                _excel_value(record["matched_quantity"]),
                _excel_value(record["matched_amount"]),
                record["source_row"],
            ] + [_excel_value(record.get("raw", {}).get(header)) for header in result.headers]
        )
    _set_widths(matched_sheet, [18, 10, 16, 16, 14] + [16] * max(0, len(result.headers) - 3))
    matched_sheet.freeze_panes = "A2"

    remaining_sheet = workbook.create_sheet("对抵后剩余数据")
    remaining_header = result.headers + [
        "原始行号", "原始数量", "原始总金额", "已对抵数量", "剩余数量", "剩余金额", "匹配ID",
    ]
    append_header(remaining_sheet, remaining_header)
    for record in result.remaining_rows:
        remaining_sheet.append(
            [_excel_value(record.get("raw", {}).get(header)) for header in result.headers]
            + [
                record["source_row"],
                _excel_value(record["original_quantity"]),
                _excel_value(record["original_amount"]),
                _excel_value(record["matched_quantity"]),
                _excel_value(record["remaining_quantity"]),
                _excel_value(record["remaining_amount"]),
                "、".join(record["match_ids"]),
            ]
        )
    _set_widths(
        remaining_sheet,
        [16] * max(0, len(result.headers)) + [14, 14, 16, 14, 14, 16, 22],
    )
    remaining_sheet.freeze_panes = "A2"

    for sheet in (logic_sheet, dashboard, segment_sheet, matched_sheet, remaining_sheet):
        sheet.sheet_view.showGridLines = True

    workbook.save(str(path))


def _set_widths(sheet, widths: Sequence[int]) -> None:
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def _excel_value(value: Any) -> Any:
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, (list, tuple)):
        return "、".join(_text(item) for item in value)
    return value
