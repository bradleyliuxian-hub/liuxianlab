"""Server-side Excel export for Journal Entry Analysis."""

from __future__ import annotations

import json
from decimal import Decimal
from pathlib import Path
from typing import Any, Iterable, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


def _literal(cell, value):
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"
    return cell


def _amount(value):
    if value in (None, ""):
        return None
    try:
        return Decimal(str(value))
    except Exception:
        return value


def _json(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True)


HEADER_FILL = PatternFill("solid", fgColor="E4F3CB")


def _write_sheet(worksheet, headers: Sequence[str], rows: Iterable[Sequence[Any]]):
    for column, header in enumerate(headers, start=1):
        cell = worksheet.cell(row=1, column=column)
        cell.value = header
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
    for row_index, row in enumerate(rows, start=2):
        for column, value in enumerate(row, start=1):
            _literal(worksheet.cell(row=row_index, column=column), value)
    worksheet.freeze_panes = "A2"
    if worksheet.max_row >= 1:
        worksheet.auto_filter.ref = worksheet.dimensions
    for column in range(1, len(headers) + 1):
        width = max(
            len(str(headers[column - 1])),
            min(48, max((len(str(worksheet.cell(row=row, column=column).value or "")) for row in range(2, min(worksheet.max_row, 200) + 1)), default=8)),
        )
        worksheet.column_dimensions[get_column_letter(column)].width = max(10, width + 2)


def build_export(path: Path, payload: Any) -> None:
    workbook = Workbook()
    default = workbook.active
    workbook.remove(default)

    vouchers = list(payload.get("vouchers") or [])

    lines_sheet = workbook.create_sheet("清洗分录")
    line_rows = []
    for voucher in vouchers:
        for line in voucher.get("lines") or []:
            line_rows.append([
                voucher.get("voucher_key"),
                voucher.get("voucher_id"),
                voucher.get("date"),
                voucher.get("period"),
                voucher.get("entity"),
                line.get("account_code"),
                line.get("account_name"),
                line.get("detail_account"),
                _amount(line.get("debit")),
                _amount(line.get("credit")),
                line.get("description"),
                line.get("currency"),
                _amount(line.get("original_amount")),
                line.get("sheet_name"),
                line.get("source_row"),
                line.get("internal_line_key"),
                line.get("line_id"),
                "balanced" if voucher.get("exactly_balanced") else ("within_tolerance" if voucher.get("is_within_tolerance") else "unbalanced"),
                ", ".join(voucher.get("validation_flags") or []),
            ])
    _write_sheet(
        lines_sheet,
        [
            "凭证Key", "凭证号", "日期", "期间", "主体", "科目编码", "科目名称",
            "明细科目", "借方", "贷方", "摘要", "币种", "原币",
            "原始Sheet", "原始行号", "内部行Key", "业务行ID", "平衡状态", "凭证校验",
        ],
        line_rows,
    )

    summary_sheet = workbook.create_sheet("凭证汇总")
    summary_rows = []
    for voucher in vouchers:
        business = voucher.get("business_classification") or {}
        contra = voucher.get("contra") or {}
        anomaly = voucher.get("anomaly") or {}
        summary_rows.append([
            voucher.get("voucher_key"),
            voucher.get("voucher_id"),
            voucher.get("date"),
            voucher.get("period"),
            voucher.get("entity"),
            voucher.get("line_count"),
            _amount(voucher.get("total_debit")),
            _amount(voucher.get("total_credit")),
            _amount(voucher.get("imbalance")),
            "YES" if voucher.get("exactly_balanced") else "NO",
            "YES" if voucher.get("is_within_tolerance") else "NO",
            ", ".join(voucher.get("validation_flags") or []),
            business.get("business_code"),
            business.get("business_label"),
            _amount(business.get("score")),
            _amount(business.get("confidence")),
            business.get("confidence_level"),
            _amount(anomaly.get("risk_score")),
            anomaly.get("risk_level"),
            "YES" if contra.get("ambiguous") else "NO",
            "YES" if contra.get("complexity_limited") else "NO",
            len(contra.get("unallocated_debit") or []) + len(contra.get("unallocated_credit") or []),
        ])
    _write_sheet(
        summary_sheet,
        [
            "凭证Key", "凭证号", "日期", "期间", "主体", "行数",
            "借方合计", "贷方合计", "差额", "严格平衡", "容差内",
            "校验标记", "业务分类", "业务名称", "分类分数", "分类置信度",
            "置信等级", "风险分数", "风险等级", "Contra歧义", "复杂度限制",
            "未分配节点数",
        ],
        summary_rows,
    )

    business_sheet = workbook.create_sheet("业务分类")
    business_rows = []
    for voucher in vouchers:
        business = voucher.get("business_classification")
        if not business:
            continue
        business_rows.append([
            voucher.get("voucher_key"),
            voucher.get("voucher_id"),
            business.get("business_code"),
            business.get("business_label"),
            _amount(business.get("score")),
            _amount(business.get("confidence")),
            business.get("confidence_level"),
            "YES" if business.get("is_hard_rule") else "NO",
            business.get("reason"),
            _json(business.get("candidates") or []),
            _json(business.get("evidence") or {}),
        ])
    _write_sheet(
        business_sheet,
        [
            "凭证Key", "凭证号", "业务分类", "业务名称", "分数", "置信度",
            "置信等级", "硬规则", "原因", "候选分类", "证据",
        ],
        business_rows,
    )

    contra_sheet = workbook.create_sheet("对方科目")
    contra_rows = []
    for voucher in vouchers:
        contra = voucher.get("contra") or {}
        for allocation in contra.get("allocations") or []:
            contra_rows.append([
                voucher.get("voucher_key"),
                voucher.get("voucher_id"),
                allocation.get("source_internal_key"),
                allocation.get("source_line_id"),
                allocation.get("source_side"),
                allocation.get("target_internal_key"),
                allocation.get("target_line_id"),
                allocation.get("target_side"),
                _amount(allocation.get("amount")),
                allocation.get("method"),
                _amount(allocation.get("confidence")),
                "allocation",
            ])
        for item in contra.get("unallocated_debit") or []:
            contra_rows.append([
                voucher.get("voucher_key"),
                voucher.get("voucher_id"),
                item.get("internal_line_key"),
                item.get("line_id"),
                "debit",
                None,
                None,
                None,
                _amount(item.get("amount")),
                None,
                None,
                "unallocated_debit",
            ])
        for item in contra.get("unallocated_credit") or []:
            contra_rows.append([
                voucher.get("voucher_key"),
                voucher.get("voucher_id"),
                item.get("internal_line_key"),
                item.get("line_id"),
                "credit",
                None,
                None,
                None,
                _amount(item.get("amount")),
                None,
                None,
                "unallocated_credit",
            ])
    _write_sheet(
        contra_sheet,
        [
            "凭证Key", "凭证号", "来源内部Key", "来源业务行ID", "来源方向",
            "目标内部Key", "目标业务行ID", "目标方向", "金额", "方法",
            "置信度", "类型",
        ],
        contra_rows,
    )

    finding_sheet = workbook.create_sheet("异常凭证")
    finding_rows = []
    for finding in payload.get("findings") or []:
        finding_rows.append([
            finding.get("id"),
            finding.get("voucher_key"),
            finding.get("voucher_id"),
            finding.get("code"),
            finding.get("type"),
            finding.get("severity"),
            _amount(finding.get("risk_score")),
            finding.get("amount"),
            finding.get("business_group"),
            finding.get("reason"),
            _json(finding.get("evidence") or {}),
            ", ".join(str(value) for value in finding.get("sheet_names") or []),
            ", ".join(str(value) for value in finding.get("source_rows") or []),
        ])
    _write_sheet(
        finding_sheet,
        [
            "Finding ID", "凭证Key", "凭证号", "规则代码", "异常类型",
            "风险等级", "风险分数", "凭证金额", "业务组", "异常原因",
            "结构化证据", "原始Sheet", "原始行号",
        ],
        finding_rows,
    )

    quality_sheet = workbook.create_sheet("数据质量问题")
    quality_rows = []
    for issue in payload.get("validation_issues") or []:
        quality_rows.append([
            issue.get("sheet_name"),
            issue.get("source_row"),
            issue.get("field"),
            issue.get("code"),
            issue.get("message"),
            issue.get("raw_value"),
            issue.get("column_index"),
            "issue",
        ])
    for warning in payload.get("row_warnings") or []:
        quality_rows.append([
            warning.get("sheet_name"),
            warning.get("source_row"),
            warning.get("field"),
            warning.get("code"),
            warning.get("message"),
            warning.get("raw_value"),
            warning.get("column_index"),
            "warning",
        ])
    _write_sheet(
        quality_sheet,
        ["Sheet", "行号", "字段", "代码", "说明", "原始值", "列号", "级别"],
        quality_rows,
    )

    path = Path(path)
    tmp = path.with_name(f".{path.name}.tmp.xlsx")
    workbook.save(tmp)
    workbook.close()

    check = load_workbook(tmp, read_only=True, data_only=False)
    try:
        _ = check.sheetnames
    finally:
        check.close()
    tmp.replace(path)
